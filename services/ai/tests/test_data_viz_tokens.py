"""
Chart and table style tokens: what a deck records about how it draws data.

Colours are the smallest part of why a chart looks like it belongs to a deck.
Reference page 22's bar chart has no value axis and no gridlines - every bar
carries its own label instead - and the renderer used to draw a gridded axis
on every non-pie chart unconditionally. So a user could upload a template whose
charts have no axes and get charts with axes and gridlines in their brand
colours.

Everything here defaults to what was drawn before, which is what makes it
additive: the first test class is the regression gate for exactly that.
"""

import io
import os
import shutil
import tempfile

import numpy as np
import pytest
from PIL import Image
from pptx import Presentation

from app.design_system import derive_data_viz_config, get_default_design_system
from app.renderer.chart_renderer import render_chart_to_png
from app.renderer.pptx_renderer import PPTXRenderer
from app.renderer.table_style import (
    LABEL_MODES, LEGEND_STYLES, TABLE_TREATMENTS, UNIT_CONVENTIONS,
    resolve_chart_style, resolve_table_style,
)

CHART = {
    "chartType": "bar",
    "categories": ["2021", "2022", "2023", "2024"],
    "series": [{"name": "Revenue ($B)", "values": [29.1, 32.3, 36.0, 37.4]}],
}


def _system(**data_viz):
    base = get_default_design_system()
    return {**base, "dataViz": {**base["dataViz"], **data_viz}}


def _render(design_system=None, chart=None, width=6.0, height=3.4):
    png = render_chart_to_png(chart or dict(CHART), design_system, width, height)
    assert png, "the chart did not render"
    return np.asarray(Image.open(io.BytesIO(png)).convert("RGBA"))


def _ink(pixels, threshold=40):
    """A boolean mask of where the transparent PNG actually has marks."""
    return pixels[:, :, 3] > threshold


def _thin_runs(indices, max_thickness=3):
    """Consecutive indices collapsed into lines, dropping thick slabs."""
    lines, run = 0, []
    for index in indices:
        if run and index == run[-1] + 1:
            run.append(index)
            continue
        if run and len(run) <= max_thickness:
            lines += 1
        run = [index]
    if run and len(run) <= max_thickness:
        lines += 1
    return lines


def _rules_of_colour(pixels, hex_value, axis, span=0.35):
    """
    How many thin rules of a given colour span most of the image.

    Colour rather than geometry, and for a specific reason: with the axis and
    gridlines gone, matplotlib's constrained layout grows the bars to fill the
    plot, so most rows *and* columns through them read as "mostly inked" and
    the anti-aliased edge of a bar is a thin full-length run indistinguishable
    from a rule. Both the gridline and axis colours are design-system tokens,
    so matching them says exactly what is being counted.
    """
    from app import colors as color_utils

    target = color_utils.parse_hex(hex_value)
    if target is None:
        return 0
    rgb = pixels[:, :, :3].astype(int)
    close = (np.abs(rgb - np.array(target)).sum(axis=2) < 60) & _ink(pixels)
    extent = close.shape[1] if axis == 1 else close.shape[0]
    runs = np.nonzero(close.sum(axis=axis) >= extent * span)[0]
    return _thin_runs(list(runs), max_thickness=4)


def _gridlines(pixels, design_system, span=0.35):
    """Horizontal rules drawn in the deck's gridline colour."""
    from app.renderer.table_style import resolve_chart_style

    return _rules_of_colour(pixels, resolve_chart_style(design_system).grid_color,
                            axis=1, span=span)


def _axis_rules(pixels, design_system, span=0.5):
    """Vertical rules drawn in the deck's axis colour - i.e. a value axis."""
    from app.renderer.table_style import resolve_chart_style

    return _rules_of_colour(pixels, resolve_chart_style(design_system).axis_color,
                            axis=0, span=span)


def _left_band_ink(pixels, fraction=0.05):
    """Ink in the leftmost slice, where the rotated axis title lives."""
    ink = _ink(pixels)
    return int(ink[:, : max(1, int(ink.shape[1] * fraction))].sum())


def _bar_columns(pixels):
    """
    Columns carrying ink through the middle of the plot, which is bars only.

    Taken from a band across the middle rather than the bottom half: the
    bottom carries the baseline and the category labels, which are there at
    any bar thickness.
    """
    ink = _ink(pixels)
    height = ink.shape[0]
    middle = ink[int(height * 0.45):int(height * 0.55), :]
    return int((middle.sum(axis=0) > 0).sum())


def _difference(a, b, threshold=24):
    """The fraction of pixels that differ between two renderings."""
    return float((np.abs(a.astype(int) - b.astype(int)).sum(axis=2)
                  > threshold).mean())


class TestTheDefaultsAreWhatWasDrawnBefore:
    """The regression gate: a deck that declares nothing is unchanged."""

    def test_a_design_system_with_no_tokens_keeps_every_old_behaviour(self):
        style = resolve_chart_style(None)

        assert style.show_value_axis is True
        assert style.show_category_axis is True
        assert style.show_value_gridlines is True
        assert style.show_category_gridlines is False
        assert style.label_mode == "auto"
        assert style.legend_style == "auto"
        assert style.unit_convention == "axis"
        assert style.bar_thickness == pytest.approx(0.78)

    def test_a_table_with_no_tokens_is_still_banded(self):
        assert resolve_table_style(None).treatment == "filled"
        assert resolve_table_style(None).header_filled is True

    def test_the_deterministic_path_still_returns_a_complete_set(self):
        """derive_data_viz_config must stay usable with no model call."""
        config = derive_data_viz_config(
            get_default_design_system()["palette"]["colors"], "Arial"
        )

        for key in ("chartColorSequence", "chartAxisColor", "tableHeaderBg",
                    "chartValueAxis", "chartLabelMode", "chartBarThickness",
                    "tableTreatment"):
            assert key in config, f"{key} is missing from the derived config"

    def test_the_default_design_system_publishes_them(self):
        data_viz = get_default_design_system()["dataViz"]

        assert data_viz["chartValueGridlines"] is True
        assert data_viz["tableTreatment"] == "filled"


class TestTokensAreReadRatherThanHardcoded:
    def test_gridlines_can_be_switched_off(self):
        # A line chart, because bars sit *over* the gridlines (set_axisbelow)
        # and hide most of them - which would measure the bars rather than the
        # tokens.
        line = {**CHART, "chartType": "line"}
        on = _system(chartValueGridlines=True)
        off = _system(chartValueGridlines=False)

        assert _gridlines(_render(on, chart=dict(line)), on) >= 3
        assert _gridlines(_render(off, chart=dict(line)), off) == 0

    def test_the_value_axis_can_be_dropped(self):
        """
        Reference page 22 drops it entirely. This is the single biggest
        visual difference between our charts and that deck's.
        """
        on = _system(chartValueAxis=True)
        off = _system(chartValueAxis=False, chartValueGridlines=False)

        assert _axis_rules(_render(on), on) >= 1
        assert _axis_rules(_render(off), off) == 0

    def test_bar_thickness_changes_the_bars(self):
        # Gridlines off, so the only ink across the middle of the plot is the
        # bars themselves.
        thin = _bar_columns(_render(_system(chartBarThickness=0.35,
                                            chartValueGridlines=False)))
        thick = _bar_columns(_render(_system(chartBarThickness=0.95,
                                             chartValueGridlines=False)))

        assert thin < thick

    def test_the_unit_can_move_off_the_axis(self):
        """
        The reference states its unit on a second subtitle line ("thousands",
        "$ billion") rather than rotated up the side of the plot.
        """
        on_axis = _left_band_ink(_render(_system(chartUnitConvention="axis")), 0.03)
        off_axis = _left_band_ink(
            _render(_system(chartUnitConvention="subtitle")), 0.03
        )

        assert off_axis < on_axis

    def test_data_labels_can_be_switched_off(self):
        labelled = _render(_system(chartLabelMode="all"))
        bare = _render(_system(chartLabelMode="none"))

        assert _difference(labelled, bare) > 0.001

    def test_only_the_terminal_point_can_be_labelled(self):
        """Reference page 110's line chart labels its final point and no other."""
        line = {**CHART, "chartType": "line"}
        every = _render(_system(chartLabelMode="all"), chart=dict(line))
        terminal = _render(_system(chartLabelMode="terminal"), chart=dict(line))
        none = _render(_system(chartLabelMode="none"), chart=dict(line))

        # One label is closer to no labels than four labels are.
        assert _difference(terminal, none) < _difference(every, none)
        assert _difference(terminal, every) > 0.0

    def test_a_declared_label_mode_survives_a_crowded_chart(self):
        """
        "auto" drops labels once they would smear. A deck whose convention is
        "label the end of every bar" means it however many bars there are.
        """
        crowded = {
            "chartType": "bar",
            "categories": [str(2000 + n) for n in range(14)],
            "series": [{"name": "Revenue", "values": [float(n) for n in range(14)]}],
        }
        auto = _render(_system(chartLabelMode="auto"), chart=dict(crowded))
        declared = _render(_system(chartLabelMode="end"), chart=dict(crowded))
        suppressed = _render(_system(chartLabelMode="none"), chart=dict(crowded))

        # "auto" gave up on this many bars, so it matches the suppressed
        # rendering; the declared mode does not.
        assert _difference(auto, suppressed) == pytest.approx(0.0, abs=0.0005)
        assert _difference(declared, suppressed) > 0.001

    def test_a_legend_can_be_suppressed(self):
        two_series = {
            "chartType": "bar", "categories": ["Q1", "Q2"],
            "series": [{"name": "Plan", "values": [8.6, 8.9]},
                       {"name": "Actual", "values": [8.7, 8.8]}],
        }
        with_legend = _render(_system(chartLegend="auto"), chart=dict(two_series))
        without = _render(_system(chartLegend="none"), chart=dict(two_series))

        assert _difference(with_legend, without) > 0.001


class TestTheReferenceDecksTokensProduceADifferentChart:
    """
    The acceptance criterion: a chart under reference-deck tokens is visually
    distinguishable from one under default tokens in an automated check, not
    just by eye.
    """

    REFERENCE_TOKENS = dict(
        chartValueAxis=False,
        chartValueGridlines=False,
        chartLabelMode="end",
        chartUnitConvention="subtitle",
        chartBarThickness=0.62,
    )

    def test_the_two_renderings_differ_substantially(self):
        default = _render(get_default_design_system())
        reference = _render(_system(**self.REFERENCE_TOKENS))

        assert default.shape == reference.shape
        assert _difference(default, reference) > 0.02, (
            "the tokens are not reaching the renderer"
        )

    def test_it_reads_as_no_axis_and_no_gridlines(self):
        system = _system(**self.REFERENCE_TOKENS)
        reference = _render(system)

        assert _axis_rules(reference, system) == 0
        assert _gridlines(reference, system) == 0


class TestTables:
    @pytest.fixture(scope="class")
    @staticmethod
    def work_dir():
        directory = tempfile.mkdtemp(prefix="ts_tokens_")
        try:
            yield directory
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @staticmethod
    def _render_table(work_dir, name, design_system):
        slides = [{
            "slideIndex": 0, "layoutId": "table",
            "contents": [
                {"type": "text", "runs": [{"text": "Segment performance"}]},
                {"type": "table", "table": {"headerRow": True, "rows": [
                    ["Segment", "Revenue"],
                    ["North America", "26.8"],
                    ["International", "7.6"]]}},
            ],
        }]
        path = os.path.join(work_dir, f"{name}.pptx")
        PPTXRenderer(template_path=None, color_theme="professional",
                     design_system=design_system).render(
            slides=slides, metadata={"title": "Tables"}, output_path=path)
        return Presentation(path).slides[0]

    def test_a_banded_table_fills_its_cells(self, work_dir):
        slide = self._render_table(work_dir, "banded", get_default_design_system())
        table = next(s.table for s in slide.shapes if s.has_table)

        assert table.cell(1, 0).fill.type is not None

    def test_a_hairline_table_leaves_its_cells_unfilled(self, work_dir):
        slide = self._render_table(work_dir, "hairline",
                                   _system(tableTreatment="hairline"))
        table = next(s.table for s in slide.shapes if s.has_table)

        from pptx.enum.dml import MSO_FILL

        assert table.cell(1, 0).fill.type == MSO_FILL.BACKGROUND

    def test_a_hairline_table_draws_its_own_rules(self, work_dir):
        from app.renderer.primitives import shape_primitive_id

        slide = self._render_table(work_dir, "ruled",
                                   _system(tableTreatment="hairline"))

        rules = [s for s in slide.shapes
                 if "table-rule" in (shape_primitive_id(s) or "")]
        assert len(rules) == 2  # under the header and under the first data row

    def test_a_banded_table_draws_no_rules(self, work_dir):
        from app.renderer.primitives import shape_primitive_id

        slide = self._render_table(work_dir, "norules", get_default_design_system())

        assert not [s for s in slide.shapes
                    if "table-rule" in (shape_primitive_id(s) or "")]


class TestTokenParsing:
    @pytest.mark.parametrize("value,expected", [
        (True, True), (False, False), ("true", True), ("no", False),
        ("hidden", False), ("visible", True), (None, True), ("nonsense", True),
    ])
    def test_a_flag_accepts_what_a_model_actually_writes(self, value, expected):
        style = resolve_chart_style(_system(chartValueAxis=value))

        assert style.show_value_axis is expected

    def test_an_unknown_enum_falls_back_rather_than_breaking(self):
        assert resolve_chart_style(_system(chartLabelMode="sideways")).label_mode \
            == "auto"
        assert resolve_table_style(_system(tableTreatment="etched")).treatment \
            == "filled"

    def test_a_negative_size_falls_back(self):
        assert resolve_chart_style(_system(chartBarThickness=-1)).bar_thickness \
            == pytest.approx(0.78)

    def test_the_vocabularies_are_what_the_extractor_offers(self):
        import inspect

        from app.agent.tools import extract_template_design_system as extractor

        source = inspect.getsource(extractor)
        for mode in LABEL_MODES - {"auto"}:
            assert f'"{mode}"' in source, f"{mode} is not offered to the extractor"
        for convention in UNIT_CONVENTIONS:
            assert f'"{convention}"' in source
        for treatment in TABLE_TREATMENTS:
            assert f'"{treatment}"' in source
        for legend in LEGEND_STYLES - {"auto"}:
            assert f'"{legend}"' in source


class TestExtractionMergesRatherThanReplaces:
    def test_a_partial_vision_answer_keeps_the_derived_defaults(self):
        """
        The model answers the colour questions reliably and the style ones
        partially. Replacing the derived set with a half-answer used to mean
        the style tokens were simply absent.
        """
        from app.agent.tools.extract_template_design_system import (
            ExtractTemplateDesignSystemTool,
        )

        parsed = {
            "palette": {"colors": [{"name": "primary", "value": "#00704A",
                                    "type": "hex"}], "schemeName": "x"},
            "dataViz": {"chartColorSequence": ["#00704A"],
                        "chartValueGridlines": False},
        }

        result = ExtractTemplateDesignSystemTool._parse_vision_response(
            __import__("json").dumps(parsed)
        )

        data_viz = result["dataViz"]
        assert data_viz["chartValueGridlines"] is False
        assert data_viz["chartBarThickness"] == pytest.approx(0.78)
        assert data_viz["tableTreatment"] == "filled"
