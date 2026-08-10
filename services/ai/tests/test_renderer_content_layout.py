"""
Tests for how the renderer shares a slide's content area between body text
and objects (charts, tables, images).

A contentType:"chart" slide still lists its key points, so it arrives with
body text AND a chart. Filling the single body placeholder with the text and
then falling the chart back to LayoutEngine's content rectangle drew the chart
directly on top of the bullets. Validation flagged the collision correctly,
but the content model could not fix it by rewording - the renderer recreated
it on every retry, which is what turned one bad slide into an endless redesign
loop that still ended in a download being offered.
"""

import os
import tempfile

from pptx import Presentation
from pptx.util import Inches

from app.renderer.pptx_renderer import PPTXRenderer

PICTURE_SHAPE_TYPE = 13


def _render(contents):
    slides = [{
        "slideIndex": 0, "layoutId": "1", "layoutName": "Default",
        "contents": contents,
    }]
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        output_path = f.name
    PPTXRenderer(template_path=None).render(
        slides=slides, metadata={"title": "T"}, output_path=output_path
    )
    try:
        return Presentation(output_path)
    finally:
        os.unlink(output_path)


def _bullets(n=5):
    return [
        {"text": f"Key point number {i} with some supporting detail",
         "bullet": True, "level": 0}
        for i in range(n)
    ]


def _title():
    return {"type": "text", "placeholderIdx": 0,
            "runs": [{"text": "A Title", "bold": True}]}


def _chart():
    return {"type": "chart", "placeholderIdx": 1, "chart": {
        "chartType": "bar", "categories": ["A", "B", "C"],
        "series": [{"name": "S", "values": [1.0, 2.0, 3.0]}]}}


class TestTextAndObjectsShareTheContentArea:
    def test_the_chart_does_not_overlap_the_body_text(self):
        slide = _render([_title(),
                         {"type": "text", "placeholderIdx": 1, "runs": _bullets()},
                         _chart()]).slides[0]

        body = next(s for s in slide.shapes
                    if s.has_text_frame and "Key point" in s.text_frame.text)
        picture = next(s for s in slide.shapes if s.shape_type == PICTURE_SHAPE_TYPE)

        body_bottom = body.top + body.height
        picture_bottom = picture.top + picture.height
        overlaps = body.top < picture_bottom and picture.top < body_bottom
        assert not overlaps, (
            f"chart ({picture.top}..{picture_bottom}) overlaps "
            f"body ({body.top}..{body_bottom})"
        )

    def test_every_shape_stays_inside_the_slide(self):
        slides = [{"slideIndex": 0, "layoutId": "1", "layoutName": "Default",
                   "contents": [_title(),
                                {"type": "text", "placeholderIdx": 1, "runs": _bullets()},
                                _chart()]}]
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name
        try:
            PPTXRenderer(template_path=None).render(
                slides=slides, metadata={"title": "T"}, output_path=output_path
            )
            prs = Presentation(output_path)
            for shape in prs.slides[0].shapes:
                assert shape.left >= 0 and shape.top >= 0
                assert shape.left + shape.width <= prs.slide_width
                assert shape.top + shape.height <= prs.slide_height, (
                    f"{shape.name!r} runs past the bottom of the slide"
                )
        finally:
            os.unlink(output_path)

    def test_a_table_alongside_text_also_gets_its_own_band(self):
        slide = _render([
            _title(),
            {"type": "text", "placeholderIdx": 1, "runs": _bullets(3)},
            {"type": "table", "placeholderIdx": 1,
             "table": {"headerRow": True, "rows": [["A", "B"], ["1", "2"]]}},
        ]).slides[0]

        body = next(s for s in slide.shapes
                    if s.has_text_frame and "Key point" in s.text_frame.text)
        table = next(s for s in slide.shapes if s.has_table)
        assert table.top >= body.top + body.height - 1, "table overlaps the bullets"

    def test_a_chart_alone_still_gets_the_whole_content_area(self):
        """The split only applies when text and an object actually compete - a
        chart-only slide must not be shrunk into a band."""
        slide = _render([_title(), _chart()]).slides[0]
        picture = next(s for s in slide.shapes if s.shape_type == PICTURE_SHAPE_TYPE)
        assert picture.height > Inches(2.5)

    def test_bullets_keep_their_placeholder_formatting_when_sharing(self):
        """The text band is the resized placeholder, not a bare text box, so
        it keeps the layout's typography and bullet definition."""
        slide = _render([_title(),
                         {"type": "text", "placeholderIdx": 1, "runs": _bullets(3)},
                         _chart()]).slides[0]
        body = next(s for s in slide.shapes
                    if s.has_text_frame and "Key point" in s.text_frame.text)
        assert body.is_placeholder
