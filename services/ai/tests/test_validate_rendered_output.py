"""
Tests for validate_rendered_output - the ground-truth check against the
actual LibreOffice-rendered PDF geometry (not an estimate).

Geometry-detection logic is tested against synthetic PDFs built directly
with PyMuPDF, so each case has exact, known bounding boxes - fast and
precise, independent of the renderer/LibreOffice. A slower end-to-end test
at the bottom exercises the real render -> PreviewGenerator -> validate
pipeline together.
"""

import time

import pytest
import fitz

from app.agent.tools.validate_rendered_output import ValidateRenderedOutputTool, _rect_intersection_area, _rect_area

PAGE_WIDTH_PT = 720.0   # 10in at 72pt/in
PAGE_HEIGHT_PT = 405.0  # 5.625in at 72pt/in


def _make_pdf(tmp_path, page_ops: list[list[tuple[str, float, float, float]]]) -> str:
    """page_ops: one list of (text, x, y, fontsize) insert_text calls per page."""
    doc = fitz.open()
    for ops in page_ops:
        page = doc.new_page(width=PAGE_WIDTH_PT, height=PAGE_HEIGHT_PT)
        for text, x, y, fontsize in ops:
            page.insert_text((x, y), text, fontsize=fontsize)
    path = str(tmp_path / f"test_{time.time_ns()}.pdf")
    doc.save(path)
    doc.close()
    return path


class TestGeometryOverflow:
    def test_text_within_bounds_produces_no_issues(self, tmp_path):
        pdf_path = _make_pdf(tmp_path, [[("Hello world", 50, 50, 14)]])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert issues == []

    def test_text_extending_past_the_right_edge_is_flagged(self, tmp_path):
        # Placed near the right edge with a long string at a large font size
        # so its bbox extends well past PAGE_WIDTH_PT.
        pdf_path = _make_pdf(tmp_path, [[
            ("This text is deliberately far too long to fit on the slide width", PAGE_WIDTH_PT - 40, 50, 24),
        ]])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert any(i["type"] == "overflow" and i["slideIndex"] == 0 for i in issues)

    def test_text_extending_past_the_bottom_edge_is_flagged(self, tmp_path):
        # A large font's descent pushes the glyph bbox well below the
        # baseline - placing the baseline just above the bottom edge makes
        # the actual rendered extent cross it. (A fully out-of-bounds
        # *origin* is clipped by PyMuPDF and never appears in extracted
        # text at all, per empirical check - this is why the baseline stays
        # in-bounds while the glyph extent is pushed out instead.)
        pdf_path = _make_pdf(tmp_path, [[("g", 50, PAGE_HEIGHT_PT - 2, 30)]])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert any(i["type"] == "overflow" for i in issues)

    def test_text_extending_past_the_top_edge_is_flagged(self, tmp_path):
        # Same technique as the bottom-edge case: a large font's ascent
        # pushes the glyph bbox above a baseline that is itself in-bounds.
        pdf_path = _make_pdf(tmp_path, [[("Near top", 10, 5, 30)]])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert any(i["type"] == "overflow" for i in issues)

    def test_slide_index_matches_the_page_the_issue_is_on(self, tmp_path):
        pdf_path = _make_pdf(tmp_path, [
            [("Fine on slide 0", 50, 50, 14)],
            [("This text is deliberately far too long to fit on the slide width", PAGE_WIDTH_PT - 40, 50, 24)],
        ])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert all(i["slideIndex"] == 1 for i in issues)
        assert len(issues) >= 1


class TestGeometryOverlap:
    def test_two_text_runs_at_nearly_the_same_position_are_flagged(self, tmp_path):
        # Two independent insert_text calls placed at near-identical
        # coordinates - a real overlap, like two placeholders both landing
        # on the same region. (PyMuPDF actually merges these into a single
        # "block" in its own text-extraction model - confirmed empirically
        # - which is exactly why the overlap check compares every span pair
        # rather than exempting same-block pairs; see the comment in
        # validate_rendered_output.py.)
        pdf_path = _make_pdf(tmp_path, [[
            ("First overlapping block of text", 50, 100, 18),
            ("Second overlapping block of text", 52, 102, 18),
        ]])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert any(i["type"] == "overlap" for i in issues)

    def test_well_separated_text_is_not_flagged(self, tmp_path):
        pdf_path = _make_pdf(tmp_path, [[
            ("Title up here", 50, 50, 18),
            ("Body content down here, well separated", 50, 300, 14),
        ]])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert not any(i["type"] == "overlap" for i in issues)

    def test_normal_multi_line_wrapping_is_not_flagged_as_overlap(self, tmp_path):
        """Consecutive lines of a normally-wrapped paragraph empirically
        overlap by ~13% (font ascent/descent bleeding into the adjacent
        line's box) - the area-ratio threshold must sit comfortably above
        that, or every multi-line text block would false-positive."""
        pdf_path = _make_pdf(tmp_path, [[
            (f"Line {i} of the same paragraph", 50, 60 + i * 14.4, 12) for i in range(5)
        ]])
        issues = ValidateRenderedOutputTool._check_geometry(pdf_path)
        assert not any(i["type"] == "overlap" for i in issues)


class TestRectHelpers:
    def test_intersection_area_of_non_overlapping_rects_is_zero(self):
        assert _rect_intersection_area((0, 0, 10, 10), (20, 20, 30, 30)) == 0.0

    def test_intersection_area_of_identical_rects_equals_their_area(self):
        assert _rect_intersection_area((0, 0, 10, 10), (0, 0, 10, 10)) == 100.0

    def test_partial_overlap_computes_the_correct_area(self):
        assert _rect_intersection_area((0, 0, 10, 10), (5, 5, 15, 15)) == 25.0

    def test_rect_area(self):
        assert _rect_area((0, 0, 10, 4)) == 40.0
        assert _rect_area((5, 5, 5, 5)) == 0.0


class TestRunToolResult:
    @pytest.mark.asyncio
    async def test_missing_file_fails_cleanly(self, tmp_path):
        result = await ValidateRenderedOutputTool.run(pptx_path=str(tmp_path / "nope.pptx"), session_id="s1")
        assert result.success is False
        assert "not found" in result.error.lower()

    @pytest.mark.asyncio
    async def test_clean_render_reports_valid(self, tmp_path):
        """End-to-end: a real, short slide through the actual renderer and
        LibreOffice must validate cleanly."""
        from app.renderer.pptx_renderer import PPTXRenderer

        renderer = PPTXRenderer(template_path=None)
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [{
                "type": "text", "placeholderIdx": 0,
                "runs": [{"text": "A short, reasonable title", "bold": True, "fontSizePt": 36}],
            }],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        session_id = f"test-validate-{time.time_ns()}"
        result = await ValidateRenderedOutputTool.run(pptx_path=output_path, session_id=session_id)

        assert result.success is True
        assert result.data["valid"] is True
        assert result.data["issues"] == []

    @pytest.mark.asyncio
    async def test_reuses_the_preview_generator_s_cached_pdf_when_available(self, tmp_path, monkeypatch):
        """Must not re-run LibreOffice when generate_previews already
        produced a PDF for this session - confirmed by making a second
        conversion attempt raise."""
        from app.renderer.pptx_renderer import PPTXRenderer
        from app.renderer.preview_generator import PreviewGenerator

        preview_dir = str(tmp_path / "previews")
        monkeypatch.setenv("PREVIEW_OUTPUT_DIR", preview_dir)

        renderer = PPTXRenderer(template_path=None)
        slides = [{
            "slideIndex": 0, "layoutId": "1", "layoutName": "Default",
            "contents": [{"type": "text", "placeholderIdx": 0, "runs": [{"text": "Hi"}]}],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        session_id = f"test-cache-{time.time_ns()}"
        preview_gen = PreviewGenerator(output_dir=preview_dir)
        preview_gen.generate_previews(output_path, session_id)
        assert preview_gen.get_pdf_path(session_id) is not None

        def _explode(*a, **k):
            raise AssertionError("should reuse the cached PDF, not convert again")

        monkeypatch.setattr(PreviewGenerator, "convert_to_pdf", staticmethod(_explode))

        result = await ValidateRenderedOutputTool.run(pptx_path=output_path, session_id=session_id)
        assert result.success is True


class TestImageAndChartGeometry:
    """
    Charts reach a slide as embedded pictures (chart_renderer renders a
    matplotlib PNG, the renderer places it with add_picture). Geometry
    checking only read text spans via get_text("dict"), so a chart running off
    the page - or sitting on top of the body text - passed validation cleanly
    and the deck was announced as ready for download.
    """

    @staticmethod
    def _pdf_with_image(tmp_path, rect: fitz.Rect, text_ops=()) -> str:
        doc = fitz.open()
        page = doc.new_page(width=PAGE_WIDTH_PT, height=PAGE_HEIGHT_PT)
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 240, 180))
        pix.set_rect(pix.irect, (200, 60, 40))
        page.insert_image(rect, pixmap=pix)
        for text, x, y, size in text_ops:
            page.insert_text((x, y), text, fontsize=size)
        path = str(tmp_path / f"img_{time.time_ns()}.pdf")
        doc.save(path)
        doc.close()
        return path

    def test_an_image_inside_the_slide_is_not_flagged(self, tmp_path):
        pdf = self._pdf_with_image(tmp_path, fitz.Rect(60, 60, 400, 300))
        assert ValidateRenderedOutputTool._check_geometry(pdf) == []

    def test_a_chart_running_off_the_right_edge_is_flagged(self, tmp_path):
        pdf = self._pdf_with_image(
            tmp_path, fitz.Rect(PAGE_WIDTH_PT - 100, 60, PAGE_WIDTH_PT + 220, 300)
        )
        issues = ValidateRenderedOutputTool._check_geometry(pdf)
        assert any(i["type"] == "overflow" for i in issues)
        assert any("chart or image" in i["detail"] for i in issues)
        assert any("right" in i["detail"] for i in issues)

    def test_a_chart_running_off_the_bottom_edge_is_flagged(self, tmp_path):
        pdf = self._pdf_with_image(
            tmp_path, fitz.Rect(60, PAGE_HEIGHT_PT - 80, 400, PAGE_HEIGHT_PT + 160)
        )
        issues = ValidateRenderedOutputTool._check_geometry(pdf)
        assert any(i["type"] == "overflow" and "bottom" in i["detail"] for i in issues)

    def test_the_message_names_the_slide_and_suggests_a_fix(self, tmp_path):
        """The detail string is fed straight back to the model, so it has to
        say what to actually do about it."""
        pdf = self._pdf_with_image(
            tmp_path, fitz.Rect(60, PAGE_HEIGHT_PT - 80, 400, PAGE_HEIGHT_PT + 160)
        )
        issue = next(i for i in ValidateRenderedOutputTool._check_geometry(pdf)
                     if i["type"] == "overflow")
        assert issue["slideIndex"] == 0
        assert "own slide" in issue["detail"] or "smaller region" in issue["detail"]

    def test_a_chart_covering_the_body_text_is_flagged_as_overlap(self, tmp_path):
        pdf = self._pdf_with_image(
            tmp_path,
            fitz.Rect(40, 40, 400, 300),
            text_ops=[("Body text underneath the chart", 60, 150, 14)],
        )
        issues = ValidateRenderedOutputTool._check_geometry(pdf)
        assert any(i["type"] == "overlap" and "chart or image" in i["detail"] for i in issues)

    def test_an_image_beside_the_text_is_not_flagged_as_overlap(self, tmp_path):
        pdf = self._pdf_with_image(
            tmp_path,
            fitz.Rect(380, 60, 680, 300),
            text_ops=[("Text well clear of the image", 40, 150, 14)],
        )
        issues = ValidateRenderedOutputTool._check_geometry(pdf)
        assert not any(i["type"] == "overlap" for i in issues)

    def test_a_hairline_of_bleed_is_tolerated(self, tmp_path):
        """Rendered chart PNGs carry transparent padding and LibreOffice rounds
        image placement coarsely; a point or two of bleed is not a defect."""
        pdf = self._pdf_with_image(
            tmp_path, fitz.Rect(60, 60, PAGE_WIDTH_PT + 1.5, 300)
        )
        assert ValidateRenderedOutputTool._check_geometry(pdf) == []


def test_a_shape_named_as_a_background_is_not_counted_as_an_overlap(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from app.agent.tools.validate_rendered_output import background_texts

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    mark = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    mark.text_frame.text = "DRAFT"
    mark.name = "ts-background-watermark"
    other = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    other.text_frame.text = "Body"
    path = str(tmp_path / "d.pptx")
    deck.save(path)
    assert background_texts(path) == frozenset({"DRAFT"})
