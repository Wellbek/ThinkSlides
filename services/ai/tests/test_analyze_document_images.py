"""
Tests for analyze_document's embedded-image cataloging (PDF/DOCX) - the
source-material half of "reuse existing images, never generate new ones".
"""

import io

import pytest
from PIL import Image

from app.agent.tools.analyze_document import (
    AnalyzeDocumentTool,
    MAX_CATALOG_IMAGES_PER_DOCUMENT,
    MIN_CATALOG_IMAGE_DIMENSION_PX,
)


def _make_png_bytes(width: int, height: int, color=(30, 112, 74)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def pdf_with_image(tmp_path):
    import fitz

    img_path = tmp_path / "photo.png"
    img_path.write_bytes(_make_png_bytes(200, 200))

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Quarterly Revenue Growth Chart")
    page.insert_image(fitz.Rect(72, 100, 272, 300), filename=str(img_path))
    pdf_path = tmp_path / "report.pdf"
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)


@pytest.fixture
def pdf_with_tiny_image(tmp_path):
    """A tiny (icon-sized) embedded image, below the cataloging size floor."""
    import fitz

    img_path = tmp_path / "icon.png"
    img_path.write_bytes(_make_png_bytes(20, 20))

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Some page text")
    page.insert_image(fitz.Rect(72, 100, 92, 120), filename=str(img_path))
    pdf_path = tmp_path / "icons.pdf"
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)


@pytest.fixture
def docx_with_image(tmp_path):
    from docx import Document

    img_path = tmp_path / "team.png"
    img_path.write_bytes(_make_png_bytes(300, 200))

    doc = Document()
    doc.add_paragraph("Meet the leadership team behind this quarter's results.")
    doc.add_picture(str(img_path))
    doc.add_paragraph("They have led the company through record growth.")
    docx_path = tmp_path / "report.docx"
    doc.save(str(docx_path))
    return str(docx_path)


class TestPdfImageCataloging:
    @pytest.mark.asyncio
    async def test_catalogs_a_real_embedded_image_with_page_text_as_context(self, pdf_with_image):
        images = await AnalyzeDocumentTool._extract_images_pdf(pdf_with_image, "report.pdf")
        assert len(images) == 1
        entry = images[0]
        assert entry["sourceType"] == "document"
        assert entry["data"].startswith("data:image/")
        assert "Quarterly Revenue" in entry["context"]
        assert entry["sourceDocument"] == "report.pdf"

    @pytest.mark.asyncio
    async def test_tiny_decorative_images_are_filtered_out(self, pdf_with_tiny_image):
        images = await AnalyzeDocumentTool._extract_images_pdf(pdf_with_tiny_image, "icons.pdf")
        assert images == []

    @pytest.mark.asyncio
    async def test_run_includes_images_in_tool_result_for_pdf(self, pdf_with_image):
        result = await AnalyzeDocumentTool.run(document_path=pdf_with_image)
        assert result.success is True
        assert len(result.data["images"]) == 1

    @pytest.mark.asyncio
    async def test_nonexistent_pdf_returns_empty_list_not_raises(self):
        images = await AnalyzeDocumentTool._extract_images_pdf("/no/such/file.pdf", "x.pdf")
        assert images == []


class TestDocxImageCataloging:
    @pytest.mark.asyncio
    async def test_catalogs_embedded_image_with_surrounding_paragraph_text(self, docx_with_image):
        images = await AnalyzeDocumentTool._extract_images_docx(docx_with_image, "report.docx")
        assert len(images) == 1
        entry = images[0]
        assert entry["sourceType"] == "document"
        assert entry["data"].startswith("data:image/")
        assert "leadership team" in entry["context"] or "record growth" in entry["context"]

    @pytest.mark.asyncio
    async def test_run_includes_images_in_tool_result_for_docx(self, docx_with_image):
        result = await AnalyzeDocumentTool.run(document_path=docx_with_image)
        assert result.success is True
        assert len(result.data["images"]) == 1

    @pytest.mark.asyncio
    async def test_docx_with_no_images_returns_empty_list(self, tmp_path):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Just text, nothing else.")
        docx_path = tmp_path / "textonly.docx"
        doc.save(str(docx_path))

        images = await AnalyzeDocumentTool._extract_images_docx(str(docx_path), "textonly.docx")
        assert images == []


def test_size_filter_rejects_below_minimum_dimensions():
    small = _make_png_bytes(MIN_CATALOG_IMAGE_DIMENSION_PX - 10, MIN_CATALOG_IMAGE_DIMENSION_PX - 10)
    assert AnalyzeDocumentTool._passes_catalog_size_filter(small) is False


def test_size_filter_accepts_above_minimum_dimensions():
    ok = _make_png_bytes(MIN_CATALOG_IMAGE_DIMENSION_PX + 10, MIN_CATALOG_IMAGE_DIMENSION_PX + 10)
    assert AnalyzeDocumentTool._passes_catalog_size_filter(ok) is True


def test_size_filter_rejects_oversized_bytes():
    from app.agent.tools import analyze_document as mod

    huge = b"x" * (mod.MAX_CATALOG_IMAGE_BYTES + 1)
    assert AnalyzeDocumentTool._passes_catalog_size_filter(huge, width=500, height=500) is False


@pytest.mark.asyncio
async def test_catalog_is_capped_per_document(tmp_path):
    import fitz

    img_path = tmp_path / "img.png"
    img_path.write_bytes(_make_png_bytes(150, 150))

    doc = fitz.open()
    for i in range(MAX_CATALOG_IMAGES_PER_DOCUMENT + 3):
        page = doc.new_page()
        page.insert_image(fitz.Rect(10, 10, 160, 160), filename=str(img_path))
    pdf_path = tmp_path / "many.pdf"
    doc.save(str(pdf_path))
    doc.close()

    images = await AnalyzeDocumentTool._extract_images_pdf(str(pdf_path), "many.pdf")
    assert len(images) == MAX_CATALOG_IMAGES_PER_DOCUMENT
