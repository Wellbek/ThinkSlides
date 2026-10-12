"""
What the AST gate refuses, and what it lets through.

The gate is the primary control on agent-written code: nothing dangerous is
caught at runtime, because nothing dangerous can be *expressed*. So these are
not "does it usually catch this" tests - each one is a way out of the sandbox
that must be a stated refusal, and the refusal must name the reason, because
the model is expected to fix its script and try again.

`app/connections/catalog.py` refused arbitrary execution once on the grounds
that this process holds other tenants' data and decrypted API keys. Every
refusal below is that sentence, made mechanical.
"""

import pytest

from app.agent.sandbox import gate

# A script that does the job it is for: opens the deck, edits it, saves it.
GOOD = '''
from pptx import Presentation
from pptx.util import Inches, Pt

deck = Presentation("deck.pptx")
slide = deck.slides[0]
box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
box.text_frame.text = "Revenue held at 14%"
box.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
deck.save("deck.pptx")
'''


class TestTheScriptItIsFor:
    def test_a_script_that_edits_the_deck_is_allowed(self):
        assert gate.check(GOOD) is None

    def test_it_may_compute_before_it_draws(self):
        assert gate.check(
            "import statistics\n"
            "import json\n"
            "values = [1.0, 2.0, 3.5]\n"
            "mean = statistics.mean(values)\n"
            "print(json.dumps({'mean': mean}))\n"
        ) is None

    def test_it_may_draw_with_matplotlib(self):
        assert gate.check(
            "import matplotlib\n"
            "matplotlib.use('Agg')\n"
            "from matplotlib import pyplot\n"
            "figure, axes = pyplot.subplots()\n"
            "axes.bar(['a', 'b'], [1, 2])\n"
            "figure.savefig('figure.png')\n"
        ) is None

    def test_it_may_read_and_write_inside_the_work_dir(self):
        assert gate.check("open('deck.pptx', 'rb').read()") is None


class TestTheWaysOut:
    """Each of these is a route off the box or into the process. Every one
    has to be refused, and the refusal has to say which."""

    @pytest.mark.parametrize("source,expected", [
        ("import os", "os"),
        ("import sys", "sys"),
        ("import subprocess", "subprocess"),
        ("import socket", "socket"),
        ("import shutil", "shutil"),
        ("import pathlib", "pathlib"),
        ("import urllib.request", "urllib"),
        ("import requests", "requests"),
        ("import ctypes", "ctypes"),
        ("import pickle", "pickle"),
        ("import importlib", "importlib"),
        ("import builtins", "builtins"),
        ("import inspect", "inspect"),
        ("import multiprocessing", "multiprocessing"),
        ("import threading", "threading"),
        ("import asyncio", "asyncio"),
    ])
    def test_a_module_outside_the_allowlist_is_refused(self, source, expected):
        problem = gate.check(source)

        assert problem is not None
        assert expected in problem

    @pytest.mark.parametrize("source", [
        "import app",
        "import app.db",
        "from app.db import models",
        "from app.config import settings",
        "from app.connections.catalog import CATALOG",
    ])
    def test_this_services_own_code_is_refused(self, source):
        """
        The one that matters most.

        `app.db` is the tenant database and `app.config` holds the decrypted
        provider keys. A script that can import either is a script with the
        run's credentials, and no amount of process isolation below this
        would help, because it would be inside the process that has them.
        """
        problem = gate.check(source)

        assert problem is not None
        assert "app" in problem

    @pytest.mark.parametrize("source", [
        "__import__('os').system('id')",
        "eval('1+1')",
        "exec('x = 1')",
        "compile('x=1', '<s>', 'exec')",
        "globals()",
        "locals()",
        "vars()",
    ])
    def test_a_name_that_turns_data_into_code_is_refused(self, source):
        assert gate.check(source) is not None

    @pytest.mark.parametrize("source", [
        "().__class__.__bases__[0].__subclasses__()",
        "(lambda: 0).__globals__['__builtins__']",
        "type(1).__mro__",
        "print.__self__",
        "GOOD.__dict__",
    ])
    def test_reaching_through_the_object_graph_is_refused(self, source):
        """The classic escape: from any object at all to every class in the
        process, and from there to the ones holding file handles."""
        assert gate.check(source) is not None

    @pytest.mark.parametrize("source", [
        "getattr(deck, name)",
        "getattr(deck, 'a' + 'b')",
        "setattr(deck, name, 1)",
    ])
    def test_addressing_an_attribute_by_computed_name_is_refused(self, source):
        """An allowlist over names cannot see a name that is a string. A
        literal `getattr(x, 'width')` is fine - it is an attribute read, and
        the attribute rules already cover it."""
        assert gate.check(source) is not None

    def test_a_literal_getattr_is_allowed(self):
        assert gate.check("getattr(deck, 'slide_width')") is None

    @pytest.mark.parametrize("source", [
        "open('/etc/passwd')",
        "open('/app/.env')",
        "open('../../etc/shadow')",
        "open('~/.aws/credentials')",
        "open('sub/../../etc/passwd')",
        "open(target)",
        "open('deck' + name)",
    ])
    def test_opening_anything_outside_the_work_dir_is_refused(self, source):
        """
        Relative, and stated outright.

        An earlier rule held every path to a `/work/` prefix. It reads better
        and cannot be made to work: a bind mount at a fixed path needs
        CAP_SYS_ADMIN, which the container hardening drops, and two runs
        cannot share one absolute path anyway. "No absolute path, no `..`, no
        `~`, no computed name" needs no privilege and is decidable from the
        string.
        """
        problem = gate.check(source)

        assert problem is not None
        assert gate.DECK_NAME in problem

    def test_global_and_nonlocal_are_refused(self):
        assert gate.check("def f():\n    global x\n    x = 1") is not None
        assert gate.check(
            "def f():\n    y = 1\n    def g():\n        nonlocal y\n"
            "        y = 2\n    g()") is not None


class TestTheShapeOfTheScript:
    def test_an_empty_script_is_refused(self):
        assert gate.check("") is not None
        assert gate.check("   \n  ") is not None

    def test_a_script_that_does_not_parse_says_where(self):
        problem = gate.check("def f(:\n    pass")

        assert problem is not None
        assert "line" in problem

    def test_a_script_past_the_length_limit_is_refused(self):
        problem = gate.check("x = 1\n" * 6000)

        assert problem is not None
        assert str(gate.MAX_CHARS) in problem

    def test_loops_nested_past_the_limit_are_refused(self):
        """Not a security control - `setrlimit` in the runner is that. This
        catches the shape before it costs a process."""
        source = ""
        for depth in range(gate.MAX_LOOP_DEPTH + 1):
            source += "    " * depth + f"for i{depth} in range(10):\n"
        source += "    " * (gate.MAX_LOOP_DEPTH + 1) + "pass\n"

        assert gate.check(source) is not None

    def test_a_fork_bomb_cannot_even_be_written(self):
        """It needs `os`, and `os` is not a module a script may import."""
        assert gate.check(
            "import os\nwhile True:\n    os.fork()") is not None

    def test_an_infinite_loop_is_the_runners_problem_not_this_ones(self):
        """
        Stated as a test because the division matters.

        `while True: pass` is legal Python that names nothing forbidden, and
        no static check can decide whether a loop terminates. It is refused
        by the wall clock and `RLIMIT_CPU` in the runner, which is the layer
        that can actually decide it. A gate that pretended to catch this
        would be a gate people trusted for something it cannot do.
        """
        assert gate.check("while True:\n    pass") is None


class TestTheRefusalIsWorthActingOn:
    """`draw_svg` states why a refusal is a sentence: the model gets to try
    again while it still has the context to."""

    def test_a_refused_import_lists_what_is_available(self):
        problem = gate.check("import os")

        assert "pptx" in problem
        assert "matplotlib" in problem

    def test_a_refusal_says_which_line(self):
        problem = gate.check("from pptx import Presentation\nimport os\n")

        assert "line 2" in problem

    def test_the_deck_path_is_named_when_a_path_is_refused(self):
        assert "deck.pptx" in gate.check("open('/etc/passwd')")
