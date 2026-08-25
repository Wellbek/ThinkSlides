"""
The Phase 0 gate: can the representation hold a page we call the bar?

This is the test the whole rebuild rests on. If EY page 4 - three columns of
prose, a serif pull-quote, an inset photograph, a magenta panel bleeding off
the right trim, a folio hung in the margin and a quote reversed out of the
panel - cannot be *written down* in the DesignDoc, then nothing built on top
of it will ever produce that page, and the representation has to change
before another line goes on top.

It ran red before it ran green, and for the right reason: the previous
representation refused sibling overlap outright, which forbids type over an
image. That single rule is why every reconstruction attempt against the old
shape was impossible rather than merely difficult.

The reference PDFs are deliberately not committed (see ``eval/corpus.py``),
so this skips when the source is not on the machine.
"""
import os

import pytest

from app.design.geometry import check_page, resolve_page
from app.design.schema import Figure, GridPlace, TextBlock

EY_PDF = os.path.expanduser(
    "~/Downloads/ey-inaugural-ey-iif-global-insurance-risk-management-survey-march23.pdf"
)

pytestmark = pytest.mark.skipif(
    not os.path.exists(EY_PDF),
    reason="reference PDF not present; see app/eval/corpus.py for the policy",
)


@pytest.fixture(scope="module")
def reconstruction(tmp_path_factory):
    from tools.build_reconstruction import (
        PAGE,
        _page_geometry,
        build_pack,
        build_page,
    )
    from app.design.schema import DesignDoc, PackRef
    from app.measure.grid import edges_from_boxes, recover_grid

    outdir = tmp_path_factory.mktemp("reconstruction")
    doc_pdf, _, width, height, blocks, images = _page_geometry(EY_PDF, PAGE)
    boxes = [
        b["bbox"] for b in blocks if (b["bbox"][2] - b["bbox"][0]) > 0.02 * width
    ]
    grid = recover_grid(edges_from_boxes(boxes, width))
    assert grid is not None, "no grid recovered; reconstruction would be guesswork"

    pack = build_pack(width, height, grid)
    page = build_page(pack, blocks, images, width, height, str(outdir), doc_pdf)
    design = DesignDoc(pack=PackRef(id=pack.id), canvas=pack.canvas,
                       grid=pack.grid, pages=(page,))
    return design, pack, page, str(outdir)


class TestTheGridIsRecoveredNotChosen:
    def test_the_recovered_grid_matches_the_page(self, reconstruction):
        """Ground truth, read off the page: the body columns sit at 0.277 and
        0.509, a pitch of 0.232, and extending that lattice leftward lands on
        the folio at 0.045 and rightward on the bleed panel at 0.745."""
        _, pack, _, _ = reconstruction
        assert pack.grid.columns == 4
        w = pack.canvas.w_in
        assert pack.grid.margin.left / w == pytest.approx(0.045, abs=0.01)
        column = pack.grid.column_width_in(pack.canvas)
        assert column / w == pytest.approx(0.203, abs=0.01)


class TestThePageIsExpressible:
    def test_it_places_without_geometry_problems(self, reconstruction):
        design, _, page, _ = reconstruction
        assert check_page(design, page) == []

    def test_the_panel_bleeds_off_three_edges(self, reconstruction):
        """The device the previous representation could not express at all."""
        design, pack, page, _ = reconstruction
        placed = {f.id: f for f in resolve_page(design, page)}
        rail = placed["bleed-panel"]
        assert rail.rect.top == pytest.approx(0.0, abs=1e-6)
        assert rail.right == pytest.approx(pack.canvas.w_in, abs=1e-6)
        assert rail.bottom == pytest.approx(pack.canvas.h_in, abs=1e-6)
        # ...and it is still a grid column, not a free rectangle.
        assert isinstance(page.frame_by_id("bleed-panel").place, GridPlace)

    def test_the_quote_is_layered_over_the_panel(self, reconstruction):
        """Type over a full-bleed image: overlap at different z, which the
        old region validator rejected as an error."""
        design, _, page, _ = reconstruction
        placed = {f.id: f for f in resolve_page(design, page)}
        quote, panel = placed["panel-quote"], placed["bleed-panel"]
        assert quote.z > panel.z
        assert quote.rect.left >= panel.rect.left
        assert quote.right <= panel.right

    def test_every_element_of_the_page_is_present(self, reconstruction):
        _, _, page, _ = reconstruction
        ids = {f.id for f in page.all_frames()}
        for expected in ("page-panel", "bleed-panel", "bleed-figure", "folio",
                         "eyebrow", "headline", "headline-rule", "pullquote",
                         "panel-quote", "body-a", "body-b", "inset-figure"):
            assert expected in ids, f"{expected} missing from the reconstruction"

    def test_the_body_columns_carry_the_publication_s_own_copy(self, reconstruction):
        """Three paragraphs in the first column, two in the second - the page
        as it is, not a summary of it."""
        _, _, page, _ = reconstruction
        first = page.frame_by_id("body-a").content[0]
        second = page.frame_by_id("body-b").content[0]
        assert isinstance(first, TextBlock) and isinstance(second, TextBlock)
        assert len(first.paragraphs) == 3
        assert len(second.paragraphs) == 2
        assert first.type_role == "body"

    def test_the_figures_cover_their_frames_rather_than_stretching(self, reconstruction):
        """A photograph stretched into a slot whose aspect it does not share
        is a distorted photograph, and no slot's aspect is chosen to match."""
        _, _, page, _ = reconstruction
        figure = page.frame_by_id("bleed-figure").content[0]
        assert isinstance(figure, Figure)
        assert figure.fit == "cover"


class TestThePackIsMeasuredNotInvented:
    def test_body_type_matches_the_publication(self, reconstruction):
        _, pack, _, _ = reconstruction
        body = pack.role("body")
        assert body.size_pt == pytest.approx(16.0)
        assert body.line_height_pt == pytest.approx(21.0)

    def test_the_legibility_floor_permits_the_publication_s_own_body_copy(
        self, reconstruction
    ):
        """The global 16pt floor this replaces is 3.0% of a 540pt canvas.
        EY sets body copy at 1.48% of its own. The old floor forbade the
        register it was supposed to be protecting."""
        _, pack, _, _ = reconstruction
        floor = pack.policy.floor_pt("body", pack.canvas)
        assert floor <= pack.role("body").size_pt

    def test_the_pack_records_how_confident_the_grid_recovery_was(self, reconstruction):
        _, pack, _, _ = reconstruction
        assert pack.grid_confidence.value > 0
        assert pack.grid_confidence.note


class TestItRenders:
    def test_the_page_renders_to_an_openable_pptx(self, reconstruction, tmp_path):
        from pptx import Presentation

        from app.renderer.designdoc_renderer import render_doc

        design, pack, _, _ = reconstruction
        out = str(tmp_path / "ey-p4.pptx")
        render_doc(design, pack, out)
        assert os.path.exists(out)

        deck = Presentation(out)
        assert len(deck.slides) == 1
        # Text is still native and editable - fidelity did not cost that.
        texts = [
            shape.text_frame.text
            for shape in deck.slides[0].shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        assert any("Executive" in t for t in texts)
        assert any("insurance sector" in t for t in texts)
