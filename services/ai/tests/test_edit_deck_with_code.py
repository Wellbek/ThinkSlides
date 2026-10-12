"""
The tool that lets the agent write its own Python.

The three gates have their own tests - `test_sandbox_gate.py` for what a
script may say, `test_sandbox_runner.py` for what the process may reach and
what the deck may contain. These cover the tool that puts them in order: that
a refusal at any gate leaves the deck exactly as it was, that a script which
passes all three actually changes the file the download reads, and that a
figure drawn from nothing is refused before any of it runs.
"""

import asyncio
import os

import pytest

from app.agent import run_context
from app.agent.tools.edit_deck_with_code import EditDeckWithCodeTool as Tool

ADDS_A_CAPTION = '''
from pptx import Presentation
from pptx.util import Inches

deck = Presentation("deck.pptx")
box = deck.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
box.text_frame.text = "Revenue held at 14%"
deck.save("deck.pptx")
'''


@pytest.fixture
def deck_path(tmp_path):
    """A rendered deck on disk, as finalize_and_render_pptx would have left it."""
    from pptx import Presentation

    path = str(tmp_path / "deck.pptx")
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)

    token = run_context.set_rendered_path(path)
    yield path
    run_context.reset_rendered_path(token)


def _text_of(path: str) -> str:
    from pptx import Presentation

    return " ".join(
        shape.text_frame.text
        for slide in Presentation(path).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def _run(**kwargs):
    kwargs.setdefault("intent", "add the caption the page is missing")
    kwargs.setdefault("inputs", "none")
    return asyncio.run(Tool.run(**kwargs))


class TestAScriptThatPassesEveryGate:
    def test_its_edit_reaches_the_deck_on_disk(self, deck_path):
        result = _run(code=ADDS_A_CAPTION, slide_index=0)

        assert result.success, result.error
        assert "Revenue held at 14%" in _text_of(deck_path)

    def test_it_edits_the_file_the_download_reads(self, deck_path):
        """Not a copy. `finalize_and_render_pptx` writes a real file and the
        loop records its path; the download, the vision check and this tool
        all read that one path, so an edit that landed anywhere else would be
        an edit nobody ever sees."""
        _run(code=ADDS_A_CAPTION)

        assert run_context.get_rendered_path() == deck_path
        assert os.path.exists(deck_path)

    def test_the_result_says_what_it_did_in_the_slides_own_terms(self, deck_path):
        result = _run(code=ADDS_A_CAPTION,
                      intent="put the headline figure above the chart")

        assert "put the headline figure above the chart" in result.data["summary"]
        assert result.data["outcome"] == "ran"


class TestEachGateLeavesTheDeckAlone:
    """A refusal at any of the three has to be inert. A tool that edited the
    deck and *then* decided it should not have would be worse than one that
    never ran."""

    def test_a_script_the_ast_gate_refuses_never_runs(self, deck_path):
        before = _text_of(deck_path)

        result = _run(code="import os\nos.system('id')")

        assert not result.success
        assert "os" in result.error
        assert result.data["outcome"] == "refused"
        assert _text_of(deck_path) == before

    def test_a_script_that_raises_leaves_the_deck_as_it_was(self, deck_path):
        before = _text_of(deck_path)

        result = _run(code="raise ValueError('I cannot do this')")

        assert not result.success
        assert "I cannot do this" in result.error
        assert result.data["outcome"] == "failed"
        assert _text_of(deck_path) == before

    def test_a_deck_that_fails_adoption_is_discarded_with_its_reason(
            self, deck_path):
        """The script ran and produced a file. The file is not a deck."""
        before = _text_of(deck_path)

        result = _run(code="open('deck.pptx', 'w').write('not a deck at all')")

        assert not result.success
        assert "not a presentation" in result.error
        assert "discarded" in result.error
        assert _text_of(deck_path) == before

    def test_a_script_that_deletes_the_deck_is_refused(self, deck_path):
        before = _text_of(deck_path)

        result = _run(code="open('deck.pptx', 'wb').write(b'')")

        assert not result.success
        assert _text_of(deck_path) == before

    def test_a_deck_grown_past_the_runs_slide_count_is_refused(self, deck_path):
        result = _run(code=(
            "from pptx import Presentation\n"
            "deck = Presentation('deck.pptx')\n"
            "layout = deck.slide_layouts[6]\n"
            "for _ in range(200):\n"
            "    deck.slides.add_slide(layout)\n"
            "deck.save('deck.pptx')\n"
        ))

        assert not result.success
        assert "slides" in result.error

    def test_there_is_nothing_to_edit_before_anything_is_rendered(self):
        token = run_context.set_rendered_path(None)
        try:
            result = _run(code=ADDS_A_CAPTION)
        finally:
            run_context.reset_rendered_path(token)

        assert not result.success
        assert "finalize_and_render_pptx" in result.error


class TestAFigureDeclaresWhereItsNumbersCameFrom:
    """
    The risk free-form code does not remove.

    A script can produce a figure that renders perfectly and is wrong -
    invented numbers that look computed. `generate_image.py` calls that the
    single worst thing this system could emit, and a vision model grading a
    rendered page cannot tell a real figure from a plausible one. So the
    check is at the boundary and it is mechanical, exactly as it is for every
    other figure: declared, or not drawn.
    """

    DRAWS_A_CHART = (
        "from pptx import Presentation\n"
        "from pptx.chart.data import CategoryChartData\n"
        "deck = Presentation('deck.pptx')\n"
        "chart_data = CategoryChartData()\n"
        "deck.save('deck.pptx')\n"
    )

    @pytest.mark.parametrize("declared", ["", "  ", "none", "N/A", "-"])
    def test_a_chart_with_nothing_declared_is_refused(self, deck_path, declared):
        result = _run(code=self.DRAWS_A_CHART, inputs=declared)

        assert not result.success
        assert "where its numbers came from" in result.error
        assert result.data["outcome"] == "refused"

    def test_a_chart_with_its_source_named_is_allowed_through(self, deck_path):
        result = _run(
            code=self.DRAWS_A_CHART,
            inputs="the revenue table on page 4 of the uploaded 10-K",
        )

        # It gets past the grounding check; whether the script itself works is
        # a different question and a different gate.
        assert "where its numbers came from" not in (result.error or "")

    def test_a_script_that_draws_no_figure_needs_no_declaration(self, deck_path):
        """The rule is about asserted numbers, not about scripts. A caption
        is not a claim that anything was measured."""
        result = _run(code=ADDS_A_CAPTION, inputs="")

        assert result.success, result.error

    @pytest.mark.parametrize("marker", [
        "chart_data", "add_chart", "add_table", "axes.bar(", "pyplot.pie(",
    ])
    def test_every_way_of_drawing_one_is_covered(self, marker):
        from app.agent.sandbox import grounding

        assert grounding.problem(f"x = 1\n{marker}\n", "") is not None
        assert grounding.problem(f"x = 1\n{marker}\n", "the 10-K") is None

    @pytest.mark.parametrize("source", [
        "import matplotlib.pyplot as plt\nplt.bar([1], [2])\n",
        "from matplotlib.pyplot import bar\nbar([1], [2])\n",
        "slide.shapes.add_chart(1, 0, 0, 1, 1, data)\n",
        "data = CategoryChartData()\n",
    ])
    def test_a_figure_is_read_off_the_script_not_its_text(self, source):
        from app.agent.sandbox import grounding

        assert grounding.problem(source, "") is not None

    def test_a_helper_named_like_a_plot_is_not_a_figure(self):
        """probe-progress-bar (#376): a progress bar drawn by a helper called
        `bar` was refused as an ungrounded chart."""
        from app.agent.sandbox import grounding

        source = ("def bar(slide, width):\n    return width\n"
                  "bar(None, 10)\nplot = 3\n")
        assert grounding.problem(source, "none") is None
