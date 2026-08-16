"""
Representative decks, rendered for real and put through the same checks the
agent's validation step runs.

Everything else in this suite tests one layer. This renders whole decks -
the shapes a generated presentation actually takes - through PPTXRenderer
and LibreOffice, then applies the geometric, content and package passes. It
is the regression net for "the output got worse", which is the failure that
unit tests cannot see and that issue #56 was opened about.

No LLM is involved: the decks are the slide JSON the agent produces, captured
in ``app.eval.fixture_decks`` so the renderer can be changed with confidence.
They live there rather than here because the style-fidelity harness renders
the same decks to score them - see ``python -m app.eval baseline``.
"""

import os
import shutil
import tempfile

import pytest

from app.agent.tools import output_qa
from app.agent.tools.validate_rendered_output import ValidateRenderedOutputTool
from app.eval.fixture_decks import DECKS, REPRESENTATIVE_PROMPTS  # noqa: F401
from app.renderer.preview_generator import PreviewGenerator
from app.renderer.pptx_renderer import PPTXRenderer


@pytest.fixture(scope="module", params=sorted(DECKS))
def rendered(request):
    """Render a deck for real and convert it the way the preview path does."""
    slides, theme = DECKS[request.param]
    work_dir = tempfile.mkdtemp(prefix=f"ts_regress_{request.param}_")
    pptx_path = os.path.join(work_dir, "deck.pptx")
    try:
        PPTXRenderer(template_path=None, color_theme=theme).render(
            slides=slides,
            metadata={"title": f"{request.param.title()} Deck"},
            output_path=pptx_path,
        )
        pdf_path = PreviewGenerator.convert_to_pdf(pptx_path, work_dir)
        assert pdf_path, f"LibreOffice could not convert the {request.param} deck"
        yield {"name": request.param, "slides": slides,
               "pptx": pptx_path, "pdf": pdf_path}
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def _describe(issues):
    return "\n".join(f"  slide {i.get('slideIndex', '?')}: {i['detail']}" for i in issues)


class TestEveryRepresentativeDeckRendersCleanly:
    def test_the_geometry_has_no_overflow_or_overlap(self, rendered):
        issues = ValidateRenderedOutputTool._check_geometry(rendered["pdf"])

        assert issues == [], f"{rendered['name']} deck:\n{_describe(issues)}"

    def test_the_copy_passes_the_content_checks(self, rendered):
        issues = output_qa.check_content(rendered["slides"])

        assert issues == [], f"{rendered['name']} deck:\n{_describe(issues)}"

    def test_the_package_is_valid(self, rendered):
        issues = output_qa.check_package(rendered["pptx"])

        assert issues == [], f"{rendered['name']} deck:\n{_describe(issues)}"

    def test_no_slide_renders_nearly_empty(self, rendered):
        advisories = output_qa.check_density(rendered["pdf"])

        assert advisories == [], f"{rendered['name']} deck:\n{_describe(advisories)}"

    def test_every_slide_made_it_into_the_file(self, rendered):
        import fitz

        doc = fitz.open(rendered["pdf"])
        try:
            assert len(doc) == len(rendered["slides"])
        finally:
            doc.close()


class TestTheDeckReadsAsOneDesignedArtifact:
    def test_the_light_dark_rhythm_is_present(self, rendered):
        """A deck of identically white pages is what issue #56 opened on."""
        from app import colors
        from pptx import Presentation

        prs = Presentation(rendered["pptx"])
        backgrounds = {
            colors.to_hex(tuple(slide.background.fill.fore_color.rgb))
            for slide in prs.slides
        }

        assert len(backgrounds) > 1, f"{rendered['name']} deck is all one background"
        assert any(colors.is_dark(b) for b in backgrounds)

    def test_every_slide_carries_the_motif(self, rendered):
        from pptx import Presentation
        from app.renderer.pptx_renderer import CHROME_ACCENT_RULE

        prs = Presentation(rendered["pptx"])
        for index, slide in enumerate(prs.slides):
            assert any((s.name or "") == CHROME_ACCENT_RULE for s in slide.shapes), (
                f"{rendered['name']} slide {index + 1} has no accent rule"
            )

    def test_body_text_is_never_left_at_the_inherited_colour(self, rendered):
        """The defect that put black subtitles on a brand-dark cover."""
        from pptx import Presentation
        from app.renderer.pptx_renderer import is_chrome_shape

        prs = Presentation(rendered["pptx"])
        for index, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if not shape.has_text_frame or is_chrome_shape(shape):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        assert run.font.color.rgb is not None, (
                            f"{rendered['name']} slide {index + 1}: "
                            f"{run.text[:40]!r} has no explicit colour"
                        )
