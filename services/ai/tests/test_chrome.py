"""
Deck chrome: the eyebrow, the brand mark, the page number, the footer rule.

The eyebrow is the piece that matters and the one that did not exist. A page
number tells a reader how far through they are; a section eyebrow tells them
*where they are*, which is what makes 184 pages read as one document. It only
works if every slide in a part names that part identically, so the section
comes from the plan rather than being re-invented per slide.

The regression that matters here is the quiet one: a deck that declares no
chrome must render the footer band it always did.
"""

import os
import shutil
import tempfile

import pytest
from pptx import Presentation

from app import colors
from app.design_system import get_default_design_system
from app.renderer import chrome as chrome_module
from app.renderer.chrome import ChromeSpec, compile_chrome
from app.renderer.deck_theme import (
    DeckTheme, ROLE_CLOSING, ROLE_CONTENT, ROLE_COVER, ROLE_SECTION,
)
from app.renderer.layout import PositionConfig
from app.renderer.pptx_renderer import PPTXRenderer, is_chrome_shape
from app.renderer.primitives import shape_primitive_id
from app.renderer.typography import TypeScale
from app.renderer.xml_utils import emu_to_inches

CANVAS = PositionConfig(left=0.0, top=0.0, width=13.333, height=7.5)


def _compile(declaration=None, role=ROLE_CONTENT, section="Applied AI",
             deck_title="Technology Trends", number=12, total=40):
    design_system = get_default_design_system()
    if declaration is not None:
        design_system = {**design_system, "chrome": declaration}
    theme = DeckTheme.from_design_system(design_system)
    return compile_chrome(
        spec=ChromeSpec.from_design_system(design_system),
        style=theme.style_for(role),
        type_scale=TypeScale.from_design_system(design_system),
        canvas=CANVAS, role=role, slide_number=number, total_slides=total,
        deck_title=deck_title, section_name=section,
    )


def _by_id(marks):
    return {mark.id: mark for mark in marks}


class TestTheDefaultIsTheFooterThatWasAlwaysThere:
    def test_a_deck_declaring_no_chrome_gets_a_label_and_a_page_number(self):
        marks = _by_id(_compile())

        assert marks["chrome-footer"].text == "Technology Trends"
        assert marks["chrome-page-number"].text == "12"

    def test_the_label_sits_left_and_the_number_right(self):
        marks = _by_id(_compile())

        assert marks["chrome-footer"].align == "left"
        assert marks["chrome-page-number"].align == "right"

    def test_no_eyebrow_or_brand_mark_unless_asked_for(self):
        marks = _by_id(_compile())

        assert "chrome-eyebrow" not in marks
        assert "chrome-brand" not in marks

    def test_the_cover_carries_none_of_it(self):
        assert _compile(role=ROLE_COVER) == []

    def test_every_mark_is_classified_as_chrome(self):
        assert all(mark.chrome for mark in _compile())


class TestTheEyebrow:
    DECLARATION = chrome_module.reference_chrome_declaration()

    def test_it_names_the_section(self):
        marks = _by_id(_compile(self.DECLARATION))

        assert marks["chrome-eyebrow"].text == "Applied AI"

    def test_it_sits_top_right(self):
        eyebrow = _by_id(_compile(self.DECLARATION))["chrome-eyebrow"]

        assert eyebrow.align == "right"
        assert eyebrow.box.top < CANVAS.height / 2

    def test_a_slide_with_no_section_falls_back_rather_than_printing_nothing(self):
        marks = _by_id(_compile(self.DECLARATION, section=""))

        assert marks["chrome-eyebrow"].text == "Technology Trends"

    def test_it_can_be_told_to_show_the_deck_title_instead(self):
        declaration = {**self.DECLARATION,
                       "eyebrow": {**self.DECLARATION["eyebrow"],
                                   "source": "deckTitle"}}

        marks = _by_id(_compile(declaration))

        assert marks["chrome-eyebrow"].text == "Technology Trends"

    def test_it_can_be_switched_off_entirely(self):
        declaration = {**self.DECLARATION,
                       "eyebrow": {**self.DECLARATION["eyebrow"], "source": "none"}}

        assert "chrome-eyebrow" not in _by_id(_compile(declaration))

    def test_it_is_limited_to_the_roles_it_was_declared_for(self):
        """A section break does not need an eyebrow naming the section."""
        assert "chrome-eyebrow" not in _by_id(
            _compile(self.DECLARATION, role=ROLE_SECTION)
        )


class TestEachPieceCanBeMovedOrDropped:
    def test_a_position_moves_the_mark(self):
        left = _by_id(_compile({"pageNumber": {"position": "bottomLeft"}}))
        right = _by_id(_compile({"pageNumber": {"position": "bottomRight"}}))

        assert left["chrome-page-number"].box.left < \
            right["chrome-page-number"].box.left

    def test_a_disabled_element_is_not_drawn(self):
        marks = _by_id(_compile({"pageNumber": {"enabled": False}}))

        assert "chrome-page-number" not in marks

    def test_a_declared_size_is_applied(self):
        marks = _by_id(_compile({"pageNumber": {"sizePt": 14}}))

        assert marks["chrome-page-number"].size_pt == 14

    def test_a_page_number_format_can_show_the_total(self):
        marks = _by_id(_compile({"pageNumberFormat": "{n} / {total}"}))

        assert marks["chrome-page-number"].text == "12 / 40"

    def test_a_broken_format_falls_back_to_the_bare_number(self):
        marks = _by_id(_compile({"pageNumberFormat": "{nope}"}))

        assert marks["chrome-page-number"].text == "12"

    def test_a_footer_rule_is_drawn_when_declared(self):
        marks = _by_id(_compile({"footerRule": {"enabled": True}}))

        rule = marks["chrome-footer-rule"]
        assert rule.kind == "rect"
        assert rule.box.width > CANVAS.width * 0.8

    def test_it_sits_above_the_footer_band_rather_than_through_it(self):
        """It used to land inside the footer's own line box, so every deck
        that declared a rule struck out its own footer text."""
        marks = _by_id(_compile({"footerRule": {"enabled": True}}))

        rule = marks["chrome-footer-rule"]
        footer = marks["chrome-footer"]
        assert rule.box.top + rule.box.height <= footer.box.top

    def test_a_footer_rule_turned_off_is_not_drawn(self):
        assert "chrome-footer-rule" not in _by_id(
            _compile({"footerRule": {"enabled": False}})
        )

    def test_suppressed_roles_are_configurable(self):
        declaration = {"suppressedRoles": [ROLE_COVER, ROLE_CLOSING]}

        assert _compile(declaration, role=ROLE_CLOSING) == []
        assert _compile(declaration, role=ROLE_CONTENT) != []


class TestTheBrandMark:
    def test_it_can_be_words(self):
        marks = _by_id(_compile({
            "brandMark": {"enabled": True, "kind": "text", "value": "Acme & Co"},
        }))

        assert marks["chrome-brand"].kind == "text"
        assert marks["chrome-brand"].text == "Acme & Co"

    def test_it_can_be_a_picture(self):
        marks = _by_id(_compile({
            "brandMark": {"enabled": True, "kind": "image",
                          "value": "https://example.test/logo.png"},
        }))

        brand = marks["chrome-brand"]
        assert brand.kind == "image"
        assert brand.image["data"] == "https://example.test/logo.png"

    def test_it_does_not_sit_on_top_of_the_page_number(self):
        """
        A logo and a number in the same corner collide at exactly the point
        in a deck where the number reaches three digits.
        """
        marks = _by_id(_compile({
            "brandMark": {"enabled": True, "kind": "text", "value": "Acme"},
            "pageNumber": {"enabled": True, "position": "bottomRight"},
        }))

        brand = marks["chrome-brand"]
        number = marks["chrome-page-number"]
        assert brand.box.left + brand.box.width <= number.box.left + \
            number.box.width - 0.4

    def test_an_empty_value_draws_nothing(self):
        marks = _by_id(_compile({"brandMark": {"enabled": True, "value": ""}}))

        assert "chrome-brand" not in marks


class TestChromeIsLegibleOnWhateverItLandsOn:
    @pytest.mark.parametrize("role", [ROLE_CONTENT, ROLE_SECTION, ROLE_CLOSING])
    def test_the_colour_comes_from_the_slides_own_style(self, role):
        design_system = get_default_design_system()
        theme = DeckTheme.from_design_system(design_system)
        style = theme.style_for(role)

        marks = _compile(chrome_module.reference_chrome_declaration(), role=role)

        for mark in marks:
            if mark.kind != "text":
                continue
            assert colors.contrast_ratio(mark.color, style.background) > 1.5


class TestChromeReachesTheRenderedFile:
    @pytest.fixture(scope="class")
    @staticmethod
    def work_dir():
        directory = tempfile.mkdtemp(prefix="ts_chrome_")
        try:
            yield directory
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @staticmethod
    def _render(work_dir, name, slides, design_system=None):
        path = os.path.join(work_dir, f"{name}.pptx")
        PPTXRenderer(template_path=None, color_theme="professional",
                     design_system=design_system).render(
            slides=slides, metadata={"title": "Technology Trends"},
            output_path=path,
        )
        return Presentation(path)

    @staticmethod
    def _slides():
        return [
            {"slideIndex": 0, "layoutId": "title",
             "contents": [{"type": "text", "runs": [{"text": "Cover"}]}]},
            {"slideIndex": 1, "layoutId": "content", "section": "Applied AI",
             "contents": [{"type": "text", "runs": [{"text": "Adoption"}]},
                          {"type": "text",
                           "runs": [{"text": "A point", "bullet": True}]}]},
            {"slideIndex": 2, "layoutId": "content", "section": "Applied AI",
             "contents": [{"type": "text", "runs": [{"text": "Investment"}]},
                          {"type": "text",
                           "runs": [{"text": "Another point", "bullet": True}]}]},
        ]

    def test_the_eyebrow_names_the_section_on_every_slide_in_it(self, work_dir):
        design_system = {**get_default_design_system(),
                         "chrome": chrome_module.reference_chrome_declaration()}
        deck = self._render(work_dir, "eyebrow", self._slides(), design_system)

        for index in (1, 2):
            eyebrow = next(
                s for s in deck.slides[index].shapes
                if shape_primitive_id(s) == "chrome-eyebrow"
            )
            assert eyebrow.text_frame.text == "Applied AI"

    def test_chrome_is_classified_as_chrome_in_the_file(self, work_dir):
        design_system = {**get_default_design_system(),
                         "chrome": chrome_module.reference_chrome_declaration()}
        deck = self._render(work_dir, "classified", self._slides(), design_system)

        marks = [s for s in deck.slides[1].shapes
                 if (shape_primitive_id(s) or "").startswith("chrome-")]
        assert marks
        assert all(is_chrome_shape(s) for s in marks)

    def test_a_deck_declaring_nothing_still_carries_its_footer(self, work_dir):
        deck = self._render(work_dir, "default", self._slides())

        text = " ".join(
            s.text_frame.text for s in deck.slides[1].shapes
            if (shape_primitive_id(s) or "").startswith("chrome-")
            and s.has_text_frame
        )
        assert "Technology Trends" in text
        assert "2" in text

    def test_chrome_sits_inside_the_canvas(self, work_dir):
        design_system = {**get_default_design_system(),
                         "chrome": chrome_module.reference_chrome_declaration()}
        deck = self._render(work_dir, "inside", self._slides(), design_system)

        for shape in deck.slides[1].shapes:
            if not (shape_primitive_id(shape) or "").startswith("chrome-"):
                continue
            assert emu_to_inches(shape.left) >= -0.01
            assert emu_to_inches(shape.left) + emu_to_inches(shape.width) \
                <= CANVAS.width + 0.01
            assert emu_to_inches(shape.top) + emu_to_inches(shape.height) \
                <= CANVAS.height + 0.01


class TestTheSectionComesFromThePlan:
    def test_the_planner_asks_for_one(self):
        from app.agent.tools import plan_presentation_structure as planner
        import inspect

        source = inspect.getsource(planner)
        assert '"section"' in source
        assert "REPEATED verbatim" in source

    def test_the_content_designer_carries_it_onto_the_slide(self):
        from app.agent.tools import design_slide_content as designer
        import inspect

        source = inspect.getsource(designer)
        assert 'outline.get("section")' in source
