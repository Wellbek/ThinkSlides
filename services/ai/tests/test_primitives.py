"""
The primitive layer: the open floor under the block vocabulary.

The point of this layer is that it is open *without* being outside the guard
rails. So the tests split in two: that a mark can express anything a slide
needs, and that a mark is checked exactly as closely as a block - contained in
its frame, legible against its ground, measured before it is drawn.
"""

import base64
import io
import os
import shutil
import tempfile

import pytest
from pptx import Presentation

from app import colors
from app.renderer import primitives as prim
from app.renderer import text_fit
from app.renderer.layout import PositionConfig
from app.renderer.pptx_renderer import PPTXRenderer, is_chrome_shape
from app.renderer.typography import TypeScale
from app.renderer.xml_utils import emu_to_inches

CANVAS = PositionConfig(left=0.0, top=0.0, width=13.333, height=7.5)
NAVY = "#051C2C"


def _context(ground=NAVY):
    return prim.PaintContext(
        ground=ground, default_color="#FFFFFF", type_scale=TypeScale(),
        default_size_pt=18.0, default_font="Arial",
    )


def _png_bytes():
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 30, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


class TestResolvingAuthoredGeometry:
    def test_fractions_resolve_against_the_canvas(self):
        [mark] = prim.resolve(
            [{"kind": "rect", "x": 0.5, "y": 0.25, "w": 0.25, "h": 0.5,
              "fill": "#FFFFFF"}],
            frame_for=None, canvas=CANVAS,
        )

        assert mark.box.left == pytest.approx(CANVAS.width * 0.5)
        assert mark.box.top == pytest.approx(CANVAS.height * 0.25)
        assert mark.box.width == pytest.approx(CANVAS.width * 0.25)

    def test_fractions_resolve_against_a_named_region(self):
        frame = PositionConfig(left=6.0, top=1.0, width=4.0, height=3.0)

        [mark] = prim.resolve(
            [{"kind": "rect", "regionId": "right", "x": 0.5, "y": 0.0,
              "w": 0.5, "h": 1.0, "fill": "#FFFFFF"}],
            frame_for=lambda _: frame, canvas=CANVAS,
        )

        assert mark.box.left == pytest.approx(8.0)
        assert mark.box.width == pytest.approx(2.0)
        assert mark.region_id == "right"

    def test_an_unknown_region_falls_back_to_the_canvas(self):
        [mark] = prim.resolve(
            [{"kind": "rect", "regionId": "ghost", "x": 0, "y": 0, "w": 1, "h": 1,
              "fill": "#FFFFFF"}],
            frame_for=lambda _: None, canvas=CANVAS,
        )

        assert mark.box.width == pytest.approx(CANVAS.width)

    def test_a_line_keeps_both_of_its_ends(self):
        [mark] = prim.resolve(
            [{"kind": "line", "x1": 0.1, "y1": 0.5, "x2": 0.9, "y2": 0.5}],
            frame_for=None, canvas=CANVAS,
        )

        assert mark.kind == "line"
        assert mark.end == pytest.approx((CANVAS.width * 0.9, CANVAS.height * 0.5))

    def test_ids_are_assigned_when_absent_and_kept_when_given(self):
        marks = prim.resolve(
            [{"kind": "rect", "fill": "#FFF"},
             {"kind": "rect", "id": "divider", "fill": "#FFF"}],
            frame_for=None, canvas=CANVAS, id_prefix="s1-mark",
        )

        assert marks[0].id == "s1-mark-0"
        assert marks[1].id == "divider"

    def test_duplicate_ids_are_made_unique(self):
        marks = prim.resolve(
            [{"kind": "rect", "id": "same", "fill": "#FFF"},
             {"kind": "rect", "id": "same", "fill": "#FFF"}],
            frame_for=None, canvas=CANVAS,
        )

        assert len({m.id for m in marks}) == 2

    def test_an_unknown_kind_is_dropped_rather_than_guessed(self):
        assert prim.resolve([{"kind": "hexagon"}], frame_for=None,
                            canvas=CANVAS) == []


class TestValidationChecksTheOpenLayerAsCloselyAsTheClosedOne:
    def test_a_mark_escaping_its_frame_is_rejected(self):
        errors, _ = prim.validate_primitives(
            [{"kind": "rect", "x": 0.8, "y": 0, "w": 0.5, "h": 1, "fill": "#FFF"}]
        )

        assert any("escapes its frame" in e for e in errors)

    def test_a_mark_bound_to_a_missing_region_is_rejected(self):
        errors, _ = prim.validate_primitives(
            [{"kind": "rect", "regionId": "ghost", "fill": "#FFF"}],
            region_ids=["real"],
        )

        assert any("ghost" in e for e in errors)

    def test_two_labels_drawn_on_top_of_each_other_are_rejected(self):
        errors, _ = prim.validate_primitives([
            {"kind": "text", "x": 0.1, "y": 0.1, "w": 0.4, "h": 0.2, "text": "One"},
            {"kind": "text", "x": 0.15, "y": 0.12, "w": 0.4, "h": 0.2, "text": "Two"},
        ])

        assert any("overlap" in e for e in errors)

    def test_labels_side_by_side_are_not_an_overlap(self):
        errors, _ = prim.validate_primitives([
            {"kind": "text", "x": 0.0, "y": 0.1, "w": 0.45, "h": 0.2, "text": "One"},
            {"kind": "text", "x": 0.5, "y": 0.1, "w": 0.45, "h": 0.2, "text": "Two"},
        ])

        assert errors == []

    def test_an_empty_text_mark_is_rejected(self):
        errors, _ = prim.validate_primitives(
            [{"kind": "text", "x": 0, "y": 0, "w": 1, "h": 0.2, "text": "  "}]
        )

        assert any("no text" in e for e in errors)

    def test_an_image_mark_without_a_source_is_rejected(self):
        errors, _ = prim.validate_primitives([{"kind": "image", "image": {}}])

        assert any("data source" in e for e in errors)

    def test_a_zero_length_line_is_rejected(self):
        errors, _ = prim.validate_primitives(
            [{"kind": "line", "x1": 0.5, "y1": 0.5, "x2": 0.5, "y2": 0.5}]
        )

        assert any("no length" in e for e in errors)

    def test_an_invisible_shape_warns_rather_than_failing(self):
        errors, warnings = prim.validate_primitives(
            [{"kind": "rect", "x": 0, "y": 0, "w": 0.5, "h": 0.5}]
        )

        assert errors == []
        assert any("renders as nothing" in w for w in warnings)

    def test_duplicate_ids_are_rejected(self):
        errors, _ = prim.validate_primitives([
            {"kind": "rect", "id": "a", "fill": "#FFF"},
            {"kind": "rect", "id": "a", "fill": "#FFF"},
        ])

        assert any("duplicate" in e for e in errors)

    def test_no_primitives_is_not_an_error(self):
        assert prim.validate_primitives(None) == ([], [])


class TestPainting:
    @pytest.fixture(scope="class")
    @staticmethod
    def slide():
        from app.renderer.default_template import get_default_template_path

        presentation = Presentation(get_default_template_path())
        yield presentation.slides.add_slide(presentation.slide_layouts[0])

    def test_each_kind_reaches_the_slide(self, slide):
        marks = prim.resolve([
            {"kind": "text", "id": "t", "x": 0.05, "y": 0.05, "w": 0.3, "h": 0.1,
             "text": "Label"},
            {"kind": "rect", "id": "r", "x": 0.4, "y": 0.05, "w": 0.2, "h": 0.1,
             "fill": "#FFFFFF"},
            {"kind": "ellipse", "id": "e", "x": 0.65, "y": 0.05, "w": 0.1, "h": 0.1,
             "fill": "#FFFFFF"},
            {"kind": "line", "id": "l", "x1": 0.05, "y1": 0.5, "x2": 0.9, "y2": 0.5,
             "stroke": "#FFFFFF", "strokeWeightPt": 1.0},
            {"kind": "image", "id": "i", "x": 0.05, "y": 0.6, "w": 0.1, "h": 0.1,
             "image": {"source": "base64",
                       "data": base64.b64encode(_png_bytes()).decode()}},
        ], frame_for=None, canvas=CANVAS)

        shapes = prim.paint_all(slide, marks, _context())

        assert len(shapes) == 5

    def test_a_shape_carries_its_primitive_id(self, slide):
        [mark] = prim.resolve(
            [{"kind": "rect", "id": "named-mark", "x": 0.05, "y": 0.8,
              "w": 0.1, "h": 0.05, "fill": "#FFFFFF"}],
            frame_for=None, canvas=CANVAS,
        )

        shape = prim.paint(slide, mark, _context())

        assert prim.shape_primitive_id(shape) == "named-mark"
        assert prim.is_content_shape(shape)
        assert not is_chrome_shape(shape)

    def test_a_shape_is_not_mistaken_for_chrome(self, slide):
        """Content named by this layer must keep being measured as content."""
        assert not is_chrome_shape(
            type("S", (), {"name": prim.content_shape_name("x")})()
        )


class TestTextIsAlwaysLegibleAgainstItsGround:
    @pytest.fixture(scope="class")
    @staticmethod
    def work_dir():
        directory = tempfile.mkdtemp(prefix="ts_prim_")
        try:
            yield directory
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    def test_a_colour_that_fails_contrast_is_corrected(self, work_dir):
        """A mark may ask for any colour; it may not ask to be invisible."""
        slides = [{
            "slideIndex": 0, "layoutId": "content",
            "regions": [{"id": "dark", "kind": "panel", "x": 0, "y": 0,
                         "w": 1.0, "h": 1.0, "ground": NAVY}],
            "primitives": [
                {"kind": "text", "id": "bad", "regionId": "dark",
                 "x": 0.1, "y": 0.4, "w": 0.6, "h": 0.1,
                 "text": "Nearly invisible", "color": "#0A2030"},
            ],
            "contents": [],
        }]
        path = os.path.join(work_dir, "contrast.pptx")
        PPTXRenderer(template_path=None, color_theme="professional").render(
            slides=slides, metadata={"title": "Contrast"}, output_path=path)

        for shape in Presentation(path).slides[0].shapes:
            if prim.shape_primitive_id(shape) != "bad":
                continue
            run = shape.text_frame.paragraphs[0].runs[0]
            rendered = colors.to_hex(tuple(run.font.color.rgb))
            assert colors.contrast_ratio(rendered, NAVY) >= colors.WCAG_AA_NORMAL
            break
        else:
            pytest.fail("the mark did not reach the slide")

    def test_a_mark_lands_where_its_fractions_put_it(self, work_dir):
        slides = [{
            "slideIndex": 0, "layoutId": "content",
            "primitives": [
                {"kind": "rect", "id": "band", "x": 0.25, "y": 0.5,
                 "w": 0.5, "h": 0.02, "fill": "#FF0000"},
            ],
            "contents": [],
        }]
        path = os.path.join(work_dir, "placed.pptx")
        PPTXRenderer(template_path=None, color_theme="professional").render(
            slides=slides, metadata={"title": "Placed"}, output_path=path)

        band = next(s for s in Presentation(path).slides[0].shapes
                    if prim.shape_primitive_id(s) == "band")
        assert emu_to_inches(band.left) == pytest.approx(CANVAS.width * 0.25, abs=0.02)
        assert emu_to_inches(band.width) == pytest.approx(CANVAS.width * 0.5, abs=0.02)

    def test_marks_are_drawn_behind_the_slides_content(self, work_dir):
        slides = [{
            "slideIndex": 0, "layoutId": "content",
            "primitives": [{"kind": "rect", "id": "panel", "x": 0.0, "y": 0.0,
                            "w": 1.0, "h": 1.0, "fill": "#EEEEEE"}],
            "contents": [{"type": "text", "runs": [{"text": "Headline"}]}],
        }]
        path = os.path.join(work_dir, "behind.pptx")
        PPTXRenderer(template_path=None, color_theme="professional").render(
            slides=slides, metadata={"title": "Behind"}, output_path=path)

        shapes = list(Presentation(path).slides[0].shapes)
        panel = next(i for i, s in enumerate(shapes)
                     if prim.shape_primitive_id(s) == "panel")
        headline = next(i for i, s in enumerate(shapes)
                        if s.has_text_frame and "Headline" in s.text_frame.text)
        assert panel < headline


class TestMeasurement:
    def test_bounds_cover_every_mark(self):
        marks = prim.resolve([
            {"kind": "rect", "x": 0.1, "y": 0.1, "w": 0.2, "h": 0.2, "fill": "#FFF"},
            {"kind": "rect", "x": 0.5, "y": 0.6, "w": 0.2, "h": 0.2, "fill": "#FFF"},
        ], frame_for=None, canvas=CANVAS)

        extent = prim.bounds(marks)

        assert extent.left == pytest.approx(CANVAS.width * 0.1)
        assert extent.top + extent.height == pytest.approx(CANVAS.height * 0.8)

    def test_bounds_of_nothing_is_none(self):
        assert prim.bounds([]) is None

    def test_translating_keeps_ids_and_moves_geometry(self):
        marks = prim.resolve(
            [{"kind": "rect", "id": "a", "x": 0, "y": 0, "w": 0.5, "h": 0.5,
              "fill": "#FFF"}],
            frame_for=None, canvas=CANVAS,
        )

        moved = prim.translated(marks, 1.0, 2.0)

        assert moved[0].id == "a"
        assert moved[0].box.left == pytest.approx(marks[0].box.left + 1.0)
        assert moved[0].box.top == pytest.approx(marks[0].box.top + 2.0)

    def test_measure_text_is_conservative_rather_than_optimistic(self):
        """
        Under-measuring costs an overlap; over-measuring costs white space.
        """
        width = 2.0
        generous = prim.measure_text("Segment performance", 18.0, width)
        bold = prim.measure_text("Segment performance", 18.0, width, bold=True)

        assert bold >= generous


class TestTheWrapEstimatorHandlesAWordWiderThanItsLine:
    """
    "$0.8T-$1.3T" in a narrow rail measured as one line, rendered as two, and
    the second landed on the caption underneath it.
    """

    def test_a_long_unbroken_word_takes_more_than_one_line(self):
        lines = text_fit.estimate_wrapped_lines("$0.8T-$1.3T", 32.0, 100.0)

        assert lines >= 2

    def test_a_word_that_fits_still_takes_one_line(self):
        assert text_fit.estimate_wrapped_lines("Revenue", 12.0, 400.0) == 1

    def test_the_estimate_scales_with_how_far_over_the_word_is(self):
        narrow = text_fit.estimate_wrapped_lines("A" * 60, 20.0, 100.0)
        wide = text_fit.estimate_wrapped_lines("A" * 60, 20.0, 400.0)

        assert narrow > wide >= 1


class TestAutofitHonoursTheFramesRealInsets:
    def test_zero_insets_fit_more_than_default_insets(self):
        blocks = [text_fit.TextBlock(text="Telecommunications", font_size_pt=20)]

        with_default = text_fit.compute_autofit_scale(blocks, 180.0, 26.0)[0]
        with_none = text_fit.compute_autofit_scale(
            blocks, 180.0, 26.0, inset_h_pt=0.0, inset_v_pt=0.0
        )[0]

        assert with_none >= with_default
