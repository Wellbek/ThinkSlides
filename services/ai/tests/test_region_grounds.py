"""
Split and banded grounds: a slide can be two canvases, and both halves have
to be legible.

The failure this prevents is specific and it is the worst kind - a deck that
opens, looks designed, and has half its text invisible. A slide used to have
exactly one background and text was coloured from it; the moment a slide has
two grounds, colouring from the slide gets one half right and one half wrong.
"""

import os
import shutil
import tempfile
import zipfile

import pytest
from pptx import Presentation

from app import colors
from app.eval.fixture_decks import DECKS
from app.renderer import deck_theme
from app.renderer.chart_renderer import _adapt_to_ground
from app.renderer.pptx_renderer import (
    CHROME_REGION_GROUND, PPTXRenderer, is_chrome_shape,
)
from app.renderer.regions import RegionTree
from app.renderer.table_style import resolve_chart_style
from app.renderer.xml_utils import emu_to_inches

CANVAS_W, CANVAS_H = 13.333, 7.5
NAVY = "#051C2C"
PAPER = "#FFFFFF"

SPLIT = [
    {"id": "left", "kind": "panel", "x": 0.0, "y": 0.0, "w": 0.5, "h": 1.0,
     "ground": PAPER, "children": [
         {"id": "left-title", "kind": "panel", "x": 0.08, "y": 0.06, "w": 0.84, "h": 0.18},
         {"id": "left-body", "kind": "panel", "x": 0.08, "y": 0.30, "w": 0.84, "h": 0.50},
     ]},
    {"id": "right", "kind": "panel", "x": 0.5, "y": 0.0, "w": 0.5, "h": 1.0,
     "ground": NAVY, "children": [
         {"id": "right-body", "kind": "panel", "x": 0.08, "y": 0.20, "w": 0.84, "h": 0.60},
     ]},
]


def _render(slides, work_dir, name, theme="professional"):
    path = os.path.join(work_dir, f"{name}.pptx")
    PPTXRenderer(template_path=None, color_theme=theme).render(
        slides=slides, metadata={"title": "Grounds"}, output_path=path,
    )
    return path


def _split_slide(**overrides):
    slide = {
        "slideIndex": 0, "layoutId": "content", "regions": SPLIT,
        "contents": [
            {"type": "text", "regionId": "left-title",
             "runs": [{"text": "On white", "bullet": False}]},
            {"type": "text", "regionId": "left-body",
             "runs": [{"text": "Body copy on the light half", "bullet": True}]},
            {"type": "text", "regionId": "right-body",
             "runs": [{"text": "Body copy on the dark half", "bullet": True}]},
        ],
    }
    slide.update(overrides)
    return slide


def _runs_with_ground(path, tree):
    """Every rendered run, paired with the ground of the region it sits in."""
    out = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if is_chrome_shape(shape) or not shape.has_text_frame:
                continue
            centre_x = emu_to_inches(shape.left) + emu_to_inches(shape.width) / 2
            centre_y = emu_to_inches(shape.top) + emu_to_inches(shape.height) / 2
            ground = PAPER
            for resolved in tree.grounded():
                rect = resolved.rect
                if (rect.left <= centre_x <= rect.left + rect.width
                        and rect.top <= centre_y <= rect.top + rect.height):
                    ground = resolved.region.ground
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip() or run.font.color.rgb is None:
                        continue
                    out.append((run.text, colors.to_hex(tuple(run.font.color.rgb)), ground))
    return out


class TestARegionCanDeclareItsOwnGround:
    @pytest.fixture(scope="class")
    @staticmethod
    def work_dir():
        directory = tempfile.mkdtemp(prefix="ts_grounds_")
        try:
            yield directory
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    def test_a_declared_ground_renders_as_a_real_fill(self, work_dir):
        path = _render([_split_slide()], work_dir, "split")

        fills = {
            colors.to_hex(tuple(shape.fill.fore_color.rgb))
            for shape in Presentation(path).slides[0].shapes
            if (shape.name or "") == CHROME_REGION_GROUND
        }
        assert NAVY in fills
        assert PAPER in fills

    def test_a_ground_covers_the_region_it_was_declared_on(self, work_dir):
        path = _render([_split_slide()], work_dir, "split")
        tree = RegionTree.from_slide({"regions": SPLIT}, CANVAS_W, CANVAS_H)
        expected = tree.rect("right")

        navy = next(
            shape for shape in Presentation(path).slides[0].shapes
            if (shape.name or "") == CHROME_REGION_GROUND
            and colors.to_hex(tuple(shape.fill.fore_color.rgb)) == NAVY
        )
        assert emu_to_inches(navy.left) == pytest.approx(expected.left, abs=0.01)
        assert emu_to_inches(navy.width) == pytest.approx(expected.width, abs=0.01)

    def test_a_ground_is_behind_the_content_it_holds(self, work_dir):
        path = _render([_split_slide()], work_dir, "split")

        shapes = list(Presentation(path).slides[0].shapes)
        last_ground = max(
            i for i, s in enumerate(shapes) if (s.name or "") == CHROME_REGION_GROUND
        )
        first_content = min(
            i for i, s in enumerate(shapes)
            if s.has_text_frame and s.text_frame.text.strip() and not is_chrome_shape(s)
        )
        assert last_ground < first_content

    def test_a_ground_is_classified_as_chrome(self, work_dir):
        path = _render([_split_slide()], work_dir, "split")

        grounds = [
            s for s in Presentation(path).slides[0].shapes
            if (s.name or "") == CHROME_REGION_GROUND
        ]
        assert grounds
        assert all(is_chrome_shape(s) for s in grounds)

    def test_grounds_paint_outermost_first(self, work_dir):
        nested = [{"id": "panel", "kind": "panel", "x": 0, "y": 0, "w": 1, "h": 1,
                   "ground": PAPER, "children": [
                       {"id": "rail", "kind": "rail", "x": 0, "y": 0,
                        "w": 0.3, "h": 1.0, "ground": NAVY}]}]
        path = _render([{"slideIndex": 0, "layoutId": "content", "regions": nested,
                         "contents": [{"type": "text", "regionId": "rail",
                                       "runs": [{"text": "In the rail"}]}]}],
                       work_dir, "nested")

        order = [
            colors.to_hex(tuple(s.fill.fore_color.rgb))
            for s in Presentation(path).slides[0].shapes
            if (s.name or "") == CHROME_REGION_GROUND
        ]
        assert order == [PAPER, NAVY]


class TestTextIsColouredForItsOwnGround:
    @pytest.fixture(scope="class")
    @staticmethod
    def rendered():
        directory = tempfile.mkdtemp(prefix="ts_grounds_text_")
        try:
            path = _render([_split_slide()], directory, "split")
            tree = RegionTree.from_slide({"regions": SPLIT}, CANVAS_W, CANVAS_H)
            yield path, tree
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    def test_the_two_halves_do_not_get_the_same_text_colour(self, rendered):
        path, tree = rendered
        by_ground = {}
        for text, color, ground in _runs_with_ground(path, tree):
            by_ground.setdefault(ground, set()).add(color)

        assert by_ground[NAVY] != by_ground[PAPER]

    def test_every_run_clears_wcag_aa_against_its_own_ground(self, rendered):
        """The acceptance test, and the specific failure this prevents."""
        path, tree = rendered

        for text, color, ground in _runs_with_ground(path, tree):
            ratio = colors.contrast_ratio(color, ground)
            assert ratio >= colors.WCAG_AA_NORMAL, (
                f"{text!r} renders at {color} on {ground} - contrast {ratio:.2f}"
            )

    def test_text_on_the_dark_half_is_light(self, rendered):
        path, tree = rendered

        on_dark = [c for _, c, g in _runs_with_ground(path, tree) if g == NAVY]
        assert on_dark
        assert all(not colors.is_dark(c) for c in on_dark)

    def test_bullets_are_tinted_for_their_own_ground(self):
        theme = deck_theme.DeckTheme.from_design_system(None)

        light = theme.style_on(PAPER)
        dark = theme.style_on(NAVY)
        assert light.bullet_color != dark.bullet_color
        assert colors.contrast_ratio(dark.bullet_color, NAVY) >= 3.0


class TestChartsKnowWhatTheyAreDrawnOn:
    def test_a_chart_on_a_dark_ground_gets_light_furniture(self):
        base = resolve_chart_style(None)

        adapted = _adapt_to_ground(base, NAVY)

        assert adapted.axis_color != base.axis_color
        assert colors.contrast_ratio(adapted.axis_color, NAVY) >= colors.WCAG_AA_NORMAL

    def test_a_chart_on_a_light_ground_is_untouched(self):
        base = resolve_chart_style(None)

        assert _adapt_to_ground(base, PAPER) is base
        assert _adapt_to_ground(base, None) is base
        assert _adapt_to_ground(base, "not a colour") is base


class TestDeckThemeResolvesByGroundNotByRole:
    def test_the_ground_decides_the_treatment(self):
        theme = deck_theme.DeckTheme.from_design_system(None)

        on_dark = theme.style_on(NAVY, deck_theme.ROLE_CONTENT)
        on_light = theme.style_on(PAPER, deck_theme.ROLE_CONTENT)
        assert colors.is_dark(on_dark.background)
        assert not colors.is_dark(on_light.background)
        assert colors.contrast_ratio(on_dark.body_color, NAVY) >= colors.WCAG_AA_NORMAL
        assert colors.contrast_ratio(on_light.body_color, PAPER) >= colors.WCAG_AA_NORMAL

    def test_an_unparseable_ground_falls_back_to_paper(self):
        theme = deck_theme.DeckTheme.from_design_system(None)

        assert theme.style_on("chartreuse-ish").background == theme.white

    @pytest.mark.parametrize("role", ["cover", "section", "content", "closing"])
    def test_style_for_is_unchanged_for_every_role(self, role):
        theme = deck_theme.DeckTheme.from_design_system(None)

        style = theme.style_for(role)
        expected_bg = theme.dark if role in deck_theme.DARK_ROLES else theme.white
        assert style.background == expected_bg
        assert style.role == (role if role in deck_theme.DARK_ROLES
                              else deck_theme.ROLE_CONTENT)

    @pytest.mark.parametrize("role", ["cover", "section", "closing"])
    def test_a_hero_role_survives_a_light_ground(self, role):
        """A design direction may put a cover on a light ground.

        The role used to be rewritten to ``content`` on any light background,
        which cost the slide all three things that key off it: chrome is
        suppressed on ``cover``, the standfirst is set as running lines rather
        than bullets on ``cover``, and the title steps up a size on any hero
        role. A light-ground cover got a footer, a page number, a brand mark
        and a bulleted agenda at body size.
        """
        theme = deck_theme.DeckTheme.from_design_system(None)

        style = theme.style_on(PAPER, role)

        assert style.role == role
        assert not colors.is_dark(style.background)
        assert colors.contrast_ratio(style.body_color, PAPER) >= colors.WCAG_AA_NORMAL

    def test_an_unknown_role_still_falls_back_to_content(self):
        theme = deck_theme.DeckTheme.from_design_system(None)

        assert theme.style_on(PAPER, "exhibit").role == deck_theme.ROLE_CONTENT
        assert theme.style_on(NAVY, "exhibit").role == deck_theme.ROLE_CONTENT


class TestSlidesWithNoRegionGroundsAreUntouched:
    @pytest.fixture(scope="class")
    @staticmethod
    def work_dir():
        directory = tempfile.mkdtemp(prefix="ts_grounds_regress_")
        try:
            yield directory
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @pytest.mark.parametrize("name", sorted(DECKS))
    def test_no_ground_shapes_are_added(self, name, work_dir):
        slides, theme = DECKS[name]
        path = _render(slides, work_dir, name, theme)

        for slide in Presentation(path).slides:
            assert not [s for s in slide.shapes
                        if (s.name or "") == CHROME_REGION_GROUND]

    @pytest.mark.parametrize("name", sorted(DECKS))
    def test_the_package_is_identical_with_an_empty_region_list(self, name, work_dir):
        slides, theme = DECKS[name]
        plain = _render(slides, work_dir, f"{name}-plain", theme)
        empty = _render([{**s, "regions": []} for s in slides],
                        work_dir, f"{name}-empty", theme)

        def parts(path):
            with zipfile.ZipFile(path) as archive:
                return {n: archive.read(n) for n in sorted(archive.namelist())
                        if n.startswith("ppt/slides/slide") and n.endswith(".xml")}

        assert parts(plain) == parts(empty)

    def test_a_region_without_a_ground_paints_nothing(self, work_dir):
        regions = [{"id": "panel", "kind": "panel", "x": 0.1, "y": 0.2,
                    "w": 0.8, "h": 0.6}]
        path = _render([{"slideIndex": 0, "layoutId": "content", "regions": regions,
                         "contents": [{"type": "text", "regionId": "panel",
                                       "runs": [{"text": "No ground here"}]}]}],
                       work_dir, "groundless")

        assert not [s for s in Presentation(path).slides[0].shapes
                    if (s.name or "") == CHROME_REGION_GROUND]
