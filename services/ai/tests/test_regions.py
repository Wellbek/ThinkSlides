"""
The region tree: does a slide's own composition survive to the rendered file,
and does the placeholder path stay exactly where it was?

The second half is the load-bearing one. Regions touch the schema, the
renderer, the validator and every tool that emits slide JSON, and every
session already persisted holds the old shape. So the regression gate here is
not "region-free slides still look fine" - it is that they render *identically*
to a run where the region machinery is not reached at all.
"""

import os
import shutil
import tempfile
import zipfile

import pytest

from app.agent.tools.validate_slide_json import ValidateSlideJsonTool
from app.eval.fixture_decks import DECKS
from app.renderer.pptx_renderer import PPTXRenderer
from app.renderer.regions import (
    RegionTree, content_region_ids, validate_regions,
)

CANVAS_W, CANVAS_H = 13.333, 7.5


def _tree(regions):
    return RegionTree.from_slide({"regions": regions}, CANVAS_W, CANVAS_H)


def _render(slides, work_dir, name="deck", theme="professional"):
    path = os.path.join(work_dir, f"{name}.pptx")
    PPTXRenderer(template_path=None, color_theme=theme).render(
        slides=slides, metadata={"title": "Region Test"}, output_path=path,
    )
    return path


def _slide_xml(pptx_path):
    """Every slide part, which is the deterministic half of the package."""
    with zipfile.ZipFile(pptx_path) as archive:
        return {
            name: archive.read(name)
            for name in sorted(archive.namelist())
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        }


# --------------------------------------------------------------------------
# Reference pages, written out as region trees. These are the proof that the
# model can express the compositions the epic exists to reproduce - each was
# read off the page and is asserted against what the page actually does.
# --------------------------------------------------------------------------

# Page 3: a dark rail one third wide holding the title, an exhibit number and
# a bolded takeaway, with a white panel beside it holding the exhibit.
PAGE_3 = [
    {"id": "rail", "kind": "rail", "x": 0.0, "y": 0.0, "w": 0.34, "h": 1.0,
     "ground": "#051C2C", "children": [
         {"id": "rail-title", "kind": "panel", "x": 0.08, "y": 0.04, "w": 0.84, "h": 0.22},
         {"id": "rail-exhibit-no", "kind": "panel", "x": 0.08, "y": 0.27, "w": 0.84, "h": 0.05},
         {"id": "rail-takeaway", "kind": "panel", "x": 0.08, "y": 0.38, "w": 0.84, "h": 0.30},
     ]},
    {"id": "exhibit", "kind": "panel", "x": 0.34, "y": 0.0, "w": 0.66, "h": 1.0,
     "ground": "#FFFFFF", "children": [
         {"id": "exhibit-title", "kind": "banner", "x": 0.06, "y": 0.13, "w": 0.88, "h": 0.06},
         {"id": "exhibit-plot", "kind": "panel", "x": 0.06, "y": 0.22, "w": 0.60, "h": 0.68},
         {"id": "exhibit-legend", "kind": "panel", "x": 0.70, "y": 0.22, "w": 0.26, "h": 0.68},
     ]},
]

# Page 22: a split canvas - white left, deep navy right - with a chart bound
# to one half and a column of icon stat cards to the other.
PAGE_22 = [
    {"id": "left", "kind": "panel", "x": 0.0, "y": 0.0, "w": 0.5, "h": 1.0,
     "ground": "#FFFFFF", "children": [
         {"id": "left-title", "kind": "panel", "x": 0.09, "y": 0.03, "w": 0.86, "h": 0.16},
         {"id": "left-lead", "kind": "panel", "x": 0.09, "y": 0.24, "w": 0.86, "h": 0.10},
         {"id": "left-chart", "kind": "panel", "x": 0.09, "y": 0.36, "w": 0.86, "h": 0.52},
     ]},
    {"id": "right", "kind": "panel", "x": 0.5, "y": 0.0, "w": 0.5, "h": 1.0,
     "ground": "#051C2C", "children": [
         {"id": "stat-1", "kind": "band", "x": 0.06, "y": 0.17, "w": 0.88, "h": 0.19},
         {"id": "stat-2", "kind": "band", "x": 0.06, "y": 0.38, "w": 0.88, "h": 0.19},
         {"id": "stat-3", "kind": "band", "x": 0.06, "y": 0.59, "w": 0.88, "h": 0.15},
         {"id": "stat-4", "kind": "band", "x": 0.06, "y": 0.76, "w": 0.88, "h": 0.19},
     ]},
]

# Page 40: seven hairline-ruled row bands, each an icon column, a label column
# and a prose column, under a title and a two-line lead.
PAGE_40 = [
    {"id": "title", "kind": "banner", "x": 0.045, "y": 0.07, "w": 0.91, "h": 0.09},
    {"id": "lead", "kind": "banner", "x": 0.045, "y": 0.18, "w": 0.91, "h": 0.10},
    {"id": "bands", "kind": "bandRow", "x": 0.045, "y": 0.30, "w": 0.91, "h": 0.62,
     "children": [
         {"id": f"band-{n}", "kind": "band",
          "x": 0.0, "y": round(n * 0.1428, 4), "w": 1.0, "h": 0.1400,
          "children": [
              {"id": f"band-{n}-icon", "kind": "column", "x": 0.0, "y": 0.0,
               "w": 0.04, "h": 1.0},
              {"id": f"band-{n}-label", "kind": "column", "x": 0.045, "y": 0.0,
               "w": 0.18, "h": 1.0},
              {"id": f"band-{n}-prose", "kind": "column", "x": 0.235, "y": 0.0,
               "w": 0.765, "h": 1.0},
          ]}
         for n in range(7)
     ]},
]

# Page 95: a four-column card row, each card an icon badge over a ruled
# heading over two paragraphs, above a full-width summary banner.
PAGE_95 = [
    {"id": "title", "kind": "banner", "x": 0.045, "y": 0.06, "w": 0.91, "h": 0.09},
    {"id": "lead", "kind": "banner", "x": 0.045, "y": 0.18, "w": 0.91, "h": 0.12},
    {"id": "cards", "kind": "grid", "x": 0.045, "y": 0.33, "w": 0.91, "h": 0.48,
     "children": [
         {"id": f"card-{n}", "kind": "column",
          "x": round(n * 0.253, 4), "y": 0.0, "w": 0.24, "h": 1.0,
          "children": [
              {"id": f"card-{n}-icon", "kind": "panel", "x": 0.0, "y": 0.0,
               "w": 0.30, "h": 0.18},
              {"id": f"card-{n}-heading", "kind": "panel", "x": 0.0, "y": 0.22,
               "w": 1.0, "h": 0.16},
              {"id": f"card-{n}-body", "kind": "panel", "x": 0.0, "y": 0.42,
               "w": 1.0, "h": 0.58},
          ]}
         for n in range(4)
     ]},
    {"id": "banner", "kind": "banner", "x": 0.045, "y": 0.84, "w": 0.91, "h": 0.08},
]

REFERENCE_PAGES = {"3": PAGE_3, "22": PAGE_22, "40": PAGE_40, "95": PAGE_95}


class TestTheReferencePagesAreExpressible:
    @pytest.mark.parametrize("page", sorted(REFERENCE_PAGES))
    def test_the_tree_validates(self, page):
        errors, _ = validate_regions(REFERENCE_PAGES[page])

        assert errors == [], f"page {page}: {errors}"

    @pytest.mark.parametrize("page", sorted(REFERENCE_PAGES))
    def test_every_region_resolves_inside_the_canvas(self, page):
        tree = _tree(REFERENCE_PAGES[page])

        for region_id in tree.ids():
            rect = tree.rect(region_id)
            assert rect.left >= -0.01 and rect.top >= -0.01
            assert rect.left + rect.width <= CANVAS_W + 0.01
            assert rect.top + rect.height <= CANVAS_H + 0.01

    def test_page_3_has_a_rail_one_third_wide(self):
        tree = _tree(PAGE_3)

        rail = tree.get("rail")
        assert rail.kind == "rail"
        assert 0.30 < rail.rect.width / CANVAS_W < 0.38
        assert rail.rect.height == pytest.approx(CANVAS_H)

    def test_page_22_splits_the_canvas_into_two_grounds(self):
        tree = _tree(PAGE_22)

        assert tree.ground("left") == "#FFFFFF"
        assert tree.ground("right") == "#051C2C"
        assert tree.rect("left").width == pytest.approx(tree.rect("right").width)

    def test_page_22_children_inherit_the_ground_of_their_half(self):
        tree = _tree(PAGE_22)

        assert tree.ground("stat-1") == "#051C2C"
        assert tree.ground("left-chart") == "#FFFFFF"

    def test_page_40_has_seven_bands_each_of_three_columns(self):
        tree = _tree(PAGE_40)

        bands = [r for r in tree.resolved.values() if r.kind == "band"]
        assert len(bands) == 7
        for n in range(7):
            assert tree.get(f"band-{n}-icon") is not None
            assert tree.get(f"band-{n}-label") is not None
            assert tree.get(f"band-{n}-prose") is not None

    def test_page_40_bands_stack_without_overlapping(self):
        tree = _tree(PAGE_40)

        tops = [tree.rect(f"band-{n}").top for n in range(7)]
        assert tops == sorted(tops)
        for n in range(6):
            assert tree.rect(f"band-{n}").top + tree.rect(f"band-{n}").height <= (
                tree.rect(f"band-{n + 1}").top + 0.001
            )

    def test_page_95_is_a_four_card_row_over_a_banner(self):
        tree = _tree(PAGE_95)

        cards = [tree.rect(f"card-{n}") for n in range(4)]
        assert len({round(c.width, 3) for c in cards}) == 1
        banner = tree.rect("banner")
        assert banner.top > max(c.top + c.height for c in cards) - 0.01
        assert banner.width / CANVAS_W > 0.85


class TestResolvingGeometry:
    def test_fractions_resolve_against_the_canvas(self):
        tree = _tree([{"id": "half", "kind": "panel", "x": 0.5, "y": 0.0,
                       "w": 0.5, "h": 1.0}])

        rect = tree.rect("half")
        assert rect.left == pytest.approx(CANVAS_W / 2)
        assert rect.width == pytest.approx(CANVAS_W / 2)
        assert rect.height == pytest.approx(CANVAS_H)

    def test_a_child_resolves_against_its_parent_not_the_canvas(self):
        tree = _tree([{"id": "p", "kind": "panel", "x": 0.5, "y": 0.5,
                       "w": 0.5, "h": 0.5, "children": [
                           {"id": "c", "kind": "panel", "x": 0.5, "y": 0.0,
                            "w": 0.5, "h": 1.0}]}])

        rect = tree.rect("c")
        assert rect.left == pytest.approx(CANVAS_W * 0.75)
        assert rect.width == pytest.approx(CANVAS_W * 0.25)

    def test_the_same_tree_scales_with_the_canvas(self):
        wide = RegionTree.from_slide({"regions": PAGE_22}, 20.0, 11.25)
        standard = _tree(PAGE_22)

        assert wide.rect("right").left / 20.0 == pytest.approx(
            standard.rect("right").left / CANVAS_W
        )

    def test_a_slide_with_no_regions_has_no_tree(self):
        assert RegionTree.from_slide({}, CANVAS_W, CANVAS_H) is None
        assert RegionTree.from_slide({"regions": []}, CANVAS_W, CANVAS_H) is None
        assert RegionTree.from_slide({"regions": "nope"}, CANVAS_W, CANVAS_H) is None

    def test_grounds_are_ordered_outermost_first(self):
        tree = _tree([{"id": "panel", "kind": "panel", "x": 0, "y": 0, "w": 1, "h": 1,
                       "ground": "#FFFFFF", "children": [
                           {"id": "rail", "kind": "rail", "x": 0, "y": 0,
                            "w": 0.3, "h": 1.0, "ground": "#051C2C"}]}])

        assert [r.id for r in tree.grounded()] == ["panel", "rail"]


class TestValidationCatchesTheThreeWaysAComposedSlideBreaks:
    def test_overlapping_siblings_are_rejected(self):
        errors, _ = validate_regions([
            {"id": "a", "kind": "panel", "x": 0.0, "y": 0.0, "w": 0.6, "h": 1.0},
            {"id": "b", "kind": "panel", "x": 0.4, "y": 0.0, "w": 0.6, "h": 1.0},
        ])

        assert any("overlap" in e for e in errors)

    def test_two_halves_sharing_an_edge_are_not_an_overlap(self):
        errors, _ = validate_regions([
            {"id": "a", "kind": "panel", "x": 0.0, "y": 0.0, "w": 0.5, "h": 1.0},
            {"id": "b", "kind": "panel", "x": 0.5, "y": 0.0, "w": 0.5, "h": 1.0},
        ])

        assert errors == []

    def test_a_region_escaping_its_parent_is_rejected(self):
        errors, _ = validate_regions([
            {"id": "a", "kind": "panel", "x": 0.6, "y": 0.0, "w": 0.6, "h": 1.0},
        ])

        assert any("escapes its parent" in e for e in errors)

    def test_a_child_escaping_its_parent_is_rejected(self):
        errors, _ = validate_regions([
            {"id": "a", "kind": "panel", "x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0,
             "children": [{"id": "c", "kind": "panel", "x": 0.9, "y": 0.0,
                           "w": 0.5, "h": 1.0}]},
        ])

        assert any("escapes its parent" in e for e in errors)

    def test_a_block_bound_to_a_region_that_does_not_exist_is_rejected(self):
        errors, _ = validate_regions(
            [{"id": "a", "kind": "panel", "x": 0, "y": 0, "w": 1, "h": 1}],
            bound_region_ids=["a", "ghost"],
        )

        assert any("ghost" in e for e in errors)

    def test_duplicate_ids_are_rejected(self):
        errors, _ = validate_regions([
            {"id": "a", "kind": "panel", "x": 0.0, "y": 0.0, "w": 0.5, "h": 1.0},
            {"id": "a", "kind": "panel", "x": 0.5, "y": 0.0, "w": 0.5, "h": 1.0},
        ])

        assert any("duplicate" in e for e in errors)

    def test_a_missing_id_is_rejected(self):
        errors, _ = validate_regions([{"kind": "panel", "x": 0, "y": 0, "w": 1, "h": 1}])

        assert any("missing 'id'" in e for e in errors)

    def test_a_zero_area_region_is_rejected(self):
        errors, _ = validate_regions(
            [{"id": "a", "kind": "panel", "x": 0, "y": 0, "w": 0, "h": 1}]
        )

        assert any("no area" in e for e in errors)

    def test_an_unknown_kind_warns_rather_than_failing(self):
        errors, warnings = validate_regions(
            [{"id": "a", "kind": "carousel", "x": 0, "y": 0, "w": 1, "h": 1}]
        )

        assert errors == []
        assert any("carousel" in w for w in warnings)

    def test_no_regions_is_not_an_error(self):
        assert validate_regions(None) == ([], [])

    def test_content_region_ids_reads_the_bindings(self):
        contents = [
            {"type": "text", "regionId": "left", "runs": []},
            {"type": "chart", "regionId": "right", "chart": {}},
            {"type": "text", "placeholderIdx": 1, "runs": []},
        ]

        assert content_region_ids(contents) == ["left", "right"]


class TestTheValidatorToolSeesRegions:
    @pytest.mark.asyncio
    async def test_a_dangling_binding_fails_validation(self):
        import json

        slide = {
            "slideIndex": 0, "layoutId": "content", "layoutName": "Content",
            "regions": [{"id": "left", "kind": "panel",
                         "x": 0, "y": 0, "w": 0.5, "h": 1.0}],
            "contents": [{"type": "text", "regionId": "right",
                          "runs": [{"text": "hello"}]}],
        }

        result = await ValidateSlideJsonTool.run(json.dumps(slide))

        assert result.success is False
        assert any("right" in e for e in result.data["errors"])

    @pytest.mark.asyncio
    async def test_a_well_formed_composed_slide_passes(self):
        import json

        slide = {
            "slideIndex": 0, "layoutId": "content", "layoutName": "Content",
            "regions": PAGE_22,
            "contents": [
                {"type": "text", "regionId": "left-title",
                 "runs": [{"text": "Why should leaders pay attention?"}]},
                {"type": "text", "regionId": "stat-1",
                 "runs": [{"text": "56%"}]},
            ],
        }

        result = await ValidateSlideJsonTool.run(json.dumps(slide))

        assert result.success is True, result.data["errors"]

    @pytest.mark.asyncio
    async def test_a_slide_with_no_regions_validates_as_before(self):
        import json

        slide = {
            "slideIndex": 0, "layoutId": "content", "layoutName": "Content",
            "contents": [{"type": "text", "placeholderIdx": 1,
                          "runs": [{"text": "hello"}]}],
        }

        result = await ValidateSlideJsonTool.run(json.dumps(slide))

        assert result.success is True


class TestTheRegionFreePathIsUntouched:
    """
    The regression gate. Region support is only additive if a slide that
    declares no regions produces the same file it produced before - not a
    similar one.
    """

    @pytest.fixture(scope="class")
    @staticmethod
    def work_dir():
        directory = tempfile.mkdtemp(prefix="ts_regions_")
        try:
            yield directory
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @pytest.mark.parametrize("name", sorted(DECKS))
    def test_an_empty_region_list_renders_the_identical_package(self, name, work_dir):
        slides, theme = DECKS[name]
        with_none = _render(slides, work_dir, f"{name}-plain", theme)
        with_empty = _render(
            [{**slide, "regions": []} for slide in slides],
            work_dir, f"{name}-empty", theme,
        )

        assert _slide_xml(with_none) == _slide_xml(with_empty)

    def test_an_unresolvable_binding_still_renders_the_content(self, work_dir):
        """Content is never dropped for naming a region that is not there."""
        slides = [{
            "slideIndex": 0, "layoutId": "content",
            "contents": [
                {"type": "text", "runs": [{"text": "Headline", "bullet": False}]},
                {"type": "text", "regionId": "ghost",
                 "runs": [{"text": "Body copy that must survive", "bullet": True}]},
            ],
            "regions": [{"id": "real", "kind": "panel",
                         "x": 0, "y": 0, "w": 1.0, "h": 1.0}],
        }]
        path = _render(slides, work_dir, "dangling")

        from pptx import Presentation
        text = " ".join(
            shape.text_frame.text
            for shape in Presentation(path).slides[0].shapes
            if shape.has_text_frame
        )
        assert "Body copy that must survive" in text


class TestRenderingIntoRegions:
    @pytest.fixture(scope="class")
    @staticmethod
    def work_dir():
        directory = tempfile.mkdtemp(prefix="ts_regions_render_")
        try:
            yield directory
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @staticmethod
    def _shapes(path):
        from pptx import Presentation
        from app.renderer.pptx_renderer import is_chrome_shape
        from app.renderer.xml_utils import emu_to_inches

        out = []
        for shape in Presentation(path).slides[0].shapes:
            if is_chrome_shape(shape) or not shape.has_text_frame:
                continue
            if not shape.text_frame.text.strip():
                continue
            out.append({
                "text": shape.text_frame.text.strip(),
                "left": emu_to_inches(shape.left),
                "top": emu_to_inches(shape.top),
                "width": emu_to_inches(shape.width),
            })
        return out

    def test_a_block_lands_inside_the_region_it_named(self, work_dir):
        slides = [{
            "slideIndex": 0, "layoutId": "content",
            "regions": PAGE_22,
            "contents": [
                {"type": "text", "regionId": "left-title",
                 "runs": [{"text": "On the left", "bullet": False}]},
                {"type": "text", "regionId": "stat-1",
                 "runs": [{"text": "On the right", "bullet": False}]},
            ],
        }]
        path = _render(slides, work_dir, "split")
        tree = _tree(PAGE_22)

        placed = {s["text"]: s for s in self._shapes(path)}
        left, right = placed["On the left"], placed["On the right"]
        assert left["left"] == pytest.approx(tree.rect("left-title").left, abs=0.01)
        assert right["left"] == pytest.approx(tree.rect("stat-1").left, abs=0.01)
        assert right["left"] > CANVAS_W / 2

    def test_two_blocks_in_one_region_stack_inside_it(self, work_dir):
        regions = [{"id": "panel", "kind": "panel",
                    "x": 0.1, "y": 0.2, "w": 0.5, "h": 0.6}]
        slides = [{
            "slideIndex": 0, "layoutId": "content", "regions": regions,
            "contents": [
                {"type": "text", "regionId": "panel",
                 "runs": [{"text": "First block", "bullet": False}]},
                {"type": "text", "regionId": "panel",
                 "runs": [{"text": "Second block", "bullet": False}]},
            ],
        }]
        path = _render(slides, work_dir, "stacked")
        tree = _tree(regions)
        rect = tree.rect("panel")

        placed = {s["text"]: s for s in self._shapes(path)}
        assert placed["First block"]["top"] < placed["Second block"]["top"]
        for shape in placed.values():
            assert shape["left"] >= rect.left - 0.01
            assert shape["top"] >= rect.top - 0.01
            assert shape["top"] <= rect.top + rect.height + 0.01

    def test_text_is_measured_against_its_region_not_the_slide(self, work_dir):
        """
        A block in a narrow rail has to wrap at the rail's width. Measuring
        against the whole content band is what would let a rail's copy run off
        the side of the region it was composed into.
        """
        regions = [{"id": "rail", "kind": "rail",
                    "x": 0.0, "y": 0.1, "w": 0.25, "h": 0.8}]
        slides = [{
            "slideIndex": 0, "layoutId": "content", "regions": regions,
            "contents": [{
                "type": "text", "regionId": "rail",
                "runs": [{"text": "A takeaway sentence long enough to need "
                                  "several lines inside a narrow rail",
                          "bullet": False}],
            }],
        }]
        path = _render(slides, work_dir, "rail")

        shape = self._shapes(path)[0]
        assert shape["width"] == pytest.approx(CANVAS_W * 0.25, abs=0.05)

    def test_a_region_bound_chart_lands_in_its_region(self, work_dir):
        from pptx import Presentation
        from app.renderer.xml_utils import emu_to_inches

        slides = [{
            "slideIndex": 0, "layoutId": "content", "regions": PAGE_22,
            "contents": [
                {"type": "text", "regionId": "left-title",
                 "runs": [{"text": "Adoption", "bullet": False}]},
                {"type": "chart", "regionId": "left-chart", "chart": {
                    "chartType": "bar", "categories": ["A", "B"],
                    "series": [{"name": "s", "values": [1.0, 2.0]}]}},
            ],
        }]
        path = _render(slides, work_dir, "chart-region")
        tree = _tree(PAGE_22)
        expected = tree.rect("left-chart")

        pictures = [s for s in Presentation(path).slides[0].shapes
                    if s.shape_type is not None and s.has_text_frame is False]
        assert pictures, "the chart did not reach the slide"
        picture = pictures[0]
        assert emu_to_inches(picture.left) == pytest.approx(expected.left, abs=0.05)
        assert emu_to_inches(picture.top) == pytest.approx(expected.top, abs=0.05)

    def test_a_composed_slide_carries_the_footer_but_not_the_catalog_rule(self, work_dir):
        """
        The accent rule is placed at a fixed point under a catalog layout's
        title. A composed slide has no title there - on a split canvas the bar
        landed across the lead paragraph - and its composition is already the
        design language the rule stands in for. Chrome that understands
        compositions is #126.
        """
        from pptx import Presentation
        from app.renderer.pptx_renderer import CHROME_ACCENT_RULE
        from app.renderer.primitives import shape_primitive_id

        slides = [
            {"slideIndex": 0, "layoutId": "title",
             "contents": [{"type": "text", "runs": [{"text": "Cover"}]}]},
            {"slideIndex": 1, "layoutId": "content", "regions": PAGE_40,
             "contents": [{"type": "text", "regionId": "title",
                           "runs": [{"text": "Composed", "bullet": False}]}]},
            {"slideIndex": 2, "layoutId": "content",
             "contents": [{"type": "text", "runs": [{"text": "Plain"}]},
                          {"type": "text", "runs": [{"text": "A point", "bullet": True}]}]},
        ]
        path = _render(slides, work_dir, "chrome")
        deck = Presentation(path)

        composed, plain = deck.slides[1], deck.slides[2]
        assert not any((s.name or "") == CHROME_ACCENT_RULE for s in composed.shapes)
        assert any(shape_primitive_id(s) == "chrome-page-number"
                   for s in composed.shapes)
        assert any((s.name or "") == CHROME_ACCENT_RULE for s in plain.shapes)
