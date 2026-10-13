"""Typed deck-wide operations (#376, layer 2).

The probes answered every unplanned request with a script against the
rendered file. These are the same requests as data: placed by anchor, kept
clear of each page's content, reported where they had to be left out, and
applied by every render so that a re-render cannot lose them.
"""
import asyncio
import io
import json

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from app.agent import pipeline, run_context
from app.agent.loop import AgentLoop
from app.agent.tools.apply_deck_operation import ApplyDeckOperationTool
from app.agent.tools.finalize_and_render_pptx import FinalizeAndRenderPptxTool
from app.renderer import deck_ops
from app.renderer.deck_ops import Box, DeckStyle, anchored, pages_for

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
W, H = 13.333, 7.5


def _deck(path, pages=3, text_boxes=(), fill=None):
    """A blank deck; ``text_boxes`` are (page, left, top, width, height,
    text), and ``fill`` paints page 1's ground with a full-page rectangle."""
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(W), Inches(H)
    for _ in range(pages):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    if fill:
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        ground = presentation.slides[0].shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height)
        ground.fill.solid()
        ground.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    for page, left, top, width, height, text in text_boxes:
        box = presentation.slides[page - 1].shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    presentation.save(path)
    return str(path)


def _shapes(path, page, prefix=deck_ops.OVERLAY_PREFIX):
    slide = Presentation(path).slides[page - 1]
    return [s for s in slide.shapes if (s.name or "").startswith(prefix)]


def _box(shape):
    return Box(shape.left / 914400, shape.top / 914400,
               shape.width / 914400, shape.height / 914400)


def _png(width=200, height=100):
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (200, 80, 40)).save(buffer, "PNG")
    return buffer.getvalue()


STYLE = DeckStyle(margin_in=0.5)


# ── the vocabulary ─────────────────────────────────────────────────────────


@pytest.mark.parametrize("raw, said", [
    ({"kind": "overlay", "element": {"type": "text", "text": "x", "color": "teal"}},
     "not a colour"),
    ({"kind": "overlay", "anchor": "left", "element": {"type": "progress"}},
     "top or the bottom"),
    ({"kind": "overlay", "layer": "watermark", "element": {"type": "qr", "data": "x"}},
     "a watermark is text or a picture"),
    ({"kind": "recolour", "to": "orange"}, "not a hex colour"),
    ({"kind": "link", "phrase": "object store", "url": "example.com"}, "not a link"),
    ({"kind": "overlay", "pages": [0], "element": {"type": "text", "text": "x"}},
     "counted from 1"),
    ({"kind": "rotate"}, "kind"),
])
def test_an_operation_that_says_something_impossible_is_refused(raw, said):
    with pytest.raises(ValueError, match=said):
        deck_ops.parse(raw)


def test_anchors_sit_in_from_the_edges():
    assert anchored("top-left", 1, 0.5, W, H, 0.3) == Box(0.3, 0.3, 1, 0.5)
    assert anchored("bottom-right", 1, 0.5, W, H, 0.3) == Box(W - 1.3, H - 0.8, 1, 0.5)
    centre = anchored("center", 2, 1, W, H, 0.3)
    assert (centre.left, centre.top) == pytest.approx(((W - 2) / 2, (H - 1) / 2))
    top = anchored("top", 2, 1, W, H, 0.3)
    assert (top.left, top.top) == pytest.approx(((W - 2) / 2, 0.3))
    right = anchored("right", 2, 1, W, H, 0.3)
    assert (right.left, right.top) == pytest.approx((W - 2.3, (H - 1) / 2))


def test_page_sets():
    assert pages_for("all", 4, frozenset({1})) == [1, 2, 3, 4]
    assert pages_for("all-but-cover", 4, frozenset({1})) == [2, 3, 4]
    # A short deck that opens on content has no cover to skip (#347).
    assert pages_for("all-but-cover", 3, frozenset()) == [1, 2, 3]
    assert pages_for("first", 4, frozenset()) == [1]
    assert pages_for("last", 4, frozenset()) == [4]
    assert pages_for([2, 9, 2], 4, frozenset()) == [2]


# ── overlays ───────────────────────────────────────────────────────────────


def test_an_overlay_lands_at_its_anchor_on_every_page(tmp_path):
    path = _deck(tmp_path / "d.pptx")
    report = deck_ops.apply(path, [{"kind": "overlay", "id": "logo", "anchor": "top-right",
                                    "element": {"type": "image", "asset": "logo.png"}}],
                            STYLE, {"logo.png": _png()})
    assert report[0]["pages"] == [1, 2, 3] and report[0]["skipped"] == []
    for page in (1, 2, 3):
        (shape,) = _shapes(path, page)
        box = _box(shape)
        # On the grid's side margin across, in the margin band down.
        assert box.right == pytest.approx(W - STYLE.side_in, abs=0.01)
        assert box.top == pytest.approx(deck_ops.EDGE_INSET_IN, abs=0.01)
        # Never stretched: the picture keeps its own 2:1 shape.
        assert box.width / box.height == pytest.approx(2.0, rel=0.02)


def test_an_overlay_shrinks_rather_than_cover_text(tmp_path):
    # Copy reaches into the corner a 1in-high logo would need, but not the
    # corner a 0.75in one needs.
    path = _deck(tmp_path / "d.pptx", pages=1,
                 text_boxes=[(1, 8.0, 1.15, 4.0, 1.0, "A headline")])
    report = deck_ops.apply(path, [{"kind": "overlay", "id": "logo", "anchor": "top-right",
                                    "height_in": 1.0,
                                    "element": {"type": "image", "asset": "logo.png"}}],
                            STYLE, {"logo.png": _png()})
    assert report[0]["pages"] == [1]
    (shape,) = _shapes(path, 1)
    assert _box(shape).height == pytest.approx(0.75, abs=0.01)


def test_a_page_with_no_room_goes_without_and_says_why(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=2,
                 text_boxes=[(2, 0.2, 0.2, 6.0, 2.0, "The headline of page two")])
    report = deck_ops.apply(path, [{"kind": "overlay", "id": "logo", "anchor": "top-left",
                                    "element": {"type": "image", "asset": "logo.png"}}],
                            STYLE, {"logo.png": _png()})
    assert report[0]["pages"] == [1]
    (skipped,) = report[0]["skipped"]
    assert skipped["page"] == 2
    assert "The headline of page two" in skipped["reason"]
    assert _shapes(path, 2) == []


def test_two_overlays_at_one_anchor_do_not_stack(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1)
    ops = [{"kind": "overlay", "id": f"t{n}", "anchor": "bottom-left", "height_in": 0.4,
            "element": {"type": "text", "text": f"Footer {n}"}} for n in (1, 2)]
    report = deck_ops.apply(path, ops, STYLE)
    assert report[0]["pages"] == [1]
    assert report[1]["pages"] == []
    assert "another added element" in report[1]["skipped"][0]["reason"]


def test_a_watermark_crosses_the_content_faintly_and_is_marked_as_background(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1,
                 text_boxes=[(1, 1.0, 1.0, 11.0, 5.0, "Body copy across the page")])
    report = deck_ops.apply(path, [{"kind": "overlay", "id": "wm", "layer": "watermark",
                                    "anchor": "center",
                                    "element": {"type": "text", "text": "DRAFT"}}], STYLE)
    assert report[0]["pages"] == [1]
    (mark,) = _shapes(path, 1, deck_ops.BACKGROUND_PREFIX)
    assert mark.rotation == pytest.approx(deck_ops.WATERMARK_ROTATION_DEG % 360)
    alpha = mark._element.find(f".//{A}alpha")
    assert int(alpha.get("val")) == int(deck_ops.WATERMARK_OPACITY * 100000)
    # Drawn after the content, so no filled shape can hide it (the probe's "DRA").
    names = [s.name for s in Presentation(path).slides[0].shapes]
    assert names[-1] == mark.name
    # Its box fits its text, with wrapping off and no auto-fit to re-wrap it.
    body = mark._element.find(f".//{A}bodyPr")
    assert body.get("wrap") == "none" and body.find(f"{A}spAutoFit") is None
    # The run's own validation does not count it as overlapping the copy.
    from app.agent.tools.validate_rendered_output import background_texts

    assert background_texts(path) == {"DRAFT"}


def test_a_footer_numbers_its_pages_and_skips_the_cover(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=3)
    deck_ops.apply(path, [{"kind": "overlay", "id": "foot", "anchor": "bottom-right",
                           "pages": "all-but-cover",
                           "element": {"type": "text", "text": "Acme · {page} / {pages}"}}],
                   STYLE)
    assert _shapes(path, 1) == []
    for page in (2, 3):
        (shape,) = _shapes(path, page)
        field = shape._element.find(f".//{A}fld")
        assert field.get("type") == "slidenum" and field.findtext(f"{A}t") == str(page)
        assert shape.text_frame.text.endswith(f"{page} / 3")


def test_a_footer_ends_where_the_copy_ends(tmp_path):
    """footer-not-cover (2026-10-13): set 0.3in from the edge, the footer
    lined up with nothing on a page whose copy starts 0.85in in."""
    path = _deck(tmp_path / "d.pptx", pages=1)
    deck_ops.apply(path, [{"kind": "overlay", "id": "foot", "anchor": "bottom-right",
                           "element": {"type": "text", "text": "Acme"}}],
                   DeckStyle(margin_in=0.6, side_in=0.85))
    (shape,) = _shapes(path, 1)
    assert _box(shape).right == pytest.approx(W - 0.85, abs=0.01)
    assert _box(shape).bottom == pytest.approx(H - deck_ops.EDGE_INSET_IN, abs=0.01)


def test_all_but_the_cover_on_a_deck_without_one_says_so(tmp_path):
    """A short deck opens on content (#347); asked to skip a cover it does
    not have, the footer goes on every page and the report says why."""
    path = _deck(tmp_path / "d.pptx", pages=3)
    report = deck_ops.apply(path, [{"kind": "overlay", "id": "foot", "anchor": "bottom-right",
                                    "pages": "all-but-cover",
                                    "element": {"type": "text", "text": "{page}"}}],
                            DeckStyle(cover_pages=frozenset()))
    assert report[0]["pages"] == [1, 2, 3]
    assert "no cover" in report[0]["note"]
    with_cover = deck_ops.apply(path, [{"kind": "overlay", "id": "foot",
                                        "anchor": "bottom-right", "pages": "all-but-cover",
                                        "element": {"type": "text", "text": "{page}"}}], STYLE)
    assert "note" not in with_cover[0]


def test_a_progress_bar_is_as_long_as_the_share_of_the_deck(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=4)
    deck_ops.apply(path, [{"kind": "overlay", "id": "pb", "anchor": "bottom",
                           "element": {"type": "progress"}}], STYLE)
    for page in (1, 2, 3, 4):
        (bar,) = _shapes(path, page)
        box = _box(bar)
        assert box.left == 0 and box.bottom == pytest.approx(H, abs=0.01)
        assert box.width == pytest.approx(W * page / 4, abs=0.01)


def test_an_accent_mark_on_an_accent_ground_is_redrawn_to_be_seen(tmp_path):
    """The recolour probe's cover is painted in the accent; an accent bar on
    it was simply not there."""
    from app import colors

    path = _deck(tmp_path / "d.pptx", pages=2, fill=STYLE.accent)
    deck_ops.apply(path, [{"kind": "overlay", "id": "pb", "anchor": "bottom",
                           "element": {"type": "progress"}}], STYLE)
    cover = "#" + str(_shapes(path, 1)[0].fill.fore_color.rgb)
    inner = "#" + str(_shapes(path, 2)[0].fill.fore_color.rgb)
    assert inner.upper() == STYLE.accent.upper()
    assert colors.contrast_ratio(cover, STYLE.accent) >= deck_ops.MARK_CONTRAST


def test_a_qr_code_encodes_its_link_and_is_one(tmp_path):
    import shutil
    import subprocess

    path = _deck(tmp_path / "d.pptx", pages=2)
    url = "https://example.com/deck"
    deck_ops.apply(path, [{"kind": "overlay", "id": "qr", "anchor": "bottom-left",
                           "pages": "last", "element": {"type": "qr", "data": url}}], STYLE)
    assert _shapes(path, 1) == []
    (code,) = _shapes(path, 2)
    assert code.click_action.hyperlink.address == url
    assert _box(code).width == pytest.approx(1.0, abs=0.01)
    if shutil.which("zbarimg"):
        png = tmp_path / "qr.png"
        png.write_bytes(code.image.blob)
        decoded = subprocess.run(["zbarimg", "-q", "--raw", str(png)],
                                 capture_output=True, text=True).stdout.strip()
        assert decoded == url


def test_applying_again_replaces_rather_than_adds(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=2)
    op = {"kind": "overlay", "id": "foot", "anchor": "bottom-left",
          "element": {"type": "text", "text": "Acme"}}
    deck_ops.apply(path, [op], STYLE)
    deck_ops.apply(path, [op], STYLE)
    assert len(_shapes(path, 1)) == 1 and len(_shapes(path, 2)) == 1


def test_a_missing_picture_costs_that_operation_only(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1)
    report = deck_ops.apply(path, [
        {"kind": "overlay", "id": "logo", "element": {"type": "image", "asset": "cup.png"}},
        {"kind": "overlay", "id": "foot", "anchor": "bottom-left",
         "element": {"type": "text", "text": "Acme"}},
    ], STYLE, {})
    assert "no picture named 'cup.png'" in report[0]["problem"]
    assert report[1]["pages"] == [1]


# ── recolour and links ─────────────────────────────────────────────────────


def test_a_recolour_reaches_the_page_and_the_theme(tmp_path):
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    path = _deck(tmp_path / "d.pptx", pages=1, fill=STYLE.accent)
    report = deck_ops.apply(path, [{"kind": "recolour", "id": "r", "role": "accent",
                                    "to": "#e4572e"}], STYLE)
    assert report[0]["repainted"] >= 1
    presentation = Presentation(path)
    assert str(presentation.slides[0].shapes[0].fill.fore_color.rgb) == "E4572E"
    theme = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    slot = etree.fromstring(theme.blob).find(f".//{A}accent1/{A}srgbClr")
    assert slot.get("val") == "E4572E"


def test_an_overlay_in_the_accent_wears_the_new_one(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1)
    deck_ops.apply(path, [
        {"kind": "overlay", "id": "dot", "anchor": "top-left",
         "element": {"type": "shape", "shape": "ellipse"}},
        {"kind": "recolour", "id": "r", "role": "accent", "to": "#E4572E"},
    ], STYLE)
    (dot,) = _shapes(path, 1)
    assert str(dot.fill.fore_color.rgb) == "E4572E"


def test_a_link_covers_only_the_phrase_and_keeps_its_colour(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=2, text_boxes=[
        (1, 1, 1, 8, 1, "From a shared disk to an Object Store."),
        (2, 1, 1, 8, 1, "Nothing to link here."),
    ])
    report = deck_ops.apply(path, [{"kind": "link", "id": "l", "phrase": "object store",
                                    "url": "https://example.com/os"}], STYLE)
    assert report[0]["pages"] == [1] and report[0]["linked"] == 1
    slide = Presentation(path).slides[0]
    runs = [r for p in slide.shapes[0].text_frame.paragraphs for r in p.runs]
    assert [r.text for r in runs] == ["From a shared disk to an ", "Object Store", "."]
    linked = runs[1]
    assert linked.hyperlink.address == "https://example.com/os"
    assert linked.font.underline is True
    click = linked._r.find(f"{A}rPr/{A}hlinkClick")
    assert click.find(f".//{{{deck_ops._HLINK_COLOR_NS}}}hlinkClr").get("val") == "tx"
    assert runs[0].hyperlink.address is None and runs[2].hyperlink.address is None


def test_a_link_to_a_phrase_the_deck_never_uses_is_reported(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1, text_boxes=[(1, 1, 1, 8, 1, "Hello")])
    report = deck_ops.apply(path, [{"kind": "link", "id": "l", "phrase": "object store",
                                    "url": "https://example.com"}], STYLE)
    assert "never mentions" in report[0]["problem"]


# ── the tool ───────────────────────────────────────────────────────────────


def test_the_tool_checks_and_hands_back_every_part():
    result = asyncio.run(ApplyDeckOperationTool.run(operation=[
        {"kind": "overlay", "anchor": "bottom-right",
         "element": {"type": "text", "text": "{page}"}},
        {"kind": "recolour", "to": "#E4572E"},
    ]))
    assert result.success
    assert [op["kind"] for op in result.data["operations"]] == ["overlay", "recolour"]


def test_one_bad_part_records_nothing():
    result = asyncio.run(ApplyDeckOperationTool.run(operation=[
        {"kind": "recolour", "to": "#E4572E"},
        {"kind": "overlay", "element": {"type": "text", "text": "x", "color": "teal"}},
    ]))
    assert not result.success and "Operation 2" in result.error


def test_the_tool_refuses_a_picture_nobody_supplied():
    token = run_context.set_assets({"logo.png": _png()})
    try:
        result = asyncio.run(ApplyDeckOperationTool.run(operation={
            "kind": "overlay", "element": {"type": "image", "asset": "cup.png"}}))
    finally:
        run_context._assets.reset(token)
    assert not result.success and "logo.png" in result.error


# ── the run ────────────────────────────────────────────────────────────────


def _loop():
    loop = AgentLoop(db=None, mode=None)
    loop.questions_declined = True
    return loop


def test_the_loop_records_replaces_and_removes_operations():
    loop = _loop()
    ws = loop.working_state
    serial = ws.design_serial
    footer = {"kind": "overlay", "element": {"type": "text", "text": "Acme"}}
    loop._record_deck_operation({"action": "add", "operations": [footer, dict(footer)]})
    assert [op["id"] for op in ws.deck_operations] == ["op1", "op2"]
    assert ws.design_serial == serial + 1
    loop._record_deck_operation({"action": "add", "operations": [
        {**footer, "id": "op1", "anchor": "bottom-left"}]})
    assert [op["id"] for op in ws.deck_operations] == ["op2", "op1"]
    assert ws.deck_operations[1]["anchor"] == "bottom-left"
    loop._record_deck_operation({"action": "remove", "id": "op2"})
    assert [op["id"] for op in ws.deck_operations] == ["op1"]
    # A removal of nothing changes nothing, and renders nothing again.
    before = ws.design_serial
    loop._record_deck_operation({"action": "remove", "id": "nope"})
    assert ws.design_serial == before


def test_operations_survive_a_pause():
    loop = _loop()
    loop._record_deck_operation({"action": "add", "operations": [
        {"kind": "recolour", "to": "#E4572E"}]})
    other = _loop()
    assert other._apply_state_dict(loop._state_dict(), "sid")
    assert other.working_state.deck_operations == loop.working_state.deck_operations


def test_the_change_step_offers_the_operations_and_names_what_is_there(monkeypatch):
    loop = _loop()
    ws = loop.working_state
    ws.amending, ws.amend_request = True, "Remove the watermark"
    ws.deck_operations = [{"kind": "overlay", "id": "op1", "layer": "watermark",
                           "anchor": "center", "pages": "all",
                           "element": {"type": "text", "text": "DRAFT"}}]
    monkeypatch.setattr(type(loop), "code_amendment_enabled", property(lambda self: False))
    phase = pipeline.PHASES_BY_NAME["change"]
    assert "apply_deck_operation" in phase.tools
    directive = phase.directive(loop)
    assert "apply_deck_operation" in directive
    assert "op1: text 'DRAFT' at the center of all pages as a watermark" in directive
    # Withheld tools are not named.
    assert "edit_deck_with_code" not in directive


def test_an_operation_during_an_amendment_finishes_the_change():
    loop = _loop()
    ws = loop.working_state
    ws.amending = True
    ws.has_called_finalize = True
    ws.rendered_serial = ws.design_serial
    loop._record_deck_operation({"action": "add", "operations": [
        {"kind": "link", "phrase": "object store", "url": "https://example.com"}]})
    assert ws.amend_done and ws.amend_edited
    # And the deck is rendered again, which is where it is applied.
    assert not ws.render_is_current


def test_the_closing_note_names_pages_an_overlay_was_left_off():
    loop = _loop()
    loop.working_state.deck_operation_report = [
        {"id": "op1", "kind": "overlay", "pages": [1],
         "skipped": [{"page": 2, "reason": "it would cover the text 'Revenue'"}]},
        {"id": "op2", "kind": "link", "problem": "the deck never mentions 'x'"},
    ]
    note = loop._deck_operation_note()
    assert "left off page 2, because it would cover the text 'Revenue'" in note
    assert "op2 (link): not applied - the deck never mentions 'x'" in note
    loop.working_state.deck_operation_report = [{"id": "op1", "kind": "overlay",
                                                 "pages": [1, 2], "skipped": []}]
    assert loop._deck_operation_note() == ""
    loop.working_state.deck_operation_report = [{
        "id": "op1", "kind": "overlay", "pages": [1, 2], "skipped": [],
        "note": "the deck has no cover, so it is on every page"}]
    note = loop._deck_operation_note()
    assert "how the request was read" in note and "no cover" in note
    assert "could not go everywhere" not in note


# ── the render ─────────────────────────────────────────────────────────────


def _render(tmp_path, name, operations, monkeypatch):
    from app.design import run as design_run
    from app.renderer.specimen import language_specimen_slides

    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr("app.agent.tools.finalize_and_render_pptx.PreviewGenerator",
                        type("NoPreviews", (), {"generate_previews": lambda self, *a: []}))
    pack = design_run.load("wti-2025")
    tokens = (run_context.set_pack(pack), run_context.set_deck_operations(operations))
    try:
        result = asyncio.run(FinalizeAndRenderPptxTool.run(
            slides_json=json.dumps(language_specimen_slides()), session_id=name))
    finally:
        run_context.reset_pack(tokens[0])
        run_context._deck_operations.reset(tokens[1])
    assert result.success, result.error
    return result, pack


def test_every_render_applies_the_recorded_operations(tmp_path, monkeypatch):
    operations = [
        {"kind": "overlay", "id": "foot", "anchor": "bottom-right", "pages": "all-but-cover",
         "element": {"type": "text", "text": "Acme · {page}"}},
        {"kind": "recolour", "id": "r", "role": "accent", "to": "#E4572E"},
    ]
    first, pack = _render(tmp_path, "one", operations, monkeypatch)
    second, _ = _render(tmp_path, "two", operations, monkeypatch)
    for result in (first, second):
        path = result.data["filePath"]
        report = {entry["id"]: entry for entry in result.data["deckOperations"]}
        pages = len(Presentation(path).slides)
        placed = set(report["foot"]["pages"])
        skipped = {entry["page"] for entry in report["foot"]["skipped"]}
        # Every page but the cover - the specimen opens on one, which
        # composition sets as a content page for the copy it carries - either
        # carries the footer or says why not.
        assert placed | skipped == set(range(2, pages + 1))
        assert placed and all(_shapes(path, page) for page in placed)
        # The pack's own accent is gone from the pages.
        old = pack.colors.accent.lstrip("#").upper()
        xml = "".join(slide.part.blob.decode() for slide in Presentation(path).slides)
        assert f'val="{old}"' not in xml


def test_a_render_without_operations_reports_none(tmp_path, monkeypatch):
    result, _ = _render(tmp_path, "plain", [], monkeypatch)
    assert result.data["deckOperations"] == []


# ── found in review ────────────────────────────────────────────────────────


@pytest.mark.parametrize("slide, cover", [
    ({"layoutName": "Title and Content"}, False),
    ({"layoutName": "Title Only"}, False),
    ({"layoutName": "Title Slide"}, True),
    ({"contentType": "title"}, True),
    ({"role": "cover"}, True),
])
def test_only_a_title_page_is_a_cover_on_the_template_path(slide, cover):
    assert FinalizeAndRenderPptxTool._is_cover(slide, 0) is cover
    assert FinalizeAndRenderPptxTool._is_cover(slide, 1) is False


def test_a_link_is_a_whole_word_and_never_spans_a_break(tmp_path):
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls

    path = _deck(tmp_path / "d.pptx", pages=1, text_boxes=[
        (1, 1, 1, 8, 1, "We maintain AI tooling, as Acme said.")])
    presentation = Presentation(path)
    paragraph = presentation.slides[0].shapes[0].text_frame.paragraphs[0]._p
    # "Acme" then a soft break then "Corp": not a mention of "Acme Corp".
    paragraph.append(parse_xml(f'<a:br {nsdecls("a")}/>'))
    run = parse_xml(f'<a:r {nsdecls("a")}><a:t>Corp results</a:t></a:r>')
    paragraph.append(run)
    presentation.save(path)
    report = deck_ops.apply(path, [
        {"kind": "link", "id": "ai", "phrase": "AI", "url": "https://example.com/ai"},
        {"kind": "link", "id": "acme", "phrase": "said. Corp", "url": "https://example.com"},
    ], STYLE)
    assert report[0]["linked"] == 1
    runs = Presentation(path).slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert [r.text for r in runs if r.hyperlink.address] == ["AI"]
    assert "never mentions" in report[1]["problem"]


def test_a_master_logo_is_content_an_overlay_keeps_clear_of(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1)
    presentation = Presentation(path)
    # A template's own wordmark, set in its master: drawn on the slide, then
    # moved into the master's shape tree.
    mark = presentation.slides[0].shapes.add_textbox(
        Inches(W - 0.5 - 2.0), Inches(0.2), Inches(2.0), Inches(1.0))
    mark.text_frame.text = "ACME"
    presentation.slide_masters[0].shapes._spTree.append(mark._element)
    presentation.save(path)
    report = deck_ops.apply(path, [{"kind": "overlay", "id": "logo", "anchor": "top-right",
                                    "element": {"type": "image", "asset": "logo.png"}}],
                            STYLE, {"logo.png": _png()})
    assert report[0]["pages"] == []
    assert "the text 'ACME'" in report[0]["skipped"][0]["reason"]


def test_a_watermark_does_not_crowd_out_a_later_overlay(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1)
    report = deck_ops.apply(path, [
        {"kind": "overlay", "id": "wm", "layer": "watermark", "anchor": "center",
         "element": {"type": "text", "text": "CONFIDENTIAL"}},
        {"kind": "overlay", "id": "dot", "anchor": "center",
         "element": {"type": "shape", "shape": "ellipse"}},
    ], STYLE)
    assert report[1]["pages"] == [1]


def test_reapplying_one_operation_leaves_another_with_a_longer_id(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1)
    ops = [{"kind": "overlay", "id": "logo-footer", "anchor": "bottom-left",
            "element": {"type": "text", "text": "Footer"}},
           {"kind": "overlay", "id": "logo", "anchor": "top-left",
            "element": {"type": "text", "text": "Logo"}}]
    deck_ops.apply(path, ops, STYLE)
    names = sorted(s.name for s in _shapes(path, 1))
    assert names == ["ts-overlay-op-logo-1", "ts-overlay-op-logo-footer-1"]


def test_removing_an_operation_the_deck_does_not_carry_is_refused():
    token = run_context.set_deck_operations([{"id": "op1", "kind": "recolour"}])
    try:
        result = asyncio.run(ApplyDeckOperationTool.run(action="remove", id="op 1"))
        kept = asyncio.run(ApplyDeckOperationTool.run(action="remove", id="op1"))
    finally:
        run_context._deck_operations.reset(token)
    assert not result.success and "op1" in result.error
    assert kept.success


def test_an_image_operation_keeps_its_picture_when_names_move(tmp_path):
    logo, photo = _png(200, 100), _png(100, 200)
    loop = _loop()
    token = run_context.set_assets({"picture_1.png": logo})
    try:
        loop._record_deck_operation({"action": "add", "operations": [
            {"kind": "overlay", "anchor": "top-right",
             "element": {"type": "image", "asset": "picture_1.png"}}]})
    finally:
        run_context._assets.reset(token)
    op = loop.working_state.deck_operations[0]
    path = _deck(tmp_path / "d.pptx", pages=1)
    # A later upload took the name; the operation still finds the logo.
    report = deck_ops.apply(path, [op], STYLE,
                            {"picture_1.png": photo, "picture_2.png": logo})
    assert report[0]["pages"] == [1]
    (shape,) = _shapes(path, 1)
    assert shape.image.blob == logo


def test_a_recolour_leaves_a_colour_another_role_shares(tmp_path):
    shared = DeckStyle(ink="#1F4E79", accent="#1F4E79")
    path = _deck(tmp_path / "d.pptx", pages=1, fill="#1F4E79",
                 text_boxes=[(1, 1, 1, 8, 1, "Body")])
    deck_ops.apply(path, [{"kind": "recolour", "id": "r", "role": "accent",
                           "to": "#E4572E"}], shared)
    assert str(Presentation(path).slides[0].shapes[0].fill.fore_color.rgb) == "1F4E79"


def test_the_last_page_is_the_deck_not_its_appendix(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=4)
    report = deck_ops.apply(path, [
        {"kind": "overlay", "id": "qr", "anchor": "bottom-left", "pages": "last",
         "element": {"type": "qr", "data": "https://example.com"}},
        {"kind": "overlay", "id": "n", "anchor": "top-left", "pages": [2, 9],
         "element": {"type": "text", "text": "x"}},
    ], DeckStyle(appended_pages=frozenset({4})))
    assert report[0]["pages"] == [3]
    assert report[1]["pages"] == [2]
    assert report[1]["skipped"] == [{"page": 9, "reason": "the deck has 4 pages"}]


# ── reading the deck, for code ─────────────────────────────────────────────


def test_a_script_is_told_every_box_the_engine_places_by(tmp_path):
    """What the overlay placement reads, handed to a script: the slide's
    content, the master's wordmark that `slide.shapes` never lists, and the
    clear bands above and below them."""
    path = _deck(tmp_path / "d.pptx", pages=2,
                 text_boxes=[(1, 1.0, 1.5, 6.0, 4.0, "Revenue held at 14%")])
    presentation = Presentation(path)
    mark = presentation.slides[0].shapes.add_textbox(
        Inches(W - 2.5), Inches(0.4), Inches(2.0), Inches(0.6))
    mark.text_frame.text = "ACME"
    presentation.slide_masters[0].shapes._spTree.append(mark._element)
    presentation.save(path)

    style = DeckStyle(accent="#E4572E", body_floor_pt=16.2, caption_floor_pt=10.5,
                      cover_pages=frozenset({1}))
    with open(path, "rb") as handle:
        layout = deck_ops.describe(handle.read(), style)

    assert layout["page"] == {"width": W, "height": H}
    assert layout["colours"]["accent"] == "#E4572E"
    assert layout["type"]["caption_floor_pt"] == 10.5
    first, second = layout["pages"]
    assert (first["cover"], second["cover"]) == (True, False)
    texts = {shape.get("text"): shape for shape in first["shapes"]}
    assert texts["Revenue held at 14%"]["blocks"] is True
    assert texts["Revenue held at 14%"]["box"] == [1.0, 1.5, 6.0, 4.0]
    assert texts["ACME"]["inherited"] is True and texts["ACME"]["blocks"] is True
    # Nothing offered covers content, inherited or not.
    for page in layout["pages"]:
        content = [Box(*shape["box"]) for shape in page["shapes"] if shape["blocks"]]
        for box in page["free"]["top"] + page["free"]["bottom"]:
            assert not any(Box(*box).overlaps(other) for other in content), box
    # Above the copy, the widest box stops short of the wordmark, whose own
    # corner leaves less than a caption line. Below, it runs from side
    # margin to side margin, from the copy to the edge inset.
    widest_top = max(first["free"]["top"], key=lambda box: box[2])
    assert widest_top == pytest.approx([0.5, 0.3, W - 3.0, 1.18], abs=0.03)
    widest_bottom = max(first["free"]["bottom"], key=lambda box: box[2])
    assert widest_bottom == pytest.approx([0.5, 5.52, W - 1.0, 1.68], abs=0.03)
    # On a page with only the wordmark, the whole top beside it is free.
    assert max(second["free"]["top"], key=lambda box: box[3]) == pytest.approx(
        [0.5, 0.3, W - 3.0, H - 0.6], abs=0.03)


def test_content_down_one_side_leaves_the_other_side_free(tmp_path):
    """A picture to the bottom edge on the right half: the left half's
    bottom is still free, which one band across the page would deny."""
    path = _deck(tmp_path / "d.pptx", pages=1, text_boxes=[
        (1, 0.5, 0.5, 5.0, 4.0, "Copy on the left"),
        (1, 7.0, 0.5, 6.0, 6.95, "A figure down the right")])
    with open(path, "rb") as handle:
        wide, tall = deck_ops.describe(handle.read(), STYLE)["pages"][0]["free"]["bottom"]
    # Under the copy, out to the figure: the widest box a footer fits in.
    left, top, width, height = wide
    assert left == pytest.approx(0.5) and left + width == pytest.approx(7.0, abs=0.03)
    assert top == pytest.approx(4.5, abs=0.05) and top + height == pytest.approx(H - 0.3)
    # And the gutter between them, clear to the top inset.
    assert tall[0] == pytest.approx(5.5) and tall[3] == pytest.approx(H - 0.6)


def test_a_band_too_thin_to_hold_anything_is_not_offered(tmp_path):
    path = _deck(tmp_path / "d.pptx", pages=1,
                 text_boxes=[(1, 0.5, 0.32, 12.0, 7.0, "Full page")])
    with open(path, "rb") as handle:
        page = deck_ops.describe(handle.read(), STYLE)["pages"][0]
    assert page["free"] == {"top": [], "bottom": []}


def test_the_style_survives_its_round_trip_through_json():
    style = DeckStyle(accent="#E4572E", cover_pages=frozenset({1, 3}),
                      appended_pages=frozenset({9}), body_floor_pt=16.2,
                      previous={"accent": frozenset({"#1F4E79"})})
    raw = json.loads(json.dumps(deck_ops.style_to_dict(style)))
    assert deck_ops.style_from_dict(raw) == style
    # A sidecar from a render before a field existed still reads.
    assert deck_ops.style_from_dict({"accent": "#000000", "gone": 1}).accent == "#000000"


def test_every_render_writes_the_style_it_was_placed_by(tmp_path, monkeypatch):
    operations = [{"kind": "recolour", "id": "r", "role": "accent", "to": "#E4572E"}]
    result, pack = _render(tmp_path, "styled", operations, monkeypatch)
    with open(deck_ops.style_path(result.data["filePath"]), encoding="utf-8") as handle:
        style = deck_ops.style_from_dict(json.load(handle))
    # As the deck now wears it: the recolour, and the pack's own floors.
    assert style.accent == "#E4572E"
    assert style.body_floor_pt == round(pack.policy.floor_pt("body", pack.canvas), 1)
    assert style.caption_floor_pt == round(pack.policy.floor_pt("caption", pack.canvas), 1)
    assert style.cover_pages


def test_a_render_with_no_operations_still_writes_its_style(tmp_path, monkeypatch):
    result, pack = _render(tmp_path, "plain", [], monkeypatch)
    with open(deck_ops.style_path(result.data["filePath"]), encoding="utf-8") as handle:
        assert json.load(handle)["accent"] == pack.colors.accent
