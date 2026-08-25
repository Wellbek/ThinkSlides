"""
An SVG document, read as native PowerPoint shapes.

`build_diagram` offers eight typed kinds and the primitive layer beneath it
had five, none of them a path - so there was no lower level to drop to when a
figure was not one of the eight. A state machine, a sequence diagram or a
bespoke quadrant simply could not be drawn.

The property that makes this worth building rather than rasterising: what
lands on the slide is *shapes*. A user can select a node, recolour it and edit
its label. A picture of a diagram is something the deck could already produce.
"""
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.renderer.layout import PositionConfig
from app.renderer.primitives import PaintContext, paint_all
from app.renderer.svg.importer import UnsupportedSVG, read_viewbox, to_primitives
from app.renderer.svg.paths import parse

PALETTE = {"accent": "#1E8449", "accent_alt": "#F4D03F",
           "ink": "#2E2E38", "ground": "#FFFFFF"}

FLOW = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 320 180">
<rect x="20" y="30" width="110" height="56" fill="#7C4DFF"/>
<text x="75" y="64" font-size="14" text-anchor="middle" fill="#ffffff">Ingest</text>
<path d="M130,58 L190,58" stroke="#333333" stroke-width="2" fill="none"/>
<polygon points="190,52 202,58 190,64" fill="#333333"/>
<rect x="202" y="30" width="100" height="56" fill="#00C853"/>
<text x="252" y="64" font-size="14" text-anchor="middle" fill="#ffffff">Compose</text>
</svg>"""


def _box():
    return PositionConfig(left=1.5, top=1.5, width=10.0, height=4.5)


class TestACurveBecomesSegmentsPowerPointCanDraw:
    def test_a_straight_path_keeps_its_points(self):
        [sub] = parse("M0,0 L100,0 L50,80 Z")
        assert len(sub.points) == 3
        assert sub.closed

    def test_a_curve_is_subdivided_until_it_looks_like_one(self):
        """`build_freeform` takes line segments and nothing else."""
        [sub] = parse("M0,0 C0,50 100,50 100,0", tolerance=0.25)
        assert len(sub.points) > 8

    def test_a_gentle_curve_costs_fewer_points_than_a_tight_one(self):
        """Adaptive, not a fixed step: a fixed step either wastes points on
        the gentle curve or visibly facets the tight one."""
        gentle = parse("M0,0 C10,2 90,2 100,0", tolerance=0.25)[0]
        tight = parse("M0,0 C0,200 100,200 100,0", tolerance=0.25)[0]
        assert len(gentle.points) < len(tight.points)

    def test_an_arc_is_flattened_too(self):
        [sub] = parse("M0,0 A50,50 0 0 1 100,0", tolerance=0.25)
        assert len(sub.points) > 8

    def test_a_pathological_path_cannot_produce_endless_points(self):
        sub = parse("M0,0 C0,100000 100,-100000 100,0", tolerance=0.0001)[0]
        assert len(sub.points) < 5000

    def test_garbage_yields_nothing_rather_than_raising(self):
        """Half a figure drawn is better than a deck that fails."""
        assert parse("not a path at all") == []
        assert parse("") == []

    def test_a_truncated_path_keeps_what_it_could_read(self):
        subs = parse("M0,0 L10,10 L20")
        assert subs and len(subs[0].points) >= 2


class TestTheDocumentHasToSayWhatShapeItIs:
    def test_a_viewbox_is_read(self):
        import xml.etree.ElementTree as ET

        root = ET.fromstring('<svg xmlns="http://www.w3.org/2000/svg" '
                             'viewBox="0 0 200 100"/>')
        assert read_viewbox(root).aspect == pytest.approx(2.0)

    def test_width_and_height_stand_in_for_one(self):
        import xml.etree.ElementTree as ET

        root = ET.fromstring('<svg xmlns="http://www.w3.org/2000/svg" '
                             'width="300" height="100"/>')
        assert read_viewbox(root).aspect == pytest.approx(3.0)

    def test_a_document_with_neither_is_refused_by_name(self):
        with pytest.raises(UnsupportedSVG) as excinfo:
            to_primitives('<svg xmlns="http://www.w3.org/2000/svg"/>', _box())
        assert "viewBox" in str(excinfo.value)

    def test_something_that_is_not_svg_is_refused(self):
        with pytest.raises(UnsupportedSVG):
            to_primitives("<html><body>hello</body></html>", _box())
        with pytest.raises(UnsupportedSVG):
            to_primitives("{not xml at all", _box())

    def test_a_figure_is_fitted_without_being_distorted(self):
        """An SVG has a shape of its own, and stretching a diagram to a slot's
        aspect is how a circle becomes an ellipse."""
        svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
               '<circle cx="50" cy="50" r="40" fill="#333"/></svg>')
        [prim] = to_primitives(svg, _box())
        assert prim.box.width == pytest.approx(prim.box.height)
        assert prim.box.width <= _box().width


class TestWhatLandsOnTheSlideIsShapes:
    def test_every_element_becomes_a_primitive(self):
        prims = to_primitives(FLOW, _box(), palette=PALETTE)
        assert [p.kind for p in prims].count("path") == 4
        assert [p.kind for p in prims].count("text") == 2

    def test_they_paint_as_freeforms_and_text_boxes_not_pictures(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        paint_all(slide, to_primitives(FLOW, _box(), palette=PALETTE),
                  PaintContext(ground="#FFFFFF", default_color="#2E2E38",
                               type_scale=1.0))
        kinds = [str(shape.shape_type) for shape in slide.shapes]
        assert sum("FREEFORM" in k for k in kinds) == 4
        assert sum("TEXT_BOX" in k for k in kinds) == 2
        assert not any("PICTURE" in k for k in kinds)

    def test_a_label_stays_editable_text(self):
        prims = to_primitives(FLOW, _box(), palette=PALETTE)
        labels = [p.runs[0].text for p in prims if p.kind == "text"]
        assert labels == ["Ingest", "Compose"]

    def test_a_closed_shape_is_filled_and_an_open_one_is_not(self):
        prims = to_primitives(FLOW, _box(), palette=PALETTE)
        rects = [p for p in prims if p.kind == "path" and p.closed[0]]
        strokes = [p for p in prims if p.kind == "path" and not p.closed[0]]
        assert all(p.fill for p in rects)
        assert all(p.stroke for p in strokes)


class TestAnImportedFigureIsInTheDecksColours:
    def test_a_source_colour_is_mapped_onto_the_pack(self):
        """A figure that arrives in Mermaid's default lilac and stays lilac is
        a figure that looks pasted in."""
        prims = to_primitives(FLOW, _box(), palette=PALETTE)
        fills = {p.fill for p in prims if p.kind == "path" and p.fill}
        assert "#7C4DFF" not in fills
        assert PALETTE["accent"] in fills

    def test_a_grey_stays_grey(self):
        """A rule and a hairline are structural rather than branded, and
        mapping a grey connector onto an accent is how a diagram ends up with
        green hairlines that mean nothing."""
        prims = to_primitives(FLOW, _box(), palette=PALETTE)
        strokes = {p.stroke for p in prims if p.stroke}
        assert "#333333" in strokes

    def test_a_deck_with_no_palette_keeps_the_source_colours(self):
        prims = to_primitives(FLOW, _box())
        assert "#7C4DFF" in {p.fill for p in prims if p.fill}

    def test_a_knockout_label_knows_what_it_sits_on(self):
        """Told the slide is white, the contrast guard flips a white label to
        dark and the node loses its knockout type."""
        prims = to_primitives(FLOW, _box(), palette=PALETTE)
        labels = [p for p in prims if p.kind == "text"]
        assert all(p.color == "#FFFFFF" for p in labels)
        assert all(p.ground == PALETTE["accent"] for p in labels)


class TestNothingExecutes:
    @pytest.mark.parametrize("hostile", [
        '<script>alert(1)</script>',
        '<foreignObject><body onload="x()"/></foreignObject>',
        '<image href="https://evil.test/a.png" width="10" height="10"/>',
        '<use href="https://evil.test/x.svg#a"/>',
    ])
    def test_a_dangerous_element_is_dropped(self, hostile):
        """An SVG can arrive from a model, from an MCP server a customer
        configured, or from a file a user uploaded, and none of those are
        places to run code from."""
        svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
               f'{hostile}<rect x="0" y="0" width="10" height="10" fill="#333"/>'
               '</svg>')
        prims = to_primitives(svg, _box())
        assert len(prims) == 1
        assert prims[0].kind == "path"

    def test_a_child_of_a_dropped_element_goes_with_it(self):
        svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
               '<foreignObject><text x="1" y="1">smuggled</text></foreignObject>'
               '</svg>')
        assert to_primitives(svg, _box()) == []

    def test_a_gradient_reference_is_not_guessed_at(self):
        """A figure drawn in a colour we invented is worse than one drawn in
        the pack's own."""
        svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
               '<rect x="0" y="0" width="10" height="10" fill="url(#grad)"/>'
               '</svg>')
        [prim] = to_primitives(svg, _box(), palette=PALETTE)
        assert prim.fill == "#000000"
