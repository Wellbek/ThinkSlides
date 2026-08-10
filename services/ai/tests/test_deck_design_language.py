"""
The built-in deck's design language: light/dark rhythm, the repeated accent
motif, the footer, and the catalog geometry actually reaching the slide.

Issue #56's finding was that the design system's palette, typography and
layout catalog were all extracted, shown to the model, and then discarded -
the only token that reached a rendered slide was ``primary``, as the title's
text colour. These tests assert the tokens land.
"""

import os
import tempfile

import pytest
from pptx import Presentation
from pptx.util import Emu

from app import colors
from app.design_system import get_default_design_system
from app.renderer.pptx_renderer import (
    CHROME_ACCENT_RULE,
    CHROME_FOOTER,
    PPTXRenderer,
    is_chrome_shape,
)

BRAND = "#00704A"


def _slide(layout_id, title, body=None):
    contents = [{
        "type": "text", "placeholderIdx": 0,
        "runs": [{"text": title, "bold": True}],
    }]
    if body:
        contents.append({
            "type": "text", "placeholderIdx": 1,
            "runs": [{"text": t, "bullet": True, "level": 0} for t in body],
        })
    return {"slideIndex": 0, "layoutId": layout_id, "layoutName": layout_id,
            "contents": contents}


@pytest.fixture
def deck():
    """A cover, a content slide and a closing slide, rendered in the brand."""
    slides = [
        _slide("title", "Cover", ["A subtitle line"]),
        _slide("content", "Where the business stands", ["First point", "Second point"]),
        _slide("conclusion", "What we recommend", ["Do the thing"]),
    ]
    renderer = PPTXRenderer(template_path=None, color_theme=BRAND)
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        output_path = f.name
    try:
        renderer.render(slides=slides, metadata={"title": "Deck Title"},
                        output_path=output_path)
        yield Presentation(output_path)
    finally:
        os.unlink(output_path)


def _background(slide) -> str:
    return colors.to_hex(tuple(slide.background.fill.fore_color.rgb))


def _chrome(slide, name):
    return [s for s in slide.shapes if (s.name or "") == name]


def _text_of(slide, name):
    return " ".join(s.text_frame.text for s in _chrome(slide, name) if s.has_text_frame)


class TestLightDarkRhythm:
    def test_cover_and_closing_are_dark_and_content_is_light(self, deck):
        """A deck of seven identical white pages is the thing being fixed."""
        cover, content, closing = deck.slides

        assert colors.is_dark(_background(cover))
        assert not colors.is_dark(_background(content))
        assert colors.is_dark(_background(closing))

    def test_text_on_a_dark_slide_stays_legible(self, deck):
        """Body text used to keep the master's black whatever was behind it."""
        cover = deck.slides[0]
        background = _background(cover)

        runs = [
            run
            for shape in cover.shapes
            if shape.has_text_frame and not is_chrome_shape(shape)
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
        ]
        assert runs

        for run in runs:
            rgb = run.font.color.rgb
            assert rgb is not None, "run inherited the master's colour on a dark slide"
            assert colors.contrast_ratio(colors.to_hex(tuple(rgb)), background) >= 4.5


class TestRepeatedMotif:
    def test_every_slide_carries_the_accent_rule(self, deck):
        for index, slide in enumerate(deck.slides):
            assert _chrome(slide, CHROME_ACCENT_RULE), f"slide {index + 1} has no motif"

    def test_the_rule_is_drawn_in_the_palette_accent(self, deck):
        palette = {
            c["name"]: c["value"]
            for c in get_default_design_system(color_theme=BRAND)["palette"]["colors"]
        }
        rule = _chrome(deck.slides[1], CHROME_ACCENT_RULE)[0]

        assert colors.to_hex(tuple(rule.fill.fore_color.rgb)) == palette["accent"].upper()

    def test_the_rule_carries_no_inherited_shadow(self, deck):
        """Autoshapes inherit the theme's effectRef, which LibreOffice drew as
        a grey shadow under a 4pt bar - a rendering artifact, not a motif."""
        from pptx.oxml.ns import qn

        rule = _chrome(deck.slides[1], CHROME_ACCENT_RULE)[0]
        assert rule._element.find(qn("p:style")) is None


class TestFooter:
    def test_content_slides_carry_the_deck_title_and_a_slide_number(self, deck):
        text = _text_of(deck.slides[1], CHROME_FOOTER)

        assert "Deck Title" in text
        assert "2" in text

    def test_the_cover_carries_no_footer(self, deck):
        """No real deck numbers its cover or repeats its own title under it."""
        assert _chrome(deck.slides[0], CHROME_FOOTER) == []


class TestCatalogGeometryIsApplied:
    def test_the_title_lands_on_the_catalog_rectangle(self, deck):
        """The model plans against the catalog; the slide must use it too.

        Deferring to the stock layout's own geometry meant the model designed
        for one rectangle and python-pptx's rescaled 4:3 Office template
        supplied another.
        """
        catalog = get_default_design_system()
        content = next(l for l in catalog["layouts"] if l["layoutId"] == "content")
        spec = next(p for p in content["placeholders"] if p["type"] == "title")

        title = next(s for s in deck.slides[1].shapes
                     if s.has_text_frame
                     and s.text_frame.text == "Where the business stands")

        assert title.left == Emu(spec["xEmu"])
        assert title.top == Emu(spec["yEmu"])
        assert title.width == Emu(spec["cxEmu"])

    def test_nothing_is_drawn_below_the_footer_band(self, deck):
        """Content running under the footer is the deck's own chrome colliding
        with its content, which no amount of rewording by the model can fix."""
        catalog = get_default_design_system()
        footer_top = catalog["spacing"]["footer"]["yEmu"]

        for index, slide in enumerate(deck.slides):
            for shape in slide.shapes:
                if is_chrome_shape(shape) or shape.top is None or shape.height is None:
                    continue
                assert shape.top + shape.height <= footer_top, (
                    f"slide {index + 1}: {shape.name} runs into the footer band"
                )


class TestInheritedStylingIsNeutralised:
    def test_a_closing_title_is_not_shouted(self, deck):
        """The stock "Section Header" layout declares cap="all" on its title,
        so "What we recommend" came back as "WHAT WE RECOMMEND"."""
        from pptx.oxml.ns import qn

        title = next(s for s in deck.slides[2].shapes
                     if s.has_text_frame and s.text_frame.text == "What we recommend")
        run = title.text_frame.paragraphs[0].runs[0]

        assert run._r.find(qn("a:rPr")).get("cap") == "none"
