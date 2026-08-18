"""
Composing a page the template never contained.

The region grammar offers what the template's own pages do, to be copied
verbatim. That is right when one of them fits and leaves nothing when none
does, and "leave regions out" caps the slide at one exhibit.

What these pin down is that the geometry stays in code. A model names a
composition and a proportion; nothing here lets it write coordinates, because
a model asked for coordinates produces a rail of 0.31 beside a panel of 0.71
and the overlap only surfaces as content drawn over content.
"""

import pytest

from app.agent.registry import ToolRegistry
from app.agent.tools.compose_regions import ComposeSlideRegionsTool
from app.renderer import compositions
from app.renderer.regions import RegionTree, validate_regions


def _ids(tree):
    return [region["id"] for region in tree]


class TestEveryCompositionResolves:
    @pytest.mark.parametrize("name", sorted(compositions.COMPOSITIONS))
    def test_it_produces_a_valid_region_tree(self, name):
        tree = compositions.compose(name)

        errors, _ = validate_regions(tree)
        assert errors == [], f"{name}: {errors}"

    @pytest.mark.parametrize("name", sorted(compositions.COMPOSITIONS))
    def test_every_part_is_bindable(self, name):
        """A region a block cannot name is dead geometry."""
        tree = compositions.compose(name)
        resolved = RegionTree.from_slide({"regions": tree}, 13.333, 7.5)

        assert resolved is not None
        for region_id in _ids(tree):
            assert resolved.rect(region_id) is not None

    @pytest.mark.parametrize("name", sorted(compositions.COMPOSITIONS))
    def test_a_footnote_band_is_always_reserved(self, name):
        """
        An exhibit page with no room for its source line is how a chart ends
        up making a claim with nothing standing behind it.
        """
        assert "footnote" in _ids(compositions.compose(name))

    @pytest.mark.parametrize("name", sorted(compositions.COMPOSITIONS))
    def test_the_title_band_is_left_clear(self, name):
        """
        A composed slide still puts its headline in the template's own title
        placeholder. A region under that placeholder is content over content.
        """
        for region in compositions.compose(name):
            if region["id"] == "rail":
                continue  # full-bleed by design: it is the page's edge
            assert region["y"] >= compositions.TITLE_BAND_H - 1e-9, region["id"]

    @pytest.mark.parametrize("name", sorted(compositions.COMPOSITIONS))
    def test_the_extremes_of_the_range_still_resolve(self, name):
        composition = compositions.COMPOSITIONS[name]
        for share in (composition.min_proportion, composition.max_proportion):
            tree = compositions.compose(name, share)
            errors, _ = validate_regions(tree)
            assert errors == [], f"{name} at {share}: {errors}"


class TestTheGeometryIsRefusedRatherThanFudged:
    def test_an_unknown_composition_lists_the_real_ones(self):
        with pytest.raises(compositions.CompositionError) as caught:
            compositions.compose("sunburst")

        assert "exhibitGrid" in str(caught.value)

    def test_a_proportion_outside_the_range_says_what_it_measures(self):
        with pytest.raises(compositions.CompositionError) as caught:
            compositions.compose("railAndPanel", 0.9)

        message = str(caught.value)
        assert "0.15" in message and "0.4" in message
        assert "rail" in message

    def test_too_few_bands_are_refused(self):
        with pytest.raises(compositions.CompositionError):
            compositions.compose("bandStack", parts=1)

    def test_too_many_bands_are_refused(self):
        with pytest.raises(compositions.CompositionError) as caught:
            compositions.compose("bandStack", parts=9)

        assert "hold an implication" in str(caught.value)

    def test_grounding_a_part_that_does_not_exist_names_the_real_parts(self):
        with pytest.raises(compositions.CompositionError) as caught:
            compositions.compose("split", grounds={"rail": "#12233F"})

        assert "left" in str(caught.value)


class TestWhatTheCompositionsAre:
    def test_a_split_divides_the_width(self):
        left, right, _footnote = compositions.compose("split", 0.6)

        assert left["id"] == "left" and right["id"] == "right"
        assert left["w"] > right["w"]

    def test_a_rail_is_full_bleed(self):
        """It carries its own ground and is the page's edge, not its body."""
        rail = compositions.compose("railAndPanel")[0]

        assert rail["x"] == 0.0 and rail["y"] == 0.0 and rail["h"] == 1.0

    def test_a_ground_reaches_the_region(self):
        tree = compositions.compose("railAndPanel", grounds={"rail": "#12233F"})

        assert tree[0]["ground"] == "#12233F"

    def test_an_exhibit_grid_has_four_quadrants(self):
        ids = _ids(compositions.compose("exhibitGrid"))

        assert ids[:4] == ["topLeft", "topRight", "bottomLeft", "bottomRight"]

    def test_bands_are_equal_and_stacked(self):
        tree = compositions.compose("bandStack", parts=4)
        bands = [r for r in tree if r["id"].startswith("band")]

        assert len(bands) == 4
        assert len({round(b["h"], 6) for b in bands}) == 1
        assert [b["y"] for b in bands] == sorted(b["y"] for b in bands)

    def test_the_dashboard_page_from_the_plan_resolves(self):
        """A figure strip over two exhibits - the board-pack page."""
        tree = compositions.compose("stripAndGrid", 0.24)

        assert _ids(tree) == ["strip", "gridLeft", "gridRight", "footnote"]
        errors, _ = validate_regions(tree)
        assert errors == []


class TestTheTool:
    pytestmark = pytest.mark.asyncio

    @pytest.mark.asyncio
    async def test_it_returns_the_ids_to_bind_against(self):
        result = await ComposeSlideRegionsTool.run(0, "stripAndGrid")

        assert result.success
        assert result.data["regionIds"] == ["strip", "gridLeft", "gridRight"]
        assert result.data["slideIndex"] == 0

    @pytest.mark.asyncio
    async def test_the_footnote_band_is_not_offered_for_binding(self):
        """It is reserved, not an exhibit slot."""
        result = await ComposeSlideRegionsTool.run(0, "split")

        assert "footnote" not in result.data["regionIds"]
        assert "footnote" in _ids(result.data["regions"])

    @pytest.mark.asyncio
    async def test_a_bad_proportion_comes_back_as_an_actionable_error(self):
        result = await ComposeSlideRegionsTool.run(0, "railAndPanel", proportion=0.95)

        assert not result.success
        assert "rail" in result.error

    @pytest.mark.asyncio
    async def test_a_string_proportion_is_accepted(self):
        """Tool-calling models routinely send "0.6" for a number field."""
        result = await ComposeSlideRegionsTool.run(0, "split", proportion="0.6")

        assert result.success

    @pytest.mark.asyncio
    async def test_grounds_sent_as_a_json_string_are_accepted(self):
        result = await ComposeSlideRegionsTool.run(
            0, "railAndPanel", grounds='{"rail": "#12233F"}'
        )

        assert result.success
        assert result.data["regions"][0]["ground"] == "#12233F"

    @pytest.mark.asyncio
    async def test_grounds_of_the_wrong_shape_are_refused(self):
        result = await ComposeSlideRegionsTool.run(0, "split", grounds="[1, 2]")

        assert not result.success

    @pytest.mark.asyncio
    async def test_it_is_registered_and_offered_while_a_slide_is_drafted(self):
        from app.agent.pipeline import SLIDE_TOOLS
        from app.agent.registry import _register_default_tools

        registry = ToolRegistry()
        _register_default_tools(registry)

        names = {tool.name for tool in registry.list_tools()}
        assert "compose_slide_regions" in names
        # Offered only while a slide is being drafted: composing a page is
        # not something the planner or the research phase has any use for.
        assert "compose_slide_regions" in SLIDE_TOOLS


class TestTheComposedTreeReachesTheSlide:
    def test_the_loop_attaches_it(self):
        from app.agent.loop import WorkingState

        state = WorkingState()
        state.prepared_regions[0] = compositions.compose("split")

        state.resolve_and_add_slide(
            {"slideIndex": 0, "layoutId": "content", "contents": []}
        )

        assert _ids(state.accumulated_slides[0]["regions"]) == [
            "left", "right", "footnote"
        ]

    def test_it_wins_over_a_tree_the_model_wrote_itself(self):
        """
        Two trees on one slide would leave every binding pointing at
        whichever survived. The composed one is the one the exhibits were
        actually bound against.
        """
        from app.agent.loop import WorkingState

        state = WorkingState()
        state.prepared_regions[0] = compositions.compose("split")

        state.resolve_and_add_slide({
            "slideIndex": 0, "layoutId": "content", "contents": [],
            "regions": [{"id": "copied", "kind": "panel",
                         "x": 0, "y": 0, "w": 1, "h": 1}],
        })

        assert "copied" not in _ids(state.accumulated_slides[0]["regions"])

    def test_an_uncomposed_slide_declares_no_regions(self):
        """The region-free path is untouched."""
        from app.agent.loop import WorkingState

        state = WorkingState()
        state.resolve_and_add_slide(
            {"slideIndex": 0, "layoutId": "content", "contents": []}
        )

        assert "regions" not in state.accumulated_slides[0]


class TestAComposedPageRenders:
    """
    The end the whole workstream exists for: two exhibits, two places, one
    page, in a real rendered file.

    Geometry that validates and then does not render is geometry that lies,
    so this goes all the way to the PPTX rather than stopping at the tree.
    """

    @staticmethod
    def _render(slides, work_dir):
        import os

        from app.renderer.pptx_renderer import PPTXRenderer

        path = os.path.join(work_dir, "composed.pptx")
        PPTXRenderer(template_path=None, color_theme="professional").render(
            slides=slides, metadata={"title": "Composed"}, output_path=path,
        )
        return path

    @staticmethod
    def _chart(region_id, slot, label):
        return {
            "type": "chart", "regionId": region_id, "slot": slot,
            "chart": {
                "chartType": "bar",
                "categories": ["2023", "2024", "2025"],
                "series": [{"name": label, "values": [1.0, 2.0, 3.0]}],
            },
        }

    def _composed_slide(self):
        return {
            "slideIndex": 0,
            "layoutId": "content",
            "layoutName": "Content",
            "regions": compositions.compose("split", 0.5),
            "contents": [
                {"type": "text", "placeholderIdx": 0,
                 "runs": [{"text": "Two exhibits, one page", "bold": True}]},
                self._chart("left", "bridge", "Revenue ($B)"),
                self._chart("right", "mix", "Share (%)"),
            ],
        }

    def test_both_exhibits_reach_the_file(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as work_dir:
            path = self._render([self._composed_slide()], work_dir)

            with zipfile.ZipFile(path) as archive:
                media = [n for n in archive.namelist()
                         if n.startswith("ppt/media/")]

        # One rendered chart image per exhibit. Before the slot key was
        # widened this page carried one.
        assert len(media) >= 2, media

    def test_the_two_exhibits_are_drawn_in_different_places(self):
        import tempfile

        from pptx import Presentation

        with tempfile.TemporaryDirectory() as work_dir:
            path = self._render([self._composed_slide()], work_dir)
            slide = Presentation(path).slides[0]

            pictures = [s for s in slide.shapes if s.shape_type == 13]

        assert len(pictures) >= 2
        lefts = sorted(p.left for p in pictures[:2])
        assert lefts[1] > lefts[0], "both exhibits landed in the same column"

    def test_the_page_still_validates_as_slide_json(self):
        import asyncio
        import json

        from app.agent.tools.validate_slide_json import ValidateSlideJsonTool

        result = asyncio.run(
            ValidateSlideJsonTool.run(json.dumps(self._composed_slide()))
        )

        assert result.success, result.error
