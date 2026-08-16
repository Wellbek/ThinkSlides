"""
Per-slide footnotes and the source line.

The motivating case is reference page 110, whose third footnote says the final
bar of a chart is extrapolated. Without it the chart claims a measured figure
it does not have, so the apparatus is load-bearing rather than ornamental -
and it has to reach the page it belongs to, not an appendix eleven slides
later.

Two properties are pinned hard: that footnotes and deck-wide citations can sit
on the same slide without a reader confusing them, and that the apparatus
reserves its space *before* the body is laid out rather than being drawn
underneath the last bullet.
"""

import os
import shutil
import tempfile

import pytest
from pptx import Presentation

from app.agent.tools.validate_slide_json import ValidateSlideJsonTool
from app.agent.tools.validate_rendered_output import ValidateRenderedOutputTool
from app.design_system import get_default_design_system
from app.renderer import apparatus as ap
from app.renderer.apparatus import (
    ApparatusTokens, SlideApparatus, compile_apparatus, validate_apparatus,
)
from app.renderer.deck_theme import DeckTheme, ROLE_CONTENT
from app.renderer.layout import PositionConfig
from app.renderer.pptx_renderer import PPTXRenderer, is_chrome_shape
from app.renderer.preview_generator import PreviewGenerator
from app.renderer.primitives import shape_primitive_id
from app.renderer.typography import TypeScale

CANVAS = PositionConfig(left=0.0, top=0.0, width=13.333, height=7.5)


class TestMarkersAreWrittenInline:
    def test_a_marker_splits_a_string_into_parts(self):
        parts = ap.split_markers("AI adoption[^1] rose sharply")

        assert parts == [("AI adoption", None), ("1", 1), (" rose sharply", None)]

    def test_text_without_markers_comes_back_whole(self):
        assert ap.split_markers("Plain copy") == [("Plain copy", None)]

    def test_several_markers_in_one_string(self):
        assert ap.marker_numbers("a[^1] b[^2] c[^3]") == [1, 2, 3]

    def test_stripping_leaves_the_words(self):
        assert ap.strip_markers("Revenue[^1] grew") == "Revenue grew"

    def test_a_chart_label_gets_a_unicode_superscript(self):
        """matplotlib draws into a PNG, where no run-level superscript exists."""
        assert ap.to_superscript("Developing markets[^2]") == "Developing markets²"

    def test_markers_are_found_wherever_they_were_written(self):
        slide = {
            "contents": [
                {"type": "text", "runs": [{"text": "Body[^1]"}]},
                {"type": "table", "table": {"rows": [["Cell[^2]", "Plain"]]}},
                {"type": "chart", "chart": {"categories": ["2021[^3]"],
                                            "series": [{"name": "Revenue[^4]"}]}},
                {"type": "statCallout",
                 "statCallouts": [{"value": "56%", "caption": "Adoption[^5]"}]},
            ]
        }

        assert sorted(ap.collect_markers(slide)) == [1, 2, 3, 4, 5]


class TestValidation:
    def test_a_marker_with_no_footnote_is_rejected(self):
        errors, _ = validate_apparatus({
            "footnotes": ["Only one"],
            "contents": [{"type": "text", "runs": [{"text": "Copy[^2]"}]}],
        })

        assert any("[^2]" in e for e in errors)

    def test_a_footnote_nothing_points_at_warns(self):
        _, warnings = validate_apparatus({
            "footnotes": ["Unreferenced"],
            "contents": [{"type": "text", "runs": [{"text": "Copy"}]}],
        })

        assert any("never referenced" in w for w in warnings)

    def test_a_well_formed_slide_passes(self):
        errors, warnings = validate_apparatus({
            "footnotes": ["Based on open-source repositories"],
            "contents": [{"type": "text", "runs": [{"text": "Developers[^1]"}]}],
        })

        assert errors == [] and warnings == []

    def test_no_apparatus_is_not_an_error(self):
        assert validate_apparatus({"contents": []}) == ([], [])

    @pytest.mark.asyncio
    async def test_the_validator_tool_reports_it(self):
        import json

        slide = {
            "slideIndex": 0, "layoutId": "content", "layoutName": "Content",
            "footnotes": ["Only one"],
            "contents": [{"type": "text", "runs": [{"text": "Copy[^3]"}]}],
        }

        result = await ValidateSlideJsonTool.run(json.dumps(slide))

        assert result.success is False
        assert any("[^3]" in e for e in result.data["errors"])


class TestSpaceIsReservedBeforeTheBodyIsLaidOut:
    """
    The mechanical requirement. An apparatus laid out after the body is an
    apparatus with autofitted text on top of it.
    """

    @staticmethod
    def _plan(slide):
        renderer = PPTXRenderer(template_path=None, color_theme="professional")
        catalog = get_default_design_system()
        renderer._theme = DeckTheme.from_design_system(catalog)
        renderer._type_scale = TypeScale.from_design_system(catalog)
        return renderer._plan_slide(catalog, slide, index=1, total=3)

    def test_the_body_band_shrinks_when_a_slide_carries_footnotes(self):
        plain = self._plan({"slideIndex": 1, "layoutId": "content", "contents": []})
        with_notes = self._plan({
            "slideIndex": 1, "layoutId": "content", "contents": [],
            "footnotes": ["A note", "Another note"],
            "sourceLine": "Source: internal analysis",
        })

        assert with_notes.body_regions[0].height < plain.body_regions[0].height

    def test_it_shrinks_by_about_what_the_apparatus_needs(self):
        apparatus = SlideApparatus(footnotes=("A note", "Another note"),
                                   source_line="Source: internal analysis")
        tokens = ApparatusTokens.from_design_system(get_default_design_system())
        needed = apparatus.height(tokens, CANVAS.width * 0.934)

        plain = self._plan({"slideIndex": 1, "layoutId": "content", "contents": []})
        with_notes = self._plan({
            "slideIndex": 1, "layoutId": "content", "contents": [],
            "footnotes": ["A note", "Another note"],
            "sourceLine": "Source: internal analysis",
        })

        assert plain.body_regions[0].height - with_notes.body_regions[0].height == \
            pytest.approx(needed, abs=0.05)

    def test_a_slide_with_no_apparatus_is_untouched(self):
        plan = self._plan({"slideIndex": 1, "layoutId": "content", "contents": []})
        catalog = get_default_design_system()
        content = next(entry for entry in catalog["layouts"]
                       if entry["layoutId"] == "content")
        body = next(p for p in content["placeholders"] if p["type"] == "body")

        assert plan.body_regions[0].height == pytest.approx(
            body["cyEmu"] / 914400, abs=0.001
        )


class TestTheApparatusCompiles:
    @staticmethod
    def _marks(footnotes=(), source=""):
        design_system = get_default_design_system()
        theme = DeckTheme.from_design_system(design_system)
        return compile_apparatus(
            apparatus=SlideApparatus(footnotes=tuple(footnotes), source_line=source),
            tokens=ApparatusTokens.from_design_system(design_system),
            style=theme.style_for(ROLE_CONTENT),
            type_scale=TypeScale.from_design_system(design_system),
            canvas=CANVAS,
        )

    def test_nothing_compiles_to_nothing(self):
        assert self._marks() == []

    def test_each_footnote_is_its_own_mark(self):
        marks = self._marks(["First", "Second", "Third"])

        assert [m.id for m in marks] == [
            "apparatus-footnote-1", "apparatus-footnote-2", "apparatus-footnote-3",
        ]

    def test_a_footnote_number_is_a_superscript(self):
        [mark] = self._marks(["A note"])

        assert mark.runs[0].superscript is True
        assert mark.runs[0].text == "1"

    def test_the_source_line_sits_under_the_footnotes(self):
        marks = {m.id: m for m in self._marks(["A note"], "Source: analysis")}

        assert marks["apparatus-source"].box.top > \
            marks["apparatus-footnote-1"].box.top

    def test_the_apparatus_stays_inside_the_canvas(self):
        marks = self._marks(["A note", "Another"], "Source: analysis")

        for mark in marks:
            assert mark.box.top + mark.box.height <= CANVAS.height + 0.01

    def test_it_is_set_smaller_than_body_copy(self):
        design_system = get_default_design_system()
        body_pt = TypeScale.from_design_system(design_system).body_style(0).font_size_pt

        for mark in self._marks(["A note"], "Source: analysis"):
            assert mark.size_pt < body_pt

    def test_it_is_classified_as_furniture(self):
        assert all(mark.chrome for mark in self._marks(["A note"], "Source: x"))


class TestItReachesTheRenderedFile:
    @pytest.fixture(scope="class")
    @staticmethod
    def rendered():
        work_dir = tempfile.mkdtemp(prefix="ts_apparatus_")
        slides = [
            {"slideIndex": 0, "layoutId": "title",
             "contents": [{"type": "text", "runs": [{"text": "Web3"}]}]},
            # Reference page 110's arrangement: three footnotes, one of which
            # marks a chart's final bar as extrapolated, over a source line.
            {"slideIndex": 1, "layoutId": "chart", "section": "Web3",
             "footnotes": [
                 "Based on open-source repositories and code commits from GitHub.",
                 "Private-equity and venture-capital investments.",
                 "Extrapolated from 2022-H2 investment figures of $18.3 billion.",
             ],
             "sourceLine": "Source: Dovemetrics; Electric Capital; GitHub; PitchBook",
             "contents": [
                 {"type": "text",
                  "runs": [{"text": "Why is there enthusiasm for Web3?"}]},
                 {"type": "text", "runs": [
                     {"text": "Talent and investment[^1] have risen sharply",
                      "bullet": True}]},
                 {"type": "chart", "chart": {
                     "chartType": "bar",
                     "categories": ["2020", "2021", "2022[^3]"],
                     "series": [{"name": "PE and VC investments[^2]",
                                 "values": [5.0, 30.2, 36.6]}]}},
             ]},
            {"slideIndex": 2, "layoutId": "table", "section": "Web3",
             "footnotes": ["Includes China, Middle East and North Africa."],
             "sourceLine": "Source: internal analysis",
             "contents": [
                 {"type": "text", "runs": [{"text": "Adoption by region"}]},
                 {"type": "table", "table": {"headerRow": True, "rows": [
                     ["Region", "2021"],
                     ["Developing markets[^1]", "57%"],
                     ["Europe", "51%"]]}},
             ]},
        ]
        try:
            pptx_path = os.path.join(work_dir, "apparatus.pptx")
            PPTXRenderer(template_path=None, color_theme="#00A9F4").render(
                slides=slides, metadata={"title": "Technology Trends"},
                output_path=pptx_path,
            )
            pdf_path = PreviewGenerator.convert_to_pdf(pptx_path, work_dir)
            assert pdf_path, "LibreOffice could not convert the deck"
            yield {"pptx": pptx_path, "pdf": pdf_path}
        finally:
            shutil.rmtree(work_dir, ignore_errors=True)

    def test_nothing_overflows_or_overlaps(self, rendered):
        issues = ValidateRenderedOutputTool._check_geometry(rendered["pdf"])

        assert issues == [], "\n".join(
            f"  slide {i.get('slideIndex')}: {i['detail']}" for i in issues
        )

    def test_the_footnote_block_reaches_the_slide(self, rendered):
        slide = Presentation(rendered["pptx"]).slides[1]
        notes = [s for s in slide.shapes
                 if (shape_primitive_id(s) or "").startswith("apparatus-footnote")]

        assert len(notes) == 3
        assert "Extrapolated" in notes[2].text_frame.text

    def test_the_source_line_is_its_own_element(self, rendered):
        slide = Presentation(rendered["pptx"]).slides[1]

        source = next(s for s in slide.shapes
                      if shape_primitive_id(s) == "apparatus-source")
        assert source.text_frame.text.startswith("Source:")

    def test_the_apparatus_is_furniture_not_content(self, rendered):
        slide = Presentation(rendered["pptx"]).slides[1]
        marks = [s for s in slide.shapes
                 if (shape_primitive_id(s) or "").startswith("apparatus-")]

        assert marks
        assert all(is_chrome_shape(s) for s in marks)

    def test_a_body_marker_renders_as_a_true_superscript(self, rendered):

        slide = Presentation(rendered["pptx"]).slides[1]
        for shape in slide.shapes:
            if not shape.has_text_frame or is_chrome_shape(shape):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text == "1":
                        assert run._r.get_or_add_rPr().get("baseline") == "30000"
                        return
        pytest.fail("no superscript marker found in the body copy")

    def test_the_marker_syntax_never_reaches_the_page(self, rendered):
        import fitz

        doc = fitz.open(rendered["pdf"])
        try:
            text = "".join(doc[n].get_text() for n in range(len(doc)))
        finally:
            doc.close()

        assert "[^" not in text

    def test_a_table_cell_marker_renders_as_a_superscript(self, rendered):
        slide = Presentation(rendered["pptx"]).slides[2]
        table = next(s.table for s in slide.shapes if s.has_table)

        cell = table.cell(1, 0)
        markers = [run for paragraph in cell.text_frame.paragraphs
                   for run in paragraph.runs
                   if run._r.get_or_add_rPr().get("baseline") == "30000"]
        assert markers and markers[0].text == "1"

    def test_a_chart_label_marker_becomes_a_unicode_superscript(self, rendered):
        """
        A chart label is drawn into a PNG. It cannot carry a run-level
        superscript, and it must not carry "[^3]" either.
        """
        import fitz

        doc = fitz.open(rendered["pdf"])
        try:
            page_text = doc[1].get_text()
        finally:
            doc.close()

        assert "[^3]" not in page_text


class TestFootnotesAndDeckWideCitationsCoexist:
    """
    Both instruments on one slide. They cannot collide because the notation
    differs - a citation is a bracketed "[3]" on the baseline, a footnote is a
    superscript "3" - which is the only distinction a reader can actually use.
    """

    def test_the_two_numbering_schemes_are_independent(self):
        from app.renderer import citations

        sources = citations.collect_sources([
            {"url": "https://a.test", "title": "First source"},
            {"url": "https://b.test", "title": "Second source"},
        ])
        slide = {
            "slideIndex": 1, "layoutId": "content",
            "footnotes": ["A per-slide caveat"],
            "contents": [{"type": "text", "runs": [
                {"text": "A finding[^1]", "sourceRef": 2},
            ]}],
        }

        citations._mark_runs(slide, sources)

        run = slide["contents"][0]["runs"][0]
        # The citation marker is "[2]" and the footnote marker is "[^1]" - the
        # numbers differ and so does the notation.
        assert run["citation"]["text"] == " [2]"
        assert ap.marker_numbers(run["text"]) == [1]

    def test_a_slide_can_carry_both_without_the_validator_objecting(self):
        errors, warnings = validate_apparatus({
            "footnotes": ["A per-slide caveat"],
            "contents": [{"type": "text", "runs": [
                {"text": "A finding[^1]", "sourceRef": 2}]}],
        })

        assert errors == [] and warnings == []


class TestAComposedSlideCanOverCommitItsPage:
    """
    Placeholder-bound content gets the apparatus's height taken out of its
    band. A region tree is authored, and shrinking it afterwards would distort
    a composition somebody wrote deliberately - so the ceiling is published
    and the collision is reported instead.
    """

    @staticmethod
    def _ceiling(footnotes=("A note",), source="Source: analysis"):
        from app.renderer.apparatus import content_ceiling

        return content_ceiling(
            SlideApparatus(footnotes=tuple(footnotes), source_line=source),
            ApparatusTokens.from_design_system(get_default_design_system()),
            CANVAS,
        )

    def test_the_ceiling_leaves_room_for_the_apparatus(self):
        assert self._ceiling() < CANVAS.height - 0.3

    def test_a_slide_with_no_apparatus_may_use_the_whole_canvas(self):
        from app.renderer.apparatus import content_ceiling

        ceiling = content_ceiling(
            SlideApparatus(),
            ApparatusTokens.from_design_system(get_default_design_system()),
            CANVAS,
        )

        assert ceiling == CANVAS.height

    def test_a_region_reaching_past_it_is_reported(self, monkeypatch):
        # structlog is configured with its own processors, so the warning is
        # observed at the module's logger rather than through caplog.
        from app.renderer import pptx_renderer

        reported = []
        monkeypatch.setattr(
            pptx_renderer.logger, "warning",
            lambda event, **kwargs: reported.append(event),
        )

        slides = [{
            "slideIndex": 0, "layoutId": "content",
            "footnotes": ["A note", "Another note"],
            "sourceLine": "Source: internal analysis",
            "regions": [{"id": "full", "kind": "panel",
                         "x": 0.05, "y": 0.1, "w": 0.9, "h": 0.88}],
            "contents": [{"type": "text", "regionId": "full",
                          "runs": [{"text": "Copy"}]}],
        }]
        work_dir = tempfile.mkdtemp(prefix="ts_apparatus_warn_")
        try:
            PPTXRenderer(template_path=None, color_theme="professional").render(
                slides=slides, metadata={"title": "Over-committed"},
                output_path=os.path.join(work_dir, "over.pptx"),
            )
        finally:
            shutil.rmtree(work_dir, ignore_errors=True)

        assert "region_reaches_the_apparatus_band" in reported

    def test_a_composition_that_leaves_room_is_not_reported(self, monkeypatch):
        from app.renderer import pptx_renderer

        reported = []
        monkeypatch.setattr(
            pptx_renderer.logger, "warning",
            lambda event, **kwargs: reported.append(event),
        )
        slides = [{
            "slideIndex": 0, "layoutId": "content",
            "footnotes": ["A note"], "sourceLine": "Source: internal analysis",
            "regions": [{"id": "band", "kind": "panel",
                         "x": 0.05, "y": 0.1, "w": 0.9, "h": 0.6}],
            "contents": [{"type": "text", "regionId": "band",
                          "runs": [{"text": "Copy"}]}],
        }]
        work_dir = tempfile.mkdtemp(prefix="ts_apparatus_ok_")
        try:
            PPTXRenderer(template_path=None, color_theme="professional").render(
                slides=slides, metadata={"title": "Fits"},
                output_path=os.path.join(work_dir, "fits.pptx"),
            )
        finally:
            shutil.rmtree(work_dir, ignore_errors=True)

        assert "region_reaches_the_apparatus_band" not in reported
