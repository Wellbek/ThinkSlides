"""A message sent after a deck is finished amends that deck (#376).

The working state used to be deleted when a run completed, so a follow-up such
as "add a picture to every page" started a new deck from the intake, and the
edit tool - which needs a rendered deck - was never offered.
"""
import asyncio
import datetime
import io

import pytest
from pptx import Presentation

from app.agent import pipeline
from app.agent.loop import AgentLoop
from app.agent.tools import edit_deck_with_code as edit_tool


def _finished_state(path: str) -> dict:
    return {
        "has_called_finalize": True, "last_rendered_path": path,
        "brief": {"brief": "b"}, "scouted": True, "consulted_before_planning": True,
        "offered_directions": True, "reviewed_outline_with_user": True,
        "presentation_structure": {"totalSlides": 1, "slides": [{"title": "T"}]},
        "accumulated_slides": [{"slideIndex": 0}],
        "design_serial": 3, "rendered_serial": 3, "validated_serial": 3,
        "validation_clean": True, "validation_attempts": 2,
        "code_edit_log": [{"source": "x = 1", "intent": "earlier"}],
    }


class _Row:
    def __init__(self, state, updated_at):
        self.state, self.updated_at = state, updated_at


class _Message:
    def __init__(self, content, created_at):
        self.content, self.created_at = content, created_at


class _Result:
    def __init__(self, message):
        self.message = message

    def scalar_one_or_none(self):
        return self.message


class _Db:
    def __init__(self, row, message):
        self.row, self.message = row, message
        self.added = []

    async def get(self, model, key):
        return self.row

    async def execute(self, statement):
        return _Result(self.message)

    def add(self, obj):
        self.added.append(obj)

    async def commit(self):
        pass

    async def rollback(self):
        pass


@pytest.fixture
def deck(tmp_path):
    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    return str(path)


def _loop(db):
    loop = AgentLoop(db=db, mode=None)
    loop.questions_declined = True
    return loop


NOW = datetime.datetime(2026, 10, 2, 12, 0, tzinfo=datetime.timezone.utc)


def test_a_later_message_starts_an_amendment(deck):
    db = _Db(_Row(_finished_state(deck), NOW),
             _Message("Add a logo to every page", NOW + datetime.timedelta(minutes=5)))
    loop = _loop(db)
    assert asyncio.run(loop._load_deck_snapshot("sid")) is True
    ws = loop.working_state
    assert ws.amending and ws.amend_request == "Add a logo to every page"
    # What the last run concluded does not apply to the changed deck.
    assert not ws.validation_clean and ws.validation_attempts == 0
    assert ws.has_called_finalize and ws.render_is_current


def test_a_message_older_than_the_snapshot_is_not_an_amendment(deck):
    db = _Db(_Row(_finished_state(deck), NOW),
             _Message("make a deck", NOW - datetime.timedelta(minutes=5)))
    loop = _loop(db)
    assert asyncio.run(loop._load_deck_snapshot("sid")) is False
    assert not loop.working_state.amending


def test_a_deck_whose_file_is_gone_is_not_amended(tmp_path):
    db = _Db(_Row(_finished_state(str(tmp_path / "gone.pptx")), NOW),
             _Message("change it", NOW + datetime.timedelta(minutes=1)))
    loop = _loop(db)
    assert asyncio.run(loop._load_deck_snapshot("sid")) is False
    assert not loop.working_state.has_called_finalize


def test_an_amendment_goes_to_the_change_phase_then_validates(deck):
    db = _Db(_Row(_finished_state(deck), NOW),
             _Message("Add a logo", NOW + datetime.timedelta(minutes=1)))
    loop = _loop(db)
    asyncio.run(loop._load_deck_snapshot("sid"))
    phase = pipeline.next_phase(loop)
    assert phase.name == "change"
    assert phase.tools == frozenset(
        {"apply_deck_operation", "edit_deck_with_code", "revise_slides"})
    assert "Add a logo" in phase.directive(loop)
    loop.working_state.amend_done = True
    assert pipeline.next_phase(loop).name == "validate"


def test_an_amendment_does_not_research_again(deck):
    db = _Db(_Row(_finished_state(deck), NOW),
             _Message("Add a logo", NOW + datetime.timedelta(minutes=1)))
    loop = _loop(db)
    asyncio.run(loop._load_deck_snapshot("sid"))
    research = pipeline.PHASES_BY_NAME["research"]
    assert not research.applies(loop)


def test_a_snapshot_is_written_and_read_back(deck):
    db = _Db(None, None)
    loop = _loop(db)
    loop.working_state.has_called_finalize = True
    loop.working_state.last_rendered_path = deck
    loop.working_state.code_edit_log = [{"source": "pass", "intent": "i"}]
    asyncio.run(loop._save_deck_snapshot("sid"))
    saved = db.added[0].state
    assert saved["last_rendered_path"] == deck
    assert saved["code_edit_log"] == [{"source": "pass", "intent": "i"}]


ADD_BOX = '''
from pptx import Presentation
from pptx.util import Inches
p = Presentation("deck.pptx")
for s in p.slides:
    s.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(1), Inches(0.4)).text_frame.text = "LOGO"
p.save("deck.pptx")
'''


def _texts(path):
    return [shape.text_frame.text for slide in Presentation(path).slides
            for shape in slide.shapes if shape.has_text_frame]


def test_recorded_edits_are_put_back_after_a_render(deck):
    """A render rebuilds the file from the slides and drops the edit; replay
    puts it back, so a requested change survives a revision round."""
    assert "LOGO" not in _texts(deck)
    problems = edit_tool.replay(deck, [{"source": ADD_BOX, "intent": "logo"}])
    assert problems == []
    assert _texts(deck) == ["LOGO"]


def test_a_replayed_edit_that_fails_is_reported_and_the_rest_stand(deck):
    log = [{"source": "import os", "intent": "bad"},
           {"source": ADD_BOX, "intent": "logo"}]
    problems = edit_tool.replay(deck, log)
    assert len(problems) == 1 and problems[0].startswith("bad")
    assert _texts(deck) == ["LOGO"]


def test_previews_are_drawn_again_from_the_edited_file(deck, monkeypatch):
    from app.renderer import preview_generator

    seen = []

    def fake(self, path, session_id):
        seen.append((path, session_id))
        return ["/fresh/1.png"]

    monkeypatch.setattr(preview_generator.PreviewGenerator, "generate_previews", fake)
    loop = _loop(_Db(None, None))
    loop.working_state.last_rendered_path = deck
    loop.working_state.pending_slide_previews = [{"slideIndex": 0, "previewUrl": "/old/1.png"}]
    loop.working_state.pending_render_result = {"previewUrls": ["/old/1.png"]}
    asyncio.run(loop._refresh_previews("sid"))
    assert seen == [(deck, "sid")]
    assert loop.working_state.pending_slide_previews[0]["previewUrl"] == "/fresh/1.png"
    assert loop.working_state.pending_render_result["previewUrls"] == ["/fresh/1.png"]


def _png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 30, 30)).save(buffer, "PNG")
    return buffer.getvalue()


ADD_PICTURE = '''
from pptx import Presentation
from pptx.util import Inches
p = Presentation("deck.pptx")
for s in p.slides:
    s.shapes.add_picture("picture_1.png", Inches(0.2), Inches(0.2), Inches(0.5), Inches(0.5))
p.save("deck.pptx")
'''


def _pictures(path):
    return sum(1 for slide in Presentation(path).slides
               for shape in slide.shapes if shape.shape_type == 13)


def test_a_script_can_place_a_picture_the_user_supplied(deck):
    import base64

    from app.agent import run_context

    loop = _loop(_Db(None, None))
    loop._install_assets([{"data": "data:image/png;base64," + base64.b64encode(_png()).decode(),
                           "context": "the client's logo"}])
    assert list(run_context.get_assets()) == ["picture_1.png"]
    assert "picture_1.png (the client's logo)" in loop.asset_directive()
    assert edit_tool.replay(deck, [{"source": ADD_PICTURE, "intent": "logo"}]) == []
    assert _pictures(deck) == 1


def test_without_pictures_the_model_is_told_to_ask_rather_than_draw_one():
    """probe-overlay-picture (2026-10-04) drew a coffee cup from ellipses and
    shipped it as the user's picture. Without one, the reply asks for it."""
    loop = _loop(_Db(None, None))
    loop._install_assets([])
    directive = loop.asset_directive()
    assert "stand-in" not in directive
    assert "attached" in directive


def test_an_amendment_sets_the_deck_path_for_the_tools(deck):
    """The path used to be set only where the deck was rendered, so a request
    arriving in a new process found no deck to edit."""
    from app.agent import run_context

    token = run_context.set_rendered_path(None)
    try:
        db = _Db(_Row(_finished_state(deck), NOW),
                 _Message("Add a logo", NOW + datetime.timedelta(minutes=1)))
        loop = _loop(db)
        asyncio.run(loop._load_deck_snapshot("sid"))
        # What `_run` does once the snapshot has made this an amendment.
        if loop.working_state.amending:
            run_context.set_rendered_path(loop.working_state.last_rendered_path)
        assert run_context.get_rendered_path() == deck
    finally:
        run_context.reset_rendered_path(token)


def test_more_pictures_than_the_sandbox_takes_are_not_offered():
    import base64

    from app.agent.sandbox import runner

    loop = _loop(_Db(None, None))
    entry = {"data": "data:image/png;base64," + base64.b64encode(_png()).decode()}
    loop._install_assets([dict(entry) for _ in range(runner.MAX_ASSETS + 5)])
    assert len(loop.asset_notes) == runner.MAX_ASSETS


def test_an_amendment_with_nothing_staged_publishes_the_edited_pages(deck, monkeypatch):
    from app.renderer import preview_generator

    monkeypatch.setattr(preview_generator.PreviewGenerator, "generate_previews",
                        lambda self, path, session_id: ["/a/1.png", "/a/2.png"])
    loop = _loop(_Db(None, None))
    loop.working_state.last_rendered_path = deck
    loop.working_state.previews_published = True
    asyncio.run(loop._refresh_previews("sid"))
    staged = loop.working_state.pending_slide_previews
    assert [p["previewUrl"] for p in staged] == ["/a/1.png", "/a/2.png"]
    assert not loop.working_state.previews_published


def test_a_failed_restore_leaves_a_clean_state(deck, monkeypatch):
    db = _Db(_Row(_finished_state(deck), NOW),
             _Message("change it", NOW + datetime.timedelta(minutes=1)))
    loop = _loop(db)
    loop.working_state.accumulated_slides = [{"half": "restored"}]
    monkeypatch.setattr(loop, "_apply_state_dict", lambda state, sid: False)
    assert asyncio.run(loop._load_deck_snapshot("sid")) is False
    assert loop.working_state.accumulated_slides == []


def test_revise_slides_sends_slides_back_with_the_users_words():
    from app.agent.tools.revise_slides import ReviseSlidesTool

    result = asyncio.run(ReviseSlidesTool.run([1], "show build times as horizontal bars"))
    assert result.success and result.data["slideIndices"] == [1]
    assert not asyncio.run(ReviseSlidesTool.run([], "x")).success
    assert not asyncio.run(ReviseSlidesTool.run([1], "  ")).success


def test_revise_slides_is_offered_only_while_amending(deck):
    loop = _loop(_Db(None, None))
    assert "revise_slides" in loop._withheld_tools()
    loop.working_state.amending = True
    assert "revise_slides" not in loop._withheld_tools()


def test_a_script_that_only_looks_is_not_an_edit(deck, monkeypatch):
    from app.agent import run_context

    token = run_context.set_rendered_path(deck)
    try:
        look = 'from pptx import Presentation\nprint(len(Presentation("deck.pptx").slides))\n'
        result = asyncio.run(edit_tool.EditDeckWithCodeTool.run(look, "count", "none"))
    finally:
        run_context.reset_rendered_path(token)
    assert result.success and result.data["changed"] is False
    assert "The script printed:\n1" in result.data["summary"]


def test_the_change_and_amend_phases_can_be_declined_in_words():
    loop = _loop(_Db(None, None))
    for name, flag in (("change", "amend_done"), ("amend", "code_amendment_done")):
        phase = pipeline.PHASES_BY_NAME[name]
        assert phase.on_decline is not None
        assert getattr(loop.working_state, flag) is False
        phase.on_decline(loop)
        assert getattr(loop.working_state, flag) is True


def _handed_over_with_a_finding(deck):
    state = _finished_state(deck)
    state.update(validation_clean=False, validation_attempts=3, validation_issues=[
        {"slideIndex": 0, "type": "aesthetic", "detail": "the lower third is empty"}])
    db = _Db(_Row(state, NOW),
             _Message("Put a DRAFT watermark across every page",
                      NOW + datetime.timedelta(minutes=1)))
    loop = _loop(db)
    asyncio.run(loop._load_deck_snapshot("sid"))
    return loop


def test_what_the_deck_was_handed_over_with_is_not_reopened(deck):
    """probe-watermark (2026-10-04): a watermark request redrafted slide 2's
    copy over an empty lower third the deck already had when it was handed
    over. The user asked for a mark, not new slides."""
    loop = _handed_over_with_a_finding(deck)
    ws = loop.working_state
    assert ws.handover_findings == {(0, "aesthetic")}
    ws.amend_done = True
    ws.has_called_validate = True
    ws.validation_issues = [
        {"slideIndex": 0, "type": "aesthetic", "detail": "the lower third is empty"}]
    assert ws.new_since_handover(ws.validation_issues) == []
    # Settled: the walker does not go back to drafting, nor to the backstop.
    assert pipeline.next_phase(loop) is None
    assert ws.compute_slides_needing_revision(
        ws.new_since_handover(ws.validation_issues)) == []


def test_what_the_change_brought_goes_back_to_the_change_not_the_ladder(deck):
    """probe-hyperlink-terms (2026-10-04): a script made a link pure blue on
    the dark cover, and the ladder re-placed the cover and then redrafted
    its copy - neither of which can recolour a link."""
    loop = _handed_over_with_a_finding(deck)
    ws = loop.working_state
    ws.amend_edited = True
    blue = {"slideIndex": 0, "type": "overlap", "detail": "the link is blue on blue"}
    assert ws.caused_by_change([blue]) == [blue]
    assert ws.actionable([blue]) == []


def test_a_slide_the_change_redrafted_goes_through_the_ladder(deck):
    loop = _handed_over_with_a_finding(deck)
    ws = loop.working_state
    ws.amend_redrafted = {0}
    crowded = {"slideIndex": 0, "type": "overflow", "detail": "too long"}
    assert ws.caused_by_change([crowded]) == []
    assert ws.actionable([crowded]) == [crowded]


def test_the_baseline_survives_a_pause(deck):
    loop = _handed_over_with_a_finding(deck)
    saved = loop._state_dict()
    other = _loop(None)
    assert other._apply_state_dict(saved, "sid")
    assert other.working_state.handover_findings == {(0, "aesthetic")}


def _closing_prompt(loop, monkeypatch):
    seen = {}

    async def _ask(prompt, messages, phase_label=""):
        seen["prompt"] = prompt
        if False:
            yield None

    monkeypatch.setattr(loop, "_build_system_prompt", lambda *a, **k: "BASE")
    monkeypatch.setattr(loop, "_ask_model", _ask)

    async def drain():
        return [e async for e in loop._closing_message({}, [])]

    asyncio.run(drain())
    return seen["prompt"]


def test_a_change_that_was_not_made_is_said_to_be_not_made(deck, monkeypatch):
    """probe-progress-bar (2026-10-04): every script was refused, the change
    step ran out of turns, and the closing prompt still ended on 'make
    exactly that change' with nothing saying it had not been made."""
    loop = _handed_over_with_a_finding(deck)
    loop.working_state.last_code_attempt = {
        "source": "x", "reason": "This script draws a figure"}
    prompt = _closing_prompt(loop, monkeypatch)
    assert "the change was not made" in prompt
    assert "Put a DRAFT watermark across every page" in prompt
    assert "This script draws a figure" in prompt


def test_a_change_that_was_made_has_no_such_note(deck, monkeypatch):
    loop = _handed_over_with_a_finding(deck)
    loop.working_state.amend_done = True
    assert "the change was not made" not in _closing_prompt(loop, monkeypatch)


def test_a_change_declined_in_words_is_not_checked_again(deck, monkeypatch):
    """probe-rotate-exhibit (2026-10-04): the request named a chart the deck
    does not have, the change was declined in words, and the untouched deck
    was validated again - a vision call per page for nothing."""
    loop = _handed_over_with_a_finding(deck)
    pipeline.PHASES_BY_NAME["change"].on_decline(loop)
    ws = loop.working_state
    assert ws.amend_done
    assert ws.validation_issues == [
        {"slideIndex": 0, "type": "aesthetic", "detail": "the lower third is empty"}]
    assert pipeline.next_phase(loop) is None
    prompt = _closing_prompt(loop, monkeypatch)
    assert "Validation still found 1 layout issue" in prompt
    assert "already on the deck before" in prompt


def test_without_a_script_nothing_is_blamed_on_one(deck):
    """Review of #376's fixes: a change made only with revise_slides
    re-renders the deck, and a new finding on another slide came back to the
    change step as 'correct your script' when no script had run."""
    loop = _handed_over_with_a_finding(deck)
    ws = loop.working_state
    ws.amend_redrafted = {0}
    moved = {"slideIndex": 2, "type": "aesthetic", "detail": "placed differently"}
    assert ws.caused_by_change([moved]) == []
    assert ws.actionable([moved]) == [moved]


def test_a_decline_after_a_redraft_keeps_what_validation_found(deck):
    loop = _handed_over_with_a_finding(deck)
    ws = loop.working_state
    ws.amend_redrafted = {0}
    found = [{"slideIndex": 0, "type": "overflow", "detail": "too long"}]
    ws.validation_issues = list(found)
    pipeline.PHASES_BY_NAME["change"].on_decline(loop)
    assert ws.validation_issues == found
