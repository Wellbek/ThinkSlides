"""
Density and the type ramp: measured off rendered pages, carried as tokens.

An extracted design system used to record a palette, three font sizes and a
layout catalog. It had no field for *how much goes on a slide*, so a user could
upload a dense analyst deck as their template and we would learn its colours
and lose its density - and density is a large part of what makes such a deck
read the way it does. It also had no way to express a serif display over a sans
body, which is the pairing most editorial decks are set in.

Nothing here calls a model. Words per page, blocks per page and point sizes are
counts read off a rendered PDF, which is what makes them something to hold a
planner to rather than an impression to argue with. The tests that need a
rendered deck build one with python-pptx and put it through the same LibreOffice
conversion the preview path uses, so the suite stays hermetic - the reference
publication the epic is measured against is third-party and is never committed.
"""

import asyncio
import os
import shutil
import tempfile

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.design_system import get_default_design_system
from app.measure import density as den
from app.measure.page_metrics import PageMetrics
from app.measure.structure import PageStructure


def _page(number: int, words: int, ink: float = 0.4,
          ramp: dict | None = None, display=("Georgia", 25.0),
          body=("Arial", 14.0)) -> PageMetrics:
    """A measured page, without needing a PDF to measure."""
    return PageMetrics(
        page_number=number, width_pt=960.0, height_pt=540.0,
        ink_coverage=ink, word_count=words, block_count=4,
        type_ramp=ramp or {"25.0": 60, "14.0": 900},
        display_face=display[0], display_size_pt=display[1],
        body_face=body[0], body_size_pt=body[1],
    )


def _structure(blocks: int, rows: int = 3, columns: int = 3) -> PageStructure:
    return PageStructure(block_count=blocks, rows=rows, columns=columns,
                         band_rows=0)


# ──────────────────────────────────────────────────────────────────────
# What a deck puts on a page
# ──────────────────────────────────────────────────────────────────────

class TestADeckSaysHowMuchItPutsOnAPage:
    def test_the_budget_is_a_range_not_a_number(self):
        pages = [_page(i, words) for i, words in
                 enumerate([140, 180, 200, 210, 230, 290], start=1)]
        structures = [_structure(b) for b in (6, 8, 10, 11, 13, 15)]

        profile, _ = den.build_profiles(pages, structures)

        assert profile.pages_measured == 6
        assert profile.words_p25 < profile.words_median < profile.words_p75
        assert profile.blocks_p25 < profile.blocks_median < profile.blocks_p75

    def test_a_cover_is_not_evidence_about_content_density(self):
        """A cover carries a title and nothing else; averaging it in halves
        the budget every content page is then written against."""
        pages = [_page(1, 6, ink=0.05), _page(2, 200), _page(3, 210)]
        structures = [_structure(1), _structure(10), _structure(11)]

        profile, _ = den.build_profiles(pages, structures)

        assert profile.pages_measured == 2
        assert profile.words_median == 205.0

    def test_blocks_come_from_the_grid_reader_not_the_line_clusterer(self):
        """``structure.count_blocks`` knows a table row from a paragraph;
        ``page_metrics.block_count`` is the coarser fallback."""
        pages = [_page(1, 200), _page(2, 200)]

        with_structure, _ = den.build_profiles(pages, [_structure(15),
                                                       _structure(15)])
        without, _ = den.build_profiles(pages, [])

        assert with_structure.blocks_median == 15
        assert without.blocks_median == 4

    def test_words_per_block_says_what_a_filled_block_reads_like(self):
        pages = [_page(1, 200), _page(2, 200)]

        profile, _ = den.build_profiles(pages, [_structure(10), _structure(10)])

        assert profile.words_per_block_median == 20.0

    def test_the_grid_shape_is_carried_too(self):
        pages = [_page(1, 200), _page(2, 200), _page(3, 200)]
        structures = [_structure(15, rows=5, columns=3),
                      _structure(12, rows=4, columns=3),
                      _structure(9, rows=3, columns=3)]

        profile, _ = den.build_profiles(pages, structures)

        assert profile.grid_rows_median == 4
        assert profile.grid_columns_median == 3

    def test_nothing_measured_yields_an_empty_profile(self):
        profile, types = den.build_profiles([], [])

        assert profile.pages_measured == 0
        assert profile.describe() == ""
        assert types.ramp_pt == []

    def test_the_profile_round_trips_through_the_bundle(self):
        pages = [_page(i, 200) for i in range(1, 5)]
        profile, _ = den.build_profiles(pages, [_structure(10)] * 4)

        again = den.DensityProfile.from_dict(profile.to_dict())

        assert again == profile

    def test_a_bundle_written_before_density_existed_still_parses(self):
        assert den.DensityProfile.from_dict(None).words_median == 0.0
        assert den.DensityProfile.from_dict({}).pages_measured == 0


class TestTheBudgetIsStatedInASentence:
    def test_measured_pages_say_so(self):
        pages = [_page(i, w) for i, w in enumerate([180, 200, 230], start=1)]
        profile, _ = den.build_profiles(pages, [_structure(b) for b in (7, 10, 13)])

        sentence = profile.describe()

        assert "measured" in sentence
        assert "3 of them" in sentence
        assert "words" in sentence and "blocks" in sentence

    def test_house_defaults_do_not_claim_to_be_measurements(self):
        sentence = den.DensityProfile.from_dict(den.default_density()).describe()

        assert sentence
        assert "measured" not in sentence

    def test_a_design_system_with_no_density_briefs_nothing(self):
        assert den.density_brief({}) == ""
        assert den.density_brief(None) == ""

    def test_the_brief_reads_the_bundle(self):
        brief = den.density_brief(get_default_design_system())

        assert brief.startswith("Pages in this design carry")


# ──────────────────────────────────────────────────────────────────────
# The type ramp
# ──────────────────────────────────────────────────────────────────────

class TestTheTypeRampIsMeasuredNotAssumed:
    def test_every_size_carrying_copy_is_a_rung(self):
        ramp = {"25.0": 300, "18.0": 200, "14.0": 2000, "12.0": 400}
        _, types = den.build_profiles([_page(1, 200, ramp=ramp)], [])

        assert types.ramp_pt == [25.0, 18.0, 14.0, 12.0]

    def test_a_stray_superscript_is_not_a_step_in_the_scale(self):
        ramp = {"25.0": 300, "14.0": 3000, "5.0": 4}
        _, types = den.build_profiles([_page(1, 200, ramp=ramp)], [])

        assert 5.0 not in types.ramp_pt

    def test_sizes_within_a_hair_of_each_other_are_one_rung(self):
        ramp = {"14.0": 1000, "14.4": 900, "25.0": 300}
        _, types = den.build_profiles([_page(1, 200, ramp=ramp)], [])

        assert types.ramp_pt == [25.0, 14.0]

    def test_the_caption_rung_is_copy_not_apparatus(self):
        """The reference's 7pt footnotes are apparatus, sized by the
        apparatus tokens. Reading them as its caption size would set every
        caption on every slide at 7pt."""
        ramp = {"25.0": 300, "14.0": 3000, "12.0": 400, "7.0": 200}
        _, types = den.build_profiles([_page(1, 200, ramp=ramp)], [])

        assert types.caption_pt == 12.0

    def test_a_ramp_with_no_rung_below_the_body_still_has_a_caption(self):
        ramp = {"25.0": 300, "14.0": 3000}
        _, types = den.build_profiles([_page(1, 200, ramp=ramp)], [])

        assert types.caption_pt and types.caption_pt < types.body_pt

    def test_the_three_roles_are_filled_from_the_ramp(self):
        ramp = {"25.0": 300, "14.0": 3000, "12.0": 400}
        _, types = den.build_profiles([_page(1, 200, ramp=ramp)], [])

        typography = types.as_typography()

        assert typography["titleFont"]["sizePt"] == 25.0
        assert typography["bodyFont"]["sizePt"] == 14.0
        assert typography["captionFont"]["sizePt"] == 12.0
        assert typography["rampPt"] == [25.0, 14.0, 12.0]

    def test_a_serif_display_over_a_sans_body_is_recorded_as_such(self):
        _, types = den.build_profiles(
            [_page(1, 200, display=("Georgia", 25.0), body=("Arial", 14.0))], []
        )

        assert types.display_face == "Georgia"
        assert types.body_face == "Arial"
        assert types.display_is_serif is True

    def test_the_profile_round_trips_through_the_bundle(self):
        _, types = den.build_profiles([_page(1, 200)], [])

        assert den.TypeProfile.from_dict(types.to_dict()) == types


# ──────────────────────────────────────────────────────────────────────
# Font names and the chains behind them
# ──────────────────────────────────────────────────────────────────────

class TestAFaceNameSurvivesTheRenderHost:
    @pytest.mark.parametrize("raw,expected", [
        ("ABCDEF+Georgia-Bold", "Georgia"),
        ("LiberationSans", "Liberation Sans"),
        ("NotoSerif", "Noto Serif"),
        ("TimesNewRoman-Italic", "Times New Roman"),
        ("Arial", "Arial"),
        ("ArialMT", "Arial"),
        ("", "Arial"),
    ])
    def test_a_pdf_font_name_becomes_a_family_a_picker_can_match(self, raw, expected):
        assert den._clean_face(raw) == expected

    @pytest.mark.parametrize("face", ["Liberation Sans", "DejaVu Serif",
                                      "Noto Serif", "Nimbus Roman", "Carlito"])
    def test_the_containers_own_fonts_are_recognised_as_substitutes(self, face):
        assert den.looks_substituted(face)

    @pytest.mark.parametrize("face", ["Georgia", "Arial", "Calibri",
                                      "Helvetica Neue", ""])
    def test_a_real_typeface_is_not_mistaken_for_one(self, face):
        assert not den.looks_substituted(face)

    def test_a_name_read_off_the_page_beats_the_theme_that_disagrees(self):
        """A deck can override its theme font on every slide, and then the
        theme is a claim about nothing."""
        assert den.resolve_face("Calibri", "Georgia") == "Georgia"

    def test_the_theme_wins_when_the_page_only_shows_a_substitute(self):
        assert den.resolve_face("Calibri", "Liberation Sans") == "Calibri"

    def test_a_substitute_that_disagrees_about_serifs_still_wins(self):
        """Substitution never turns a serif into a sans, so a serif page
        under a sans theme means the theme is not what the deck uses."""
        assert den.resolve_face("Calibri", "Noto Serif") == "Noto Serif"

    def test_a_substitute_is_not_the_name_written_into_a_deck(self):
        """Naming this container's font hands the reader a deck naming a
        font their machine does not have either."""
        assert den.preferred_face("Noto Serif") == "Georgia"
        assert den.preferred_face("Liberation Sans") == "Arial"
        assert den.preferred_face("Georgia") == "Georgia"

    def test_a_chain_starts_with_the_face_and_ends_generic(self):
        chain = den.fallback_chain("Georgia")

        assert chain[0] == "Georgia"
        assert chain[-1] == "serif"
        assert len(chain) == len(set(chain))

    @pytest.mark.parametrize("face,generic", [
        ("Georgia", "serif"), ("Helvetica Neue", "sans-serif"),
        ("Consolas", "monospace"), ("Some Unknown Face", "sans-serif"),
    ])
    def test_a_chain_stays_inside_its_own_family(self, face, generic):
        assert den.fallback_chain(face)[-1] == generic


# ──────────────────────────────────────────────────────────────────────
# Measuring a real template
# ──────────────────────────────────────────────────────────────────────

def _build_template(path: str, *, display=("Georgia", 25), body=("Arial", 14),
                    slides: int = 3) -> str:
    """A small deck set in a named pairing, to be measured as a template."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        head = slide.shapes.add_textbox(Inches(0.8), Inches(0.6),
                                        Inches(11.5), Inches(1.1))
        run = head.text_frame.paragraphs[0].add_run()
        run.text = f"A headline that states a finding, number {index + 1}"
        run.font.name, run.font.size, run.font.bold = display[0], Pt(display[1]), True

        copy = slide.shapes.add_textbox(Inches(0.8), Inches(2.0),
                                        Inches(11.5), Inches(4.5))
        frame = copy.text_frame
        frame.word_wrap = True
        for line in range(6):
            paragraph = frame.paragraphs[0] if line == 0 else frame.add_paragraph()
            written = paragraph.add_run()
            written.text = (
                "Body copy carrying a complete sentence with a figure of "
                f"{12 + line} percent in it, long enough to read as prose."
            )
            written.font.name, written.font.size = body[0], Pt(body[1])
    prs.save(path)
    return path


@pytest.fixture(scope="module")
def work_dir():
    path = tempfile.mkdtemp(prefix="ts_density_")
    try:
        yield path
    finally:
        shutil.rmtree(path, ignore_errors=True)


@pytest.fixture(scope="module")
def serif_over_sans(work_dir):
    return _build_template(os.path.join(work_dir, "serif_over_sans.pptx"))


@pytest.fixture(scope="module")
def extracted(serif_over_sans):
    """The design system a real extraction run produces for that template."""
    from app.agent.tools.extract_template_design_system import (
        ExtractTemplateDesignSystemTool,
    )

    result = asyncio.run(
        ExtractTemplateDesignSystemTool.run(template_path=serif_over_sans)
    )
    assert result.success
    return result.data


class TestATemplateIsMeasuredNotAsked:
    def test_a_pptx_is_profiled_through_the_pages_it_renders_to(self, serif_over_sans):
        profile, types = den.profile_template(serif_over_sans)

        assert profile.pages_measured == 3
        assert profile.words_median > 0
        # The sizes are exact: a point size is not substituted at render time.
        assert types.display_pt == 25.0
        assert types.body_pt == 14.0

    def test_a_template_with_no_pages_is_not_given_an_invented_density(self, work_dir):
        empty = os.path.join(work_dir, "layouts_only.pptx")
        Presentation().save(empty)

        assert den.profile_template(empty) is None

    def test_only_a_sample_of_a_long_deck_is_measured(self):
        pages = den.sample_pages(180)

        assert len(pages) == den.MAX_PROFILE_PAGES
        assert pages[0] == 1 and pages[-1] == 180
        assert pages == sorted(set(pages))

    def test_a_short_deck_is_measured_whole(self):
        assert den.sample_pages(5) == [1, 2, 3, 4, 5]
        assert den.sample_pages(0) == []


class TestTheSerifOverSansPairingRoundTrips:
    """The acceptance criterion of #119: a template's pairing has to survive
    extraction and arrive at the renderer as two different faces."""

    def test_extraction_records_the_measured_density(self, extracted):
        assert extracted["density"]["pagesMeasured"] == 3
        assert extracted["density"]["wordsPerSlide"]["median"] > 0

    def test_extraction_records_the_measured_sizes(self, extracted):
        typography = extracted["typography"]

        assert typography["titleFont"]["sizePt"] == 25.0
        assert typography["bodyFont"]["sizePt"] == 14.0
        assert typography["rampPt"] == [25.0, 14.0]

    def test_the_display_is_a_serif_and_the_body_is_not(self, extracted):
        from app.measure.page_metrics import _is_serif

        typography = extracted["typography"]

        assert typography["displayIsSerif"] is True
        assert _is_serif(typography["titleFont"]["typeface"])
        assert not _is_serif(typography["bodyFont"]["typeface"])

    def test_every_face_carries_the_stack_to_try(self, extracted):
        for role in ("titleFont", "bodyFont", "captionFont"):
            chain = extracted["typography"][role]["fallbacks"]
            assert chain and chain[-1] in ("serif", "sans-serif", "monospace")

    def test_the_renderer_receives_the_pairing(self, extracted):
        from app.renderer.typography import TypeScale

        scale = TypeScale.from_design_system(extracted)

        assert scale.display_is_serif is True
        assert scale.title_font != scale.body_font
        assert scale.title_style().font_name == scale.title_font
        assert scale.body_style().font_name == scale.body_font
        assert scale.title_fallbacks and scale.body_fallbacks


class TestTheChainIsAppliedWhereWeRasterise:
    """A PPTX names a face and leaves substitution to whoever opens it. A
    chart is a PNG made here, so the chain has to be walked here or a serif
    deck gets sans chart labels."""

    def test_a_serif_deck_gets_serif_chart_type(self):
        from app.renderer.chart_renderer import _resolve_font

        chain = den.fallback_chain("Georgia")
        resolved = _resolve_font("Georgia", chain)

        assert resolved != "DejaVu Sans" or "DejaVu Sans" in chain
        assert den.family_of(resolved) == "serif"

    def test_the_chart_style_carries_the_decks_recorded_chain(self):
        from app.renderer.table_style import resolve_chart_style

        design_system = {
            "typography": {
                "bodyFont": {"typeface": "Georgia", "sizePt": 14,
                             "fallbacks": ["Georgia", "Liberation Serif", "serif"]},
            },
        }

        style = resolve_chart_style(design_system)

        assert style.font_name == "Georgia"
        assert style.font_fallbacks == ["Georgia", "Liberation Serif", "serif"]

    def test_a_design_system_with_no_chain_gets_one_from_the_face(self):
        from app.renderer.table_style import resolve_chart_style, resolve_table_style

        design_system = {"typography": {"bodyFont": {"typeface": "Georgia"}}}

        assert resolve_chart_style(design_system).font_fallbacks[-1] == "serif"
        assert resolve_table_style(design_system).font_fallbacks[-1] == "serif"


# ──────────────────────────────────────────────────────────────────────
# The tokens travel
# ──────────────────────────────────────────────────────────────────────

class TestTheBuiltInDesignSystemCarriesBoth:
    def test_density_is_a_token_like_any_other(self):
        density = get_default_design_system()["density"]

        assert density["wordsPerSlide"]["p25"] < density["wordsPerSlide"]["p75"]
        assert density["blocksPerSlide"]["median"] > 0

    def test_the_built_in_density_is_not_dressed_up_as_a_measurement(self):
        assert get_default_design_system()["density"]["pagesMeasured"] == 0

    def test_every_built_in_face_carries_a_chain(self):
        typography = get_default_design_system()["typography"]

        for role in ("titleFont", "bodyFont", "captionFont"):
            assert typography[role]["fallbacks"][0] == "Arial"
            assert typography[role]["fallbacks"][-1] == "sans-serif"

    def test_the_built_in_ramp_is_published(self):
        typography = get_default_design_system()["typography"]

        assert typography["rampPt"][0] == typography["titleFont"]["sizePt"]
        assert typography["displayIsSerif"] is False

    def test_the_renderer_still_gets_its_historic_defaults(self):
        from app.renderer.typography import TypeScale

        scale = TypeScale.from_design_system(get_default_design_system())

        assert (scale.title_font, scale.body_font) == ("Arial", "Arial")
        assert (scale.title_size_pt, scale.body_size_pt) == (36, 20)


class TestThePlannerAndDesignerReadTheBudget:
    def test_the_planner_states_the_budget_as_a_target(self):
        from app.agent.tools.plan_presentation_structure import _density_guidance

        guidance = _density_guidance(get_default_design_system())

        assert "Density target" in guidance
        assert den.density_brief(get_default_design_system()) in guidance

    def test_the_designer_states_the_same_budget(self):
        from app.agent.tools.design_slide_content import (
            _density_guidance as designer_guidance,
        )

        guidance = designer_guidance(get_default_design_system())

        assert den.density_brief(get_default_design_system()) in guidance

    def test_neither_invents_a_budget_when_none_was_recorded(self):
        from app.agent.tools.design_slide_content import (
            _density_guidance as designer_guidance,
        )
        from app.agent.tools.plan_presentation_structure import _density_guidance

        assert _density_guidance({}) == ""
        assert designer_guidance({}) == ""

    def test_a_measured_template_is_what_the_planner_is_held_to(self):
        pages = [_page(i, w) for i, w in enumerate([180, 208, 235], start=1)]
        profile, _ = den.build_profiles(pages, [_structure(b) for b in (7, 10, 13)])
        from app.agent.tools.plan_presentation_structure import _density_guidance

        guidance = _density_guidance({"density": profile.to_dict()})

        assert "measured" in guidance
        assert "208" in guidance
