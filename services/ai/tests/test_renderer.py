"""
Tests for the PPTX renderer.
"""

import pytest
import os
import pathlib
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from app.renderer.pptx_renderer import PPTXRenderer, NS, is_chrome_shape
from app.design_system import COLOR_THEMES
from app.renderer.default_template import _hex_to_rgbcolor


class TestPPTXRendererDefault:
    """Test rendering with the default theme (no template)."""

    @pytest.fixture
    def renderer(self):
        return PPTXRenderer(template_path=None)

    def test_render_empty_presentation(self, renderer):
        """Test rendering a presentation with no slides."""
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name

        try:
            result = renderer.render(
                slides=[],
                metadata={"title": "Test"},
                output_path=output_path,
            )

            assert os.path.exists(result)
            assert os.path.getsize(result) > 0
        finally:
            os.unlink(output_path)

    def test_render_single_slide(self, renderer):
        """Test rendering a single text slide."""
        slides = [
            {
                "slideIndex": 0,
                "layoutId": "1",
                "layoutName": "Default",
                "contents": [
                    {
                        "type": "text",
                        "placeholderIdx": 0,
                        "runs": [
                            {"text": "Hello World", "bold": True},
                        ],
                    }
                ],
            }
        ]

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name

        try:
            result = renderer.render(
                slides=slides,
                metadata={"title": "Test Presentation", "author": "Test Author"},
                output_path=output_path,
            )

            assert os.path.exists(result)
            assert os.path.getsize(result) > 0

            # Verify it's a valid PPTX by opening with python-pptx
            from pptx import Presentation
            prs = Presentation(result)
            assert len(prs.slides) == 1
        finally:
            os.unlink(output_path)

    def test_render_multiple_slides(self, renderer):
        """Test rendering multiple slides."""
        slides = [
            {
                "slideIndex": 0,
                "layoutId": "1",
                "layoutName": "Default",
                "contents": [
                    {
                        "type": "text",
                        "placeholderIdx": 0,
                        "runs": [{"text": "Title Slide", "bold": True}],
                    }
                ],
            },
            {
                "slideIndex": 1,
                "layoutId": "1",
                "layoutName": "Default",
                "contents": [
                    {
                        "type": "text",
                        "placeholderIdx": 0,
                        "runs": [
                            {"text": "Point 1", "bullet": True, "level": 0},
                            {"text": "Point 2", "bullet": True, "level": 0},
                        ],
                    }
                ],
            },
        ]

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name

        try:
            result = renderer.render(
                slides=slides,
                metadata={"title": "Multi-Slide Test"},
                output_path=output_path,
            )

            assert os.path.exists(result)
            from pptx import Presentation
            prs = Presentation(result)
            assert len(prs.slides) == 2
        finally:
            os.unlink(output_path)


def _write_minimal_rels(rels_path: str, extra_relationships: str = "") -> None:
    os.makedirs(os.path.dirname(rels_path), exist_ok=True)
    with open(rels_path, "w", encoding="utf-8") as f:
        f.write(
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
            '<Relationship Id="rId1" Type=".../slideMaster" Target="slideMasters/slideMaster1.xml"/>\n'
            f"{extra_relationships}"
            "</Relationships>"
        )


class TestPresentationRelsRidAllocation:
    """Regression tests: rId{i+2} was hardcoded, assuming rId1 is always the
    slideMaster and rId2..N+1 are free - any template whose
    presentation.xml.rels already used those IDs for something else (a
    theme, printerSettings, ...) got duplicate Id attributes, which is
    invalid OOXML and makes PowerPoint report the file as needing repair.
    """

    @pytest.fixture
    def renderer(self):
        return PPTXRenderer(template_path=None)

    def test_new_slide_rids_do_not_collide_with_existing_relationships(self, renderer, tmp_path):
        working_dir = str(tmp_path)
        rels_path = os.path.join(working_dir, "ppt", "_rels", "presentation.xml.rels")
        # A theme relationship already occupies rId2 - the exact ID the old
        # hardcoded scheme would have reassigned to the first new slide.
        _write_minimal_rels(
            rels_path,
            '<Relationship Id="rId2" Type=".../theme" Target="theme/theme1.xml"/>\n',
        )

        rids = renderer._update_presentation_rels(working_dir, slide_count=2)

        assert rids == ["rId3", "rId4"]

        tree = ET.parse(rels_path)
        ids = [rel.get("Id") for rel in tree.getroot()]
        assert len(ids) == len(set(ids)), f"duplicate relationship IDs: {ids}"
        assert "rId2" in ids  # the pre-existing theme relationship survives untouched

    def test_creates_sldidlst_when_the_template_omits_it_entirely(self, renderer, tmp_path):
        """Regression: our own generated default template (0 slides) omits
        <p:sldIdLst> rather than emitting it empty. The renderer used to
        bail out silently here, producing a 0-slide package that still
        reported success - the "downloaded pptx is empty" failure mode."""
        working_dir = str(tmp_path)
        rels_path = os.path.join(working_dir, "ppt", "_rels", "presentation.xml.rels")
        _write_minimal_rels(rels_path)
        pres_path = os.path.join(working_dir, "ppt", "presentation.xml")
        os.makedirs(os.path.dirname(pres_path), exist_ok=True)
        with open(pres_path, "w", encoding="utf-8") as f:
            f.write(
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
                '<p:sldSz cx="9144000" cy="6858000"/>'
                "</p:presentation>"
            )

        renderer._update_presentation_xml(working_dir, slide_count=2)

        tree = ET.parse(pres_path)
        root = tree.getroot()
        sldIdLst = root.find("p:sldIdLst", NS)
        assert sldIdLst is not None
        assert len(list(sldIdLst)) == 2
        # Schema order: sldIdLst must land between sldMasterIdLst and sldSz.
        tags = [el.tag.split("}")[-1] for el in root]
        assert tags.index("sldMasterIdLst") < tags.index("sldIdLst") < tags.index("sldSz")

    def test_presentation_xml_and_rels_agree_on_the_same_rids(self, renderer, tmp_path):
        working_dir = str(tmp_path)
        rels_path = os.path.join(working_dir, "ppt", "_rels", "presentation.xml.rels")
        _write_minimal_rels(
            rels_path,
            '<Relationship Id="rId2" Type=".../theme" Target="theme/theme1.xml"/>\n',
        )
        pres_path = os.path.join(working_dir, "ppt", "presentation.xml")
        os.makedirs(os.path.dirname(pres_path), exist_ok=True)
        with open(pres_path, "w", encoding="utf-8") as f:
            f.write(
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                "<p:sldIdLst/></p:presentation>"
            )

        renderer._update_presentation_xml(working_dir, slide_count=2)

        tree = ET.parse(pres_path)
        sldIdLst = tree.getroot().find(".//p:sldIdLst", NS)
        xml_rids = [sldId.get(f"{{{NS['r']}}}id") for sldId in sldIdLst]

        rels_tree = ET.parse(rels_path)
        rels_ids = {rel.get("Id") for rel in rels_tree.getroot()}

        assert xml_rids == ["rId3", "rId4"]
        assert set(xml_rids).issubset(rels_ids)


class TestCloneLayoutPreservesFormatting:
    """Regression: bullets beyond what the template placeholder itself had
    paragraphs for were created with no pPr/rPr at all, so they silently
    lost the template's bullet marker, indent, and typography - making
    multi-point slides look thinner/plainer than the template's design.
    """

    @pytest.fixture
    def renderer(self):
        return PPTXRenderer(template_path=None)

    @pytest.fixture
    def default_template_layout2(self, tmp_path):
        """The bundled default template's slideLayout2.xml ("Title and
        Content") ships 5 pre-existing body paragraphs (levels 0-4) under
        placeholder idx=1 - real template content, not a synthetic fixture."""
        from app.renderer.default_template import get_default_template_path

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        working_dir = tmp_path / "extracted"
        with zipfile.ZipFile(template_path) as zf:
            zf.extractall(working_dir)
        return str(working_dir / "ppt" / "slideLayouts" / "slideLayout2.xml")

    def test_extra_bullets_inherit_the_template_s_paragraph_and_run_formatting(
        self, renderer, default_template_layout2, tmp_path
    ):
        # 7 runs against a placeholder whose layout only has 5 paragraphs -
        # runs 6 and 7 must be freshly created and cloned-formatted.
        runs = [{"text": f"Point {i}", "bullet": True, "level": 0} for i in range(7)]
        slide_data = {
            "slideIndex": 0,
            "layoutId": "2",
            "contents": [{"type": "text", "placeholderIdx": 1, "runs": runs}],
        }

        xml_bytes, _image_rels, _link_rels = renderer._clone_layout_as_slide(
            default_template_layout2, slide_data, str(tmp_path), slide_num=1
        )

        root = ET.fromstring(xml_bytes)
        txBody = None
        for sp in root.findall(".//p:sp", NS):
            ph = sp.find(".//p:ph", NS)
            if ph is not None and ph.get("idx") == "1":
                txBody = sp.find(".//p:txBody", NS)
                break
        assert txBody is not None

        paragraphs = txBody.findall("a:p", NS)
        assert len(paragraphs) == 7

        # The two newly-created paragraphs (index 5, 6) must carry cloned
        # formatting, not bare <a:p><a:r><a:t> with nothing else.
        for extra_para in paragraphs[5:]:
            assert extra_para.find("a:pPr", NS) is not None, "new paragraph lost the template's pPr (bullet/indent)"
            run = extra_para.find("a:r", NS)
            assert run is not None
            assert run.find("a:rPr", NS) is not None, "new run lost the template's rPr (font/size/color)"
            assert run.find("a:t", NS).text == "Point 5" or run.find("a:t", NS).text == "Point 6"

        # And the reused paragraphs (0-4) still just get their text swapped in.
        for i, para in enumerate(paragraphs[:5]):
            run = para.find("a:r", NS)
            assert run.find("a:t", NS).text == f"Point {i}"


class TestPPTXRendererFromTemplate:
    """End-to-end coverage for the template-clone path (_render_from_template),
    which previously had zero test coverage even though it's the path real
    uploaded templates and the default template both go through."""

    @pytest.fixture
    def template_path(self, tmp_path):
        from app.renderer.default_template import get_default_template_path
        return get_default_template_path(output_dir=str(tmp_path / "templates"))

    def test_render_with_more_bullets_than_the_layout_has_paragraphs(self, template_path, tmp_path):
        renderer = PPTXRenderer(template_path=template_path)
        runs = [{"text": f"Point {i}", "bullet": True, "level": 0} for i in range(7)]
        slides = [
            {
                "slideIndex": 0,
                "layoutId": "2",
                "layoutName": "Title and Content",
                "contents": [{"type": "text", "placeholderIdx": 1, "runs": runs}],
            }
        ]
        output_path = str(tmp_path / "out.pptx")

        result = renderer.render(slides=slides, metadata={"title": "Test"}, output_path=output_path)

        assert os.path.exists(result)
        assert os.path.getsize(result) > 0

        from pptx import Presentation
        prs = Presentation(result)
        assert len(prs.slides) == 1

        # No duplicate relationship IDs anywhere in the repacked package.
        with zipfile.ZipFile(result) as zf:
            rels_xml = zf.read("ppt/_rels/presentation.xml.rels")
            rels_root = ET.fromstring(rels_xml)
            ids = [rel.get("Id") for rel in rels_root]
            assert len(ids) == len(set(ids)), f"duplicate relationship IDs: {ids}"
class TestPPTXRendererColorTheme:
    """A color_theme choice (see app.design_system.COLOR_THEMES) must
    actually change the rendered deck on the no-template path - otherwise
    asking the user for a preference and then ignoring it would be worse
    than not asking at all.

    The theme now reaches a slide two ways: the cover's background is the
    palette's ``dark``, and a content slide's title is its ``primary``. Both
    are asserted, because the earlier version of this test only ever looked
    at a title-only slide - which is a *cover*, and so renders a white title
    on brand-dark in every theme.
    """

    def _deck(self):
        return [
            {
                "slideIndex": 0, "layoutId": "title", "layoutName": "Title Slide",
                "contents": [{
                    "type": "text", "placeholderIdx": 0,
                    "runs": [{"text": "Cover", "bold": True}],
                }],
            },
            {
                "slideIndex": 1, "layoutId": "content", "layoutName": "Title and Content",
                "contents": [
                    {"type": "text", "placeholderIdx": 0,
                     "runs": [{"text": "Title", "bold": True}]},
                    {"type": "text", "placeholderIdx": 1,
                     "runs": [{"text": "A point", "bullet": True, "level": 0}]},
                ],
            },
        ]

    def _render(self, color_theme):
        """Return ``(cover_background_rgb, content_title_run)``."""
        renderer = PPTXRenderer(template_path=None, color_theme=color_theme)
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name
        try:
            result = renderer.render(
                slides=self._deck(), metadata={"title": "T"}, output_path=output_path,
            )
            from pptx import Presentation
            prs = Presentation(result)
            background = prs.slides[0].background.fill.fore_color.rgb
            for shape in prs.slides[1].shapes:
                if shape.has_text_frame and shape.text_frame.text == "Title":
                    return background, shape.text_frame.paragraphs[0].runs[0]
            raise AssertionError("title run not found")
        finally:
            os.unlink(output_path)

    @staticmethod
    def _palette(theme, slot):
        return _hex_to_rgbcolor(
            next(c["value"] for c in COLOR_THEMES[theme] if c["name"] == slot)
        )

    def test_warm_theme_applies_the_warm_palette(self):
        background, run = self._render("warm")
        assert background == self._palette("warm", "dark")
        assert run.font.color.rgb == self._palette("warm", "primary")

    def test_no_theme_defaults_to_the_professional_palette(self):
        background, run = self._render(None)
        assert background == self._palette("professional", "dark")
        assert run.font.color.rgb == self._palette("professional", "primary")

    def test_unknown_theme_falls_back_to_professional(self):
        _, run = self._render("not-a-real-theme")
        assert run.font.color.rgb == self._palette("professional", "primary")

    def test_different_themes_produce_different_decks(self):
        warm_bg, warm_run = self._render("warm")
        mono_bg, mono_run = self._render("monochrome")
        assert warm_run.font.color.rgb != mono_run.font.color.rgb
        assert warm_bg != mono_bg

    def test_the_cover_is_set_on_the_brand_dark_rather_than_white(self):
        """The light/dark rhythm is the design language's load-bearing part.

        Every slide on the same white page is what made a generated deck read
        as seven unrelated documents rather than one artifact.
        """
        background, _ = self._render("professional")
        assert background != _hex_to_rgbcolor("#FFFFFF")


class TestPPTXRendererDefaultShapeSizing:
    """Regression: PositionConfig's fields are documented as inches (plain
    floats - see layout.py), but _add_text_box/_add_image used to pass them
    straight to python-pptx's add_textbox/add_picture, which interpret a
    bare number as EMU (1 inch = 914400 EMU). A "1.5 inch" position became a
    shape 1.5 EMU wide - sub-pixel, invisible - on every slide of the
    no-template render path. python-pptx (and this test) can still read the
    (garbage) values back, which is exactly why this shipped unnoticed:
    only an actual rendering engine (PowerPoint/LibreOffice) showed a blank
    slide. Confirmed by converting a rendered file with LibreOffice headless
    and finding zero extractable text before this fix, full text after.

    Slide content now goes into the layout's own placeholders, so
    ``_add_text_box`` is only reached as a fallback (a body item with no
    placeholder left to fill). These tests therefore exercise it directly and
    via that overflow path, which is where the unit bug could still recur.
    """

    def test_text_box_is_sized_in_inches_not_raw_emu(self):
        from pptx import Presentation
        from pptx.util import Inches
        from app.renderer.layout import PositionConfig

        renderer = PPTXRenderer(template_path=None)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
        renderer._add_text_box(
            slide,
            [{"text": "Title", "bold": True, "fontSizePt": 36}],
            PositionConfig(left=1.0, top=1.0, width=8.0, height=1.2),
            is_title=True,
        )
        shape = next(s for s in slide.shapes if s.has_text_frame)

        # The old bug produced shapes a handful of EMU wide/tall; a
        # correctly-converted inch-based size is always comfortably
        # above a tenth of an inch for any real slide element.
        assert shape.width > Inches(0.1)
        assert shape.height > Inches(0.1)
        assert shape.width == Inches(8.0)
        assert shape.left == Inches(1.0)

    def test_body_items_beyond_the_available_placeholders_do_not_stack(self):
        """Overflow body blocks fall back to explicitly positioned text boxes.
        Each must advance a vertical cursor - drawing them all at the same
        calculated content position would silently overlap them."""
        from pptx import Presentation
        from pptx.util import Inches

        renderer = PPTXRenderer(template_path=None)
        # Four body items: more than any stock layout has body placeholders,
        # so the last ones take the fallback path.
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Title"}]},
                {"type": "text", "placeholderIdx": 1, "runs": [{"text": "Block one", "bullet": True}]},
                {"type": "text", "placeholderIdx": 2, "runs": [{"text": "Block two", "bullet": True}]},
                {"type": "text", "placeholderIdx": 3, "runs": [{"text": "Block three", "bullet": True}]},
            ],
        }]
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name
        try:
            renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)
            prs = Presentation(output_path)
            overflow = [s for s in prs.slides[0].shapes
                        if s.has_text_frame and not s.is_placeholder
                        and not is_chrome_shape(s)]
            assert overflow, "expected at least one fallback text box"
            for shape in overflow:
                assert shape.top < Inches(1000), "top offset leaked raw EMU into the inches field"
                assert shape.height > Inches(0.1)
            tops = sorted(s.top for s in overflow)
            assert len(set(tops)) == len(tops), "fallback blocks were drawn on top of each other"
        finally:
            os.unlink(output_path)


class TestTextFitAutofit:
    """Unit coverage for app.renderer.text_fit, the shrink-to-fit estimator
    used to stop long text from visually overlapping other content - the
    root cause confirmed by rendering long text with LibreOffice and seeing
    real overlap even though validate_font_size only ever enforced a
    minimum size, never shrank anything."""

    def test_short_text_needs_no_shrinking(self):
        from app.renderer import text_fit
        blocks = [text_fit.TextBlock(text="Short line", font_size_pt=18)]
        scale, ln_reduction = text_fit.compute_autofit_scale(blocks, box_width_pt=400, box_height_pt=200)
        assert scale == 1.0
        assert ln_reduction == 0

    def test_long_text_in_a_small_box_is_shrunk(self):
        from app.renderer import text_fit
        long_text = " ".join(["word"] * 400)
        blocks = [text_fit.TextBlock(text=long_text, font_size_pt=20)]
        scale, _ = text_fit.compute_autofit_scale(blocks, box_width_pt=300, box_height_pt=100)
        assert scale < 1.0
        assert scale >= text_fit.MIN_FONT_SCALE

    def test_zero_size_box_is_left_alone(self):
        """No usable geometry (e.g. a placeholder with no explicit xfrm) must
        not be treated as "infinitely small" - that would force every
        placeholder down to the minimum scale for no reason."""
        from app.renderer import text_fit
        blocks = [text_fit.TextBlock(text="Anything", font_size_pt=18)]
        scale, ln_reduction = text_fit.compute_autofit_scale(blocks, box_width_pt=0, box_height_pt=0)
        assert scale == 1.0
        assert ln_reduction == 0


class TestCloneLayoutAutofit:
    """The template-clone path must never touch <a:rPr> (that's the whole
    point of cloning - see the module docstring), so overflow protection has
    to be the OOXML-native <a:normAutofit> render-time hint instead."""

    @pytest.fixture
    def renderer(self):
        return PPTXRenderer(template_path=None)

    @pytest.fixture
    def small_boxed_layout(self, tmp_path):
        """A synthetic single-placeholder layout with an explicit, small
        <a:xfrm> so the fit estimate has real geometry to judge against."""
        xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sldLayout xmlns:a="{NS['a']}" xmlns:r="{NS['r']}" xmlns:p="{NS['p']}">
  <p:cSld>
    <p:spTree>
      <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
      <p:grpSpPr/>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="2" name="Body"/>
          <p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>
          <p:nvPr><p:ph type="body" idx="1"/></p:nvPr>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="457200" y="457200"/><a:ext cx="914400" cy="457200"/></a:xfrm>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:pPr/><a:r><a:rPr sz="2000"/><a:t>placeholder</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:sldLayout>"""
        layout_path = tmp_path / "slideLayout9.xml"
        layout_path.write_text(xml)
        return str(layout_path)

    def test_overlong_text_gets_normautofit_without_touching_rpr(self, renderer, small_boxed_layout, tmp_path):
        # The box is 1in x 0.5in at 20pt - far too small for several long
        # bullet points at that size, so this must trigger autofit.
        runs = [
            {"text": "This is a fairly long bullet point that will not fit", "bullet": True},
            {"text": "Here is another long bullet point stretching the box", "bullet": True},
            {"text": "And a third one to make sure it clearly overflows", "bullet": True},
        ]
        slide_data = {
            "slideIndex": 0,
            "layoutId": "9",
            "contents": [{"type": "text", "placeholderIdx": 1, "runs": runs}],
        }

        xml_bytes, _, _ = renderer._clone_layout_as_slide(small_boxed_layout, slide_data, str(tmp_path), slide_num=1)
        root = ET.fromstring(xml_bytes)

        sp = next(sp for sp in root.findall(".//p:sp", NS) if sp.find(".//p:ph", NS).get("idx") == "1")
        txBody = sp.find(".//p:txBody", NS)
        bodyPr = txBody.find("a:bodyPr", NS)
        assert bodyPr is not None
        normAutofit = bodyPr.find("a:normAutofit", NS)
        assert normAutofit is not None, "overflowing placeholder must get a normAutofit hint"
        assert int(normAutofit.get("fontScale")) < 100000

    def test_short_text_gets_no_autofit_element(self, renderer, small_boxed_layout, tmp_path):
        slide_data = {
            "slideIndex": 0,
            "layoutId": "9",
            "contents": [{"type": "text", "placeholderIdx": 1, "runs": [{"text": "Hi", "bullet": True}]}],
        }
        xml_bytes, _, _ = renderer._clone_layout_as_slide(small_boxed_layout, slide_data, str(tmp_path), slide_num=1)
        root = ET.fromstring(xml_bytes)
        sp = next(sp for sp in root.findall(".//p:sp", NS) if sp.find(".//p:ph", NS).get("idx") == "1")
        bodyPr = sp.find(".//p:txBody/a:bodyPr", NS)
        assert bodyPr is None or bodyPr.find("a:normAutofit", NS) is None


class TestCloneLayoutImages:
    """Regression: the template-clone path used to silently drop every
    non-text content item (`if content.get("type") != "text": continue`),
    so an image the LLM placed on a slide never made it into the output
    whenever a user-uploaded template was in use - the primary path. Images
    must now render there too."""

    # A minimal valid 1x1 transparent PNG.
    ONE_PX_PNG_B64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="

    @pytest.fixture
    def renderer(self):
        return PPTXRenderer(template_path=None)

    @pytest.fixture
    def default_template_layout2(self, tmp_path):
        from app.renderer.default_template import get_default_template_path
        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        working_dir = tmp_path / "extracted"
        with zipfile.ZipFile(template_path) as zf:
            zf.extractall(working_dir)
        return str(working_dir / "ppt" / "slideLayouts" / "slideLayout2.xml"), str(working_dir)

    def test_base64_image_is_added_as_a_picture_shape(self, renderer, default_template_layout2, tmp_path):
        layout_file, working_dir = default_template_layout2
        slide_data = {
            "slideIndex": 0,
            "layoutId": "2",
            "contents": [{
                "type": "image",
                "image": {"source": "base64", "data": self.ONE_PX_PNG_B64, "altText": "Test image"},
            }],
        }

        xml_bytes, image_rels, _ = renderer._clone_layout_as_slide(layout_file, slide_data, working_dir, slide_num=1)

        assert len(image_rels) == 1
        assert image_rels[0]["rId"] == "rId2"
        assert image_rels[0]["target"].startswith("../media/image")

        root = ET.fromstring(xml_bytes)
        pics = root.findall(".//p:pic", NS)
        assert len(pics) == 1
        blip = pics[0].find(".//a:blip", NS)
        assert blip.get(f"{{{NS['r']}}}embed") == "rId2"

        # The image bytes actually landed on disk.
        media_dir = os.path.join(working_dir, "ppt", "media")
        assert os.path.exists(media_dir)
        assert len(os.listdir(media_dir)) >= 1

    def test_data_uri_prefixed_base64_is_decoded(self, renderer, default_template_layout2, tmp_path):
        layout_file, working_dir = default_template_layout2
        slide_data = {
            "slideIndex": 0,
            "layoutId": "2",
            "contents": [{
                "type": "image",
                "image": {"source": "base64", "data": f"data:image/png;base64,{self.ONE_PX_PNG_B64}"},
            }],
        }
        xml_bytes, image_rels, _ = renderer._clone_layout_as_slide(layout_file, slide_data, working_dir, slide_num=1)
        assert len(image_rels) == 1

    def test_unresolvable_image_is_skipped_not_crashed(self, renderer, default_template_layout2, tmp_path):
        layout_file, working_dir = default_template_layout2
        slide_data = {
            "slideIndex": 0,
            "layoutId": "2",
            "contents": [{"type": "image", "image": {"source": "base64", "data": "not-valid-base64!!"}}],
        }
        xml_bytes, image_rels, _ = renderer._clone_layout_as_slide(layout_file, slide_data, working_dir, slide_num=1)
        assert image_rels == []
        root = ET.fromstring(xml_bytes)
        assert root.findall(".//p:pic", NS) == []

    def test_end_to_end_render_with_template_and_image_produces_valid_pptx(self, tmp_path):
        """Full render() through the template path with an image content
        item must produce a file python-pptx (and therefore PowerPoint/
        LibreOffice) can open, with the picture actually present."""
        from app.renderer.default_template import get_default_template_path
        from pptx import Presentation

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        renderer = PPTXRenderer(template_path=template_path)
        slides = [{
            "slideIndex": 0,
            "layoutId": "2",
            "layoutName": "Title and Content",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Title"}]},
                {"type": "image", "image": {"source": "base64", "data": self.ONE_PX_PNG_B64}},
            ],
        }]
        output_path = str(tmp_path / "out.pptx")
        result = renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        assert os.path.exists(result)
        prs = Presentation(result)
        assert len(prs.slides) == 1
        picture_shapes = [s for s in prs.slides[0].shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
        assert len(picture_shapes) == 1


class TestTableRendering:
    """Tables were previously schema-legal (validate_slide_json accepted
    them, design_slide_content's prompt told the LLM to use them) but
    functionally dead - no add_table call existed anywhere in the renderer
    on either render path. Both paths must now produce a real table, styled
    from the run's design system rather than python-pptx defaults."""

    TABLE_CONTENT = {
        "type": "table",
        "table": {
            "headerRow": True,
            "rows": [["Region", "Revenue"], ["EMEA", "$4.2M"], ["APAC", "$3.1M"]],
        },
    }

    BRAND_DESIGN_SYSTEM = {
        "palette": {"colors": [{"name": "primary", "value": "#00704A", "type": "hex"}], "schemeName": "brand"},
    }

    def test_default_path_produces_a_real_table_shape(self, tmp_path):
        from pptx import Presentation

        renderer = PPTXRenderer(template_path=None, design_system=self.BRAND_DESIGN_SYSTEM)
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Regional Revenue", "bold": True, "fontSizePt": 36}]},
                self.TABLE_CONTENT,
            ],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        prs = Presentation(output_path)
        tables = [s for s in prs.slides[0].shapes if s.has_table]
        assert len(tables) == 1
        table = tables[0].table
        assert len(table.rows) == 3
        assert len(table.columns) == 2
        assert table.cell(0, 0).text == "Region"
        assert table.cell(1, 1).text == "$4.2M"

    def test_default_path_header_row_uses_the_design_system_primary_color(self, tmp_path):
        from pptx import Presentation

        renderer = PPTXRenderer(template_path=None, design_system=self.BRAND_DESIGN_SYSTEM)
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [self.TABLE_CONTENT],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        prs = Presentation(output_path)
        table = next(s for s in prs.slides[0].shapes if s.has_table).table
        header_cell = table.cell(0, 0)
        assert header_cell.fill.fore_color.rgb == _hex_to_rgbcolor("#00704A")

    def test_template_path_produces_a_graphic_frame_table(self, tmp_path):
        from app.renderer.default_template import get_default_template_path

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        renderer = PPTXRenderer(template_path=template_path, design_system=self.BRAND_DESIGN_SYSTEM)
        slides = [{
            "slideIndex": 0,
            "layoutId": "2",
            "layoutName": "Title and Content",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Title"}]},
                self.TABLE_CONTENT,
            ],
        }]
        output_path = str(tmp_path / "out.pptx")
        result = renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        from pptx import Presentation
        prs = Presentation(result)
        assert len(prs.slides) == 1
        tables = [s for s in prs.slides[0].shapes if s.has_table]
        assert len(tables) == 1
        table = tables[0].table
        assert table.cell(0, 0).text == "Region"
        assert table.cell(2, 1).text == "$3.1M"

    def test_no_design_system_falls_back_to_default_colors(self, tmp_path):
        """A run with no extracted design system (e.g. Stage D not landed
        yet, or extraction genuinely produced nothing) must still render a
        table, just with the built-in default colors - never crash."""
        from pptx import Presentation

        renderer = PPTXRenderer(template_path=None, design_system=None)
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [self.TABLE_CONTENT],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)
        prs = Presentation(output_path)
        assert len([s for s in prs.slides[0].shapes if s.has_table]) == 1

    def test_empty_rows_are_skipped_not_crashed(self, tmp_path):
        from pptx import Presentation

        renderer = PPTXRenderer(template_path=None)
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [{"type": "table", "table": {"rows": []}}],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)
        prs = Presentation(output_path)
        assert len([s for s in prs.slides[0].shapes if s.has_table]) == 0


class TestChartRendering:
    """Charts were previously schema-legal but functionally dead - no chart
    rendering existed on either render path. Both must now produce a real
    picture (matplotlib PNG), styled from the run's design system."""

    CHART_CONTENT = {
        "type": "chart",
        "chart": {
            "chartType": "bar",
            "categories": ["Q1", "Q2", "Q3"],
            "series": [{"name": "Revenue", "values": [10, 20, 15]}],
        },
    }

    BRAND_DESIGN_SYSTEM = {
        "palette": {"colors": [{"name": "primary", "value": "#00704A", "type": "hex"}]},
    }

    def test_default_path_produces_a_picture_for_the_chart(self, tmp_path):
        from pptx import Presentation

        renderer = PPTXRenderer(template_path=None, design_system=self.BRAND_DESIGN_SYSTEM)
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Revenue", "bold": True, "fontSizePt": 36}]},
                self.CHART_CONTENT,
            ],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        prs = Presentation(output_path)
        pictures = [s for s in prs.slides[0].shapes if s.shape_type == 13]
        assert len(pictures) == 1

    def test_template_path_produces_a_picture_for_the_chart(self, tmp_path):
        from app.renderer.default_template import get_default_template_path
        from pptx import Presentation

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        renderer = PPTXRenderer(template_path=template_path, design_system=self.BRAND_DESIGN_SYSTEM)
        slides = [{
            "slideIndex": 0,
            "layoutId": "2",
            "layoutName": "Title and Content",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Title"}]},
                self.CHART_CONTENT,
            ],
        }]
        output_path = str(tmp_path / "out.pptx")
        result = renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)

        prs = Presentation(result)
        assert len(prs.slides) == 1
        pictures = [s for s in prs.slides[0].shapes if s.shape_type == 13]
        assert len(pictures) == 1

    def test_unrenderable_chart_is_skipped_not_crashed(self, tmp_path):
        """Empty chart data must not raise or produce an empty picture shape."""
        from pptx import Presentation

        renderer = PPTXRenderer(template_path=None)
        slides = [{
            "slideIndex": 0,
            "layoutId": "1",
            "layoutName": "Default",
            "contents": [{"type": "chart", "chart": {"chartType": "bar", "series": []}}],
        }]
        output_path = str(tmp_path / "out.pptx")
        renderer.render(slides=slides, metadata={"title": "T"}, output_path=output_path)
        prs = Presentation(output_path)
        assert len([s for s in prs.slides[0].shapes if s.shape_type == 13]) == 0

    def test_template_path_chart_registers_a_media_relationship(self, tmp_path):
        from app.renderer.default_template import get_default_template_path

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        renderer = PPTXRenderer(template_path=template_path, design_system=self.BRAND_DESIGN_SYSTEM)
        working_dir = tmp_path / "extracted"
        with zipfile.ZipFile(template_path) as zf:
            zf.extractall(working_dir)
        layout_file = str(working_dir / "ppt" / "slideLayouts" / "slideLayout2.xml")

        slide_data = {"slideIndex": 0, "layoutId": "2", "contents": [self.CHART_CONTENT]}
        xml_bytes, image_rels, _ = renderer._clone_layout_as_slide(layout_file, slide_data, str(working_dir), slide_num=1)

        assert len(image_rels) == 1
        root = ET.fromstring(xml_bytes)
        assert len(root.findall(".//p:pic", NS)) == 1


class TestDefaultPathPlaceholderRendering:
    """
    Content on the no-template path must land in the layout's own
    placeholders.

    ``prs.slides.add_slide(layout)`` copies the layout's placeholders onto the
    slide. The renderer used to ignore them and stack free-floating text boxes
    at LayoutEngine coordinates instead, so every generated slide carried an
    empty "Title"/"Subtitle" placeholder *plus* unaligned text boxes drawn over
    the top of them - the "boring, misaligned, looks like an unfinished
    template" output this guards against.
    """

    @staticmethod
    def _render(slides):
        from pptx import Presentation

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name
        PPTXRenderer(template_path=None).render(
            slides=slides, metadata={"title": "T"}, output_path=output_path
        )
        try:
            return Presentation(output_path)
        finally:
            os.unlink(output_path)

    @staticmethod
    def _content_slide(runs):
        return {
            "slideIndex": 0, "layoutId": "1", "layoutName": "Default",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "A Title", "bold": True}]},
                {"type": "text", "placeholderIdx": 1, "runs": runs},
            ],
        }

    def test_body_text_fills_a_placeholder_rather_than_a_floating_box(self):
        prs = self._render([self._content_slide([
            {"text": "First point", "bullet": True, "level": 0},
            {"text": "Second point", "bullet": True, "level": 0},
        ])])
        slide = prs.slides[0]
        filled = [s for s in slide.shapes
                  if s.has_text_frame and s.text_frame.text.strip()
                  and not is_chrome_shape(s)]
        assert filled, "slide has no text at all"
        assert all(s.is_placeholder for s in filled), (
            "text was rendered into free-floating boxes instead of the layout's placeholders"
        )

    def test_no_empty_placeholder_survives_into_the_output(self):
        """An unfilled placeholder is not invisible - PowerPoint shows its
        'Click to add title' prompt, which is what made decks look
        half-finished."""
        prs = self._render([self._content_slide([
            {"text": "Only point", "bullet": True, "level": 0},
        ])])
        empty = [s for s in prs.slides[0].shapes
                 if s.is_placeholder and s.has_text_frame
                 and not s.text_frame.text.strip()]
        assert empty == []

    def test_a_bulleted_content_slide_does_not_use_the_title_layout(self):
        """infer_layout_type used to classify any slide with <=2 content items
        as a title slide - which is the shape of *every* ordinary content
        slide - so whole decks rendered on the title layout."""
        prs = self._render([self._content_slide([
            {"text": "First point", "bullet": True, "level": 0},
            {"text": "Second point", "bullet": True, "level": 0},
            {"text": "Third point", "bullet": True, "level": 0},
        ])])
        assert "title slide" not in (prs.slides[0].slide_layout.name or "").lower()

    def test_a_real_title_slide_still_uses_the_title_layout(self):
        prs = self._render([{
            "slideIndex": 0, "layoutId": "1", "layoutName": "Default",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Deck Title", "bold": True}]},
                {"type": "text", "placeholderIdx": 1, "runs": [{"text": "A subtitle", "bullet": False}]},
            ],
        }])
        assert "title slide" in (prs.slides[0].slide_layout.name or "").lower()

    def test_bulleted_runs_get_a_real_bullet_glyph(self):
        """p.level alone only sets indent - a text placeholder renders no
        marker without an explicit <a:buChar>, which is why output came out as
        plain lines of text."""
        from pptx.oxml.ns import qn

        prs = self._render([self._content_slide([
            {"text": "First point", "bullet": True, "level": 0},
            {"text": "Nested point", "bullet": True, "level": 1},
        ])])
        body = next(s for s in prs.slides[0].shapes
                    if s.is_placeholder and "First point" in s.text_frame.text)
        chars = [
            p._p.find(qn("a:pPr")).find(qn("a:buChar")).get("char")
            for p in body.text_frame.paragraphs
        ]
        assert all(c for c in chars), f"missing bullet glyphs: {chars}"
        assert chars[0] != chars[1], "nested level should use a distinct marker"

    def test_a_title_run_is_explicitly_unbulleted(self):
        from pptx.oxml.ns import qn

        prs = self._render([self._content_slide([
            {"text": "Only point", "bullet": True, "level": 0},
        ])])
        title = next(s for s in prs.slides[0].shapes
                     if s.is_placeholder and s.text_frame.text.strip() == "A Title")
        pPr = title.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
        assert pPr.find(qn("a:buNone")) is not None


class TestDefaultTemplateGeometry:
    """
    python-pptx's built-in template is 4:3 and its placeholders are positioned
    for that canvas. Resizing the canvas to 16:9 without rescaling them left
    the body placeholders hanging off the bottom of the slide.
    """

    def test_every_layout_placeholder_fits_inside_the_slide(self):
        from pptx import Presentation
        from app.renderer.default_template import DefaultTemplateGenerator

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            path = f.name
        try:
            DefaultTemplateGenerator.generate(path)
            prs = Presentation(path)
            for layout in prs.slide_masters[0].slide_layouts:
                for shape in layout.shapes:
                    if shape.top is None or shape.height is None:
                        continue
                    assert shape.top + shape.height <= prs.slide_height, (
                        f"{layout.name!r}/{shape.name!r} extends "
                        f"{shape.top + shape.height - prs.slide_height} EMU below the slide"
                    )
        finally:
            os.unlink(path)

    def test_rescaling_preserves_nonzero_geometry(self):
        """Setting `left` creates an <a:xfrm> whose other components default
        to zero, so a read-modify-write interleaved across the four attributes
        collapses top/width/height to 0."""
        from pptx import Presentation
        from app.renderer.default_template import DefaultTemplateGenerator

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            path = f.name
        try:
            DefaultTemplateGenerator.generate(path)
            prs = Presentation(path)
            layout = next(lay for lay in prs.slide_masters[0].slide_layouts
                          if (lay.name or "").lower() == "title and content")
            for shape in layout.shapes:
                assert shape.width and shape.width > 0, f"{shape.name!r} lost its width"
                assert shape.height and shape.height > 0, f"{shape.name!r} lost its height"
        finally:
            os.unlink(path)


class TestBrandCoherenceOnTheDefaultPath:
    """
    Tables and charts are styled from ``design_system`` because, unlike text,
    they have no placeholder formatting to inherit. With no template there may
    be no design system at all, in which case they used to fall back to
    table_style's built-in defaults - so a deck rendered in a brand colour got
    brand-coloured titles and bullets but a default-navy table header.
    """

    @staticmethod
    def _render_table_deck(color_theme):
        from pptx import Presentation

        slides = [{
            "slideIndex": 0, "layoutId": "1", "layoutName": "Default",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Comparison", "bold": True}]},
                {"type": "table", "placeholderIdx": 1, "table": {
                    "headerRow": True,
                    "rows": [["Year", "Revenue"], ["2024", "$36B"], ["2025", "$37B"]],
                }},
            ],
        }]
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name
        PPTXRenderer(template_path=None, color_theme=color_theme).render(
            slides=slides, metadata={"title": "T"}, output_path=output_path
        )
        try:
            prs = Presentation(output_path)
            table = next(s.table for s in prs.slides[0].shapes if s.has_table)
            return table.cell(0, 0).fill.fore_color.rgb
        finally:
            os.unlink(output_path)

    def test_table_header_follows_a_brand_hex_color_theme(self):
        from pptx.dml.color import RGBColor

        header_rgb = self._render_table_deck("#00704A")

        assert header_rgb == RGBColor(0x00, 0x70, 0x4A), (
            f"table header {header_rgb} ignored the brand colour theme"
        )

    def test_different_color_themes_produce_different_table_headers(self):
        assert self._render_table_deck("warm") != self._render_table_deck("monochrome")

    def test_an_explicit_design_system_still_wins_over_the_theme_default(self):
        """A template-derived design system must not be overwritten by the
        no-template fallback."""
        from pptx import Presentation
        from pptx.dml.color import RGBColor

        explicit = {
            "palette": {"colors": [{"name": "primary", "value": "#AA0000", "type": "hex"}]},
            "dataViz": {"tableHeaderBg": "#AA0000", "tableHeaderFg": "#FFFFFF"},
        }
        slides = [{
            "slideIndex": 0, "layoutId": "1", "layoutName": "Default",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "T", "bold": True}]},
                {"type": "table", "placeholderIdx": 1,
                 "table": {"headerRow": True, "rows": [["A", "B"], ["1", "2"]]}},
            ],
        }]
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            output_path = f.name
        try:
            PPTXRenderer(
                template_path=None, color_theme="#00704A", design_system=explicit
            ).render(slides=slides, metadata={"title": "T"}, output_path=output_path)
            prs = Presentation(output_path)
            table = next(s.table for s in prs.slides[0].shapes if s.has_table)
            assert table.cell(0, 0).fill.fore_color.rgb == RGBColor(0xAA, 0x00, 0x00)
        finally:
            os.unlink(output_path)


class TestTheCachedDefaultTemplateIsVersioned:
    """
    The default template is derived from code but cached as data on a
    persistent volume. Under its old unversioned name a deployed geometry fix
    never reached any environment that had generated one before: it kept
    serving a months-old file whose body placeholder hung 77pt below the
    bottom of the slide, so every chart on it was flagged as overflowing no
    matter how often the agent redesigned the slide.
    """

    def test_the_filename_carries_the_version(self, tmp_path):
        from app.renderer.default_template import (
            DEFAULT_TEMPLATE_VERSION, get_default_template_path,
        )

        path = get_default_template_path(str(tmp_path))
        assert f"_v{DEFAULT_TEMPLATE_VERSION}" in os.path.basename(path)

    def test_a_stale_unversioned_template_is_not_reused(self, tmp_path):
        from pptx import Presentation
        from app.renderer.default_template import get_default_template_path

        # A file from before versioning, with the broken 4:3 geometry.
        stale = tmp_path / "default_template.pptx"
        prs = Presentation()
        prs.save(str(stale))
        stale_bytes = stale.read_bytes()

        path = get_default_template_path(str(tmp_path))

        assert os.path.basename(path) != "default_template.pptx"
        assert pathlib.Path(path).read_bytes() != stale_bytes

    def test_superseded_templates_are_cleaned_up(self, tmp_path):
        from pptx import Presentation
        from app.renderer.default_template import get_default_template_path

        Presentation().save(str(tmp_path / "default_template.pptx"))
        Presentation().save(str(tmp_path / "default_template_v1.pptx"))

        path = get_default_template_path(str(tmp_path))

        remaining = sorted(p.name for p in tmp_path.glob("default_template*.pptx"))
        assert remaining == [os.path.basename(path)]

    def test_the_generated_template_fits_the_slide(self, tmp_path):
        """The actual defect the version bump exists to ship."""
        from pptx import Presentation
        from app.renderer.default_template import get_default_template_path

        prs = Presentation(get_default_template_path(str(tmp_path)))
        for layout in prs.slide_masters[0].slide_layouts:
            for shape in layout.shapes:
                if shape.top is None or shape.height is None:
                    continue
                assert shape.top + shape.height <= prs.slide_height, (
                    f"{layout.name!r}/{shape.name!r} hangs below the slide"
                )

    def test_a_second_call_reuses_the_current_version(self, tmp_path):
        from app.renderer.default_template import get_default_template_path

        first = get_default_template_path(str(tmp_path))
        mtime = os.path.getmtime(first)
        second = get_default_template_path(str(tmp_path))

        assert first == second
        assert os.path.getmtime(second) == mtime, "template was needlessly regenerated"


# A minimal valid 1x1 transparent PNG.
ONE_PX_PNG_B64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="


class TestHyperlinks:
    """
    A citation the reader cannot click is a citation they have to retype. Two
    render paths, two mechanisms: python-pptx for the default theme, raw XML
    (plus a package relationship) for the template clone.
    """

    @staticmethod
    def _cited_slide() -> dict:
        return {
            "slideIndex": 0,
            "layoutId": "2",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Revenue"}]},
                {
                    "type": "text",
                    "placeholderIdx": 1,
                    "runs": [
                        {
                            "text": "Revenue grew 15%",
                            "bullet": True,
                            "citation": {"text": " [1]", "link": "https://a.example/q3"},
                        },
                        {
                            "text": "[1] Q3 results - a.example - 2026-03-11",
                            "link": "https://a.example/q3",
                        },
                    ],
                },
            ],
        }

    def test_default_path_links_the_marker_without_linking_the_bullet(self, tmp_path):
        """The "[1]" is its own run: clicking the sentence must not navigate."""
        from pptx import Presentation

        output = str(tmp_path / "out.pptx")
        PPTXRenderer(template_path=None).render(
            slides=[self._cited_slide()], metadata={"title": "T"}, output_path=output
        )

        prs = Presentation(output)
        paragraphs = [
            p
            for shape in prs.slides[0].shapes
            if shape.has_text_frame
            for p in shape.text_frame.paragraphs
        ]
        bullet = next(p for p in paragraphs if p.text.startswith("Revenue grew"))

        assert len(bullet.runs) == 2
        assert bullet.runs[0].text == "Revenue grew 15%"
        assert bullet.runs[0].hyperlink.address is None
        assert bullet.runs[1].text == " [1]"
        assert bullet.runs[1].hyperlink.address == "https://a.example/q3"

    def test_default_path_links_a_whole_reference_line(self, tmp_path):
        from pptx import Presentation

        output = str(tmp_path / "out.pptx")
        PPTXRenderer(template_path=None).render(
            slides=[self._cited_slide()], metadata={"title": "T"}, output_path=output
        )

        prs = Presentation(output)
        runs = [
            r
            for shape in prs.slides[0].shapes
            if shape.has_text_frame
            for p in shape.text_frame.paragraphs
            for r in p.runs
        ]
        reference = next(r for r in runs if r.text.startswith("[1] Q3 results"))
        assert reference.hyperlink.address == "https://a.example/q3"

    def test_template_path_writes_hlinkclick_and_an_external_relationship(self, tmp_path):
        from app.renderer.default_template import get_default_template_path

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        output = str(tmp_path / "out.pptx")
        PPTXRenderer(template_path=template_path).render(
            slides=[self._cited_slide()], metadata={"title": "T"}, output_path=output
        )

        with zipfile.ZipFile(output) as zf:
            slide_xml = zf.read("ppt/slides/slide1.xml").decode()
            rels_xml = zf.read("ppt/slides/_rels/slide1.xml.rels").decode()

        root = ET.fromstring(slide_xml)
        links = root.findall(".//a:hlinkClick", NS)
        assert links, "no hyperlink reached the cloned slide"
        linked_ids = {link.get(f"{{{NS['r']}}}id") for link in links}

        rels = ET.fromstring(rels_xml)
        external = {
            rel.get("Id"): rel
            for rel in rels
            if rel.get("Type", "").endswith("/hyperlink")
        }
        assert linked_ids <= set(external), "a hlinkClick pointed at no relationship"
        for rel in external.values():
            # Without TargetMode PowerPoint reads the URL as a path inside the
            # package and calls the file corrupt.
            assert rel.get("TargetMode") == "External"
            assert rel.get("Target") == "https://a.example/q3"

    def test_template_path_reuses_one_relationship_per_url(self, tmp_path):
        """Both runs cite the same page; two rels for one URL is just bloat."""
        from app.renderer.default_template import get_default_template_path

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        renderer = PPTXRenderer(template_path=template_path)
        working_dir = tmp_path / "extracted"
        with zipfile.ZipFile(template_path) as zf:
            zf.extractall(working_dir)
        layout_file = str(working_dir / "ppt" / "slideLayouts" / "slideLayout2.xml")

        _xml, _images, link_rels = renderer._clone_layout_as_slide(
            layout_file, self._cited_slide(), str(working_dir), slide_num=1
        )

        assert [rel["target"] for rel in link_rels] == ["https://a.example/q3"]

    def test_template_path_link_ids_cannot_collide_with_image_ids(self, tmp_path):
        from app.renderer.default_template import get_default_template_path

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        renderer = PPTXRenderer(template_path=template_path)
        working_dir = tmp_path / "extracted"
        with zipfile.ZipFile(template_path) as zf:
            zf.extractall(working_dir)
        layout_file = str(working_dir / "ppt" / "slideLayouts" / "slideLayout2.xml")

        slide = self._cited_slide()
        slide["contents"].append({
            "type": "image",
            "image": {
                "source": "base64",
                "data": ONE_PX_PNG_B64,
            },
        })

        _xml, image_rels, link_rels = renderer._clone_layout_as_slide(
            layout_file, slide, str(working_dir), slide_num=1
        )

        assert image_rels and link_rels
        assert not ({r["rId"] for r in image_rels} & {r["rId"] for r in link_rels})

    def test_a_slide_with_no_links_gets_no_hyperlink_relationships(self, tmp_path):
        from app.renderer.default_template import get_default_template_path

        template_path = get_default_template_path(output_dir=str(tmp_path / "templates"))
        output = str(tmp_path / "out.pptx")
        PPTXRenderer(template_path=template_path).render(
            slides=[{
                "slideIndex": 0,
                "layoutId": "2",
                "contents": [
                    {"type": "text", "placeholderIdx": 1, "runs": [{"text": "Plain"}]}
                ],
            }],
            metadata={"title": "T"},
            output_path=output,
        )

        with zipfile.ZipFile(output) as zf:
            rels_xml = zf.read("ppt/slides/_rels/slide1.xml.rels").decode()
        assert "/hyperlink" not in rels_xml


class TestTemplatePackageIntegrity:
    """
    A deck built on an uploaded template has to be a file the rest of the
    world can open. Three defects here meant it was not: the slide master's
    relationship was deleted, [Content_Types].xml came back with a generated
    namespace prefix, and every title placeholder went unfilled. LibreOffice
    refused the package outright ("source file could not be loaded"), which
    also meant validate_rendered_output could never inspect one.
    """

    @pytest.fixture
    def template_path(self, tmp_path):
        from app.renderer.default_template import get_default_template_path
        return get_default_template_path(output_dir=str(tmp_path / "templates"))

    @staticmethod
    def _render(template_path, tmp_path, slides):
        output = str(tmp_path / "out.pptx")
        PPTXRenderer(template_path=template_path).render(
            slides=slides, metadata={"title": "T"}, output_path=output
        )
        return output

    def _titled_slide(self):
        return {
            "slideIndex": 0,
            "layoutId": "2",
            "contents": [
                {"type": "text", "placeholderIdx": 0, "runs": [{"text": "Sources"}]},
                {
                    "type": "text",
                    "placeholderIdx": 1,
                    "runs": [
                        {"text": f"Point {i}", "bullet": True, "level": 0}
                        for i in range(6)
                    ],
                },
            ],
        }

    def test_the_slide_master_relationship_survives(self, template_path, tmp_path):
        """presentation.xml points at it by rId; deleting it dangles the reference."""
        output = self._render(template_path, tmp_path, [self._titled_slide()])

        with zipfile.ZipFile(output) as zf:
            rels = ET.fromstring(zf.read("ppt/_rels/presentation.xml.rels"))
            presentation = ET.fromstring(zf.read("ppt/presentation.xml"))

        declared = {rel.get("Id") for rel in rels}
        master_ids = {
            m.get(f"{{{NS['r']}}}id")
            for m in presentation.findall("p:sldMasterIdLst/p:sldMasterId", NS)
        }
        assert master_ids
        assert master_ids <= declared

    def test_content_types_keeps_its_own_default_namespace(self, template_path, tmp_path):
        """<ns0:Types> is namespace-equivalent and LibreOffice still rejects it."""
        output = self._render(template_path, tmp_path, [self._titled_slide()])

        with zipfile.ZipFile(output) as zf:
            content_types = zf.read("[Content_Types].xml").decode()

        assert "ns0:" not in content_types
        assert 'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"' in content_types

    def test_rels_files_still_use_the_relationships_default_namespace(
        self, template_path, tmp_path
    ):
        """The content-types write must not leave the prefix map disturbed."""
        output = self._render(template_path, tmp_path, [self._titled_slide()])

        with zipfile.ZipFile(output) as zf:
            rels = zf.read("ppt/_rels/presentation.xml.rels").decode()

        assert "ns0:" not in rels

    def test_the_title_placeholder_is_filled(self, template_path, tmp_path):
        """
        A title placeholder is identified by type="title" and omits idx, so
        reading "no idx" as "no match" left every slide showing the layout's
        own prompt text.
        """
        output = self._render(template_path, tmp_path, [self._titled_slide()])

        with zipfile.ZipFile(output) as zf:
            slide = zf.read("ppt/slides/slide1.xml").decode()

        assert "Sources" in slide
        assert "Click to edit Master title style" not in slide

    def test_flat_bullets_do_not_render_as_a_cascade(self, template_path, tmp_path):
        """
        The layout's own paragraphs run levels 0-4 and were reused in order,
        so six level-0 bullets came out indented five deep.
        """
        output = self._render(template_path, tmp_path, [self._titled_slide()])

        with zipfile.ZipFile(output) as zf:
            root = ET.fromstring(zf.read("ppt/slides/slide1.xml"))

        body = next(
            sp for sp in root.findall(".//p:sp", NS)
            if (sp.find(".//p:ph", NS) is not None
                and sp.find(".//p:ph", NS).get("idx") == "1")
        )
        levels = [
            (p.find("a:pPr", NS).get("lvl") if p.find("a:pPr", NS) is not None else None)
            for p in body.findall(".//a:p", NS)
        ]
        assert set(levels) == {None}, f"expected every bullet at level 0, got {levels}"

    def test_a_requested_level_is_still_honoured(self, template_path, tmp_path):
        slide = self._titled_slide()
        slide["contents"][1]["runs"][2]["level"] = 1
        output = self._render(template_path, tmp_path, [slide])

        with zipfile.ZipFile(output) as zf:
            root = ET.fromstring(zf.read("ppt/slides/slide1.xml"))

        body = next(
            sp for sp in root.findall(".//p:sp", NS)
            if (sp.find(".//p:ph", NS) is not None
                and sp.find(".//p:ph", NS).get("idx") == "1")
        )
        levels = [
            (p.find("a:pPr", NS).get("lvl") if p.find("a:pPr", NS) is not None else None)
            for p in body.findall(".//a:p", NS)
        ]
        assert levels[2] == "1"

    def test_autofit_uses_the_masters_geometry_when_the_layout_inherits_it(
        self, template_path, tmp_path
    ):
        """
        A layout placeholder with no <a:xfrm> takes its size from the master.
        Giving up there meant the placeholders most likely to be overfilled
        were the only ones that never got a scale hint.
        """
        renderer = PPTXRenderer(template_path=template_path)
        working_dir = tmp_path / "extracted"
        with zipfile.ZipFile(template_path) as zf:
            zf.extractall(working_dir)
        layout_file = str(working_dir / "ppt" / "slideLayouts" / "slideLayout2.xml")

        boxes = renderer._master_placeholder_boxes(layout_file)

        assert boxes, "no master geometry was resolved"
        assert boxes[1]["cx"] > 0 and boxes[1]["cy"] > 0
        # Cached: the second call must not reparse.
        assert renderer._master_placeholder_boxes(layout_file) is boxes
