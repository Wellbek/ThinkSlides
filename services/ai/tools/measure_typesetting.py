"""
How wrong is our text measurement, against the renderer that actually draws it?

Builds a corpus of text boxes at known widths and sizes, renders them through
the same LibreOffice path the previews use, reads back how many lines each one
really broke into, and compares that against the old character-count estimate
and against real font metrics.

    PYTHONPATH=. .venv/bin/python tools/measure_typesetting.py
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile

import fitz
from pptx import Presentation
from pptx.util import Inches, Pt

from app.renderer.typeset.shape import wrap

SAMPLES = [
    "Illinois illiterate lilies fill in it, it is initial, it is illicit.",
    "WMWMW MMWWM QUAWWM MOMMA WOMBAT AWKWARD WAXWORKS.",
    "The results of our inaugural survey of insurance industry CROs, "
    "conducted with the IIF, demonstrate the expanding and diversifying "
    "risks insurers face now.",
    "Revenue held flat but the mix moved, and the movement is the finding.",
    "AI investments see unprecedented growth, reflecting rising importance.",
    "a b c d e f g h i j k l m n o p q r s t u v w x y z",
    "Semiconductor market surges, fuelled by demand for advanced technologies "
    "across every major manufacturing region.",
    "Short line.",
]
WIDTHS_IN = (1.6, 2.4, 3.2, 4.8)
SIZES_PT = (9.0, 12.0, 16.0, 24.0)
FAMILY = "Arial"
OLD_RATIO = 0.52


def old_estimate(text: str, width_pt: float, size_pt: float) -> int:
    """The current estimator: characters times a fixed average width."""
    per_line = max(1, int(width_pt / (size_pt * OLD_RATIO)))
    words, lines, current = text.split(), 0, 0
    for word in words:
        need = len(word) + (1 if current else 0)
        if current + need <= per_line:
            current += need
        else:
            lines += 1
            current = len(word)
    return lines + (1 if current else 0)


def main() -> int:
    work = tempfile.mkdtemp(prefix="typeset_")
    try:
        # One case per slide, in a box far taller than any of them needs.
        # Sharing a slide made the boxes short enough that LibreOffice
        # clipped the overflow, and counting *rendered* lines then reported
        # a 19-line block as 4 - measuring the box, not the wrap.
        deck = Presentation()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(24.0)
        blank = deck.slide_layouts[6]
        cases = []
        for text in SAMPLES:
            for width in WIDTHS_IN:
                for size in SIZES_PT:
                    slide = deck.slides.add_slide(blank)
                    box = slide.shapes.add_textbox(
                        Inches(0.3), Inches(0.3), Inches(width), Inches(23.0))
                    frame = box.text_frame
                    frame.word_wrap = True
                    frame.margin_left = frame.margin_right = 0
                    frame.margin_top = frame.margin_bottom = 0
                    para = frame.paragraphs[0]
                    run = para.add_run()
                    run.text = text
                    run.font.size = Pt(size)
                    run.font.name = FAMILY
                    cases.append({
                        "text": text, "width_pt": width * 72.0,
                        "size": size, "slide": len(deck.slides) - 1,
                    })
        pptx_path = os.path.join(work, "typeset.pptx")
        deck.save(pptx_path)

        from app.renderer.preview_generator import PreviewGenerator
        pdf_path = PreviewGenerator.convert_to_pdf(pptx_path, work)
        if not pdf_path:
            print("LibreOffice could not convert the corpus", file=sys.stderr)
            return 2
        doc = fitz.open(pdf_path)

        old_err = new_err = exact_old = exact_new = 0
        total = 0
        misses = []
        for case in cases:
            page = doc[case["slide"]]
            truth = 0
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    if any(s["text"].strip() for s in line["spans"]):
                        truth += 1
            if truth == 0:
                continue
            total += 1
            old = old_estimate(case["text"], case["width_pt"], case["size"])
            new = len(wrap(case["text"], case["width_pt"], FAMILY, case["size"]))
            old_err += abs(old - truth)
            new_err += abs(new - truth)
            exact_old += old == truth
            exact_new += new == truth
            if new != truth:
                misses.append({"w": case["width_pt"], "size": case["size"],
                               "truth": truth, "new": new, "text": case["text"]})

        if os.environ.get("SHOW_MISSES"):
            print("\n--- disagreements (new estimator) ---")
            for row in misses[:24]:
                print(f"  w={row['w']:5.0f}pt size={row['size']:4.1f} "
                      f"truth={row['truth']} ours={row['new']} "
                      f":: {row['text'][:44]!r}")
            print()
        print(f"cases measured against LibreOffice: {total}\n")
        print(f"{'estimator':28s} {'exact':>8s} {'mean |error| lines':>20s}")
        print(f"{'current (0.52 char ratio)':28s} {exact_old/total:7.1%} "
              f"{old_err/total:20.3f}")
        print(f"{'real font metrics':28s} {exact_new/total:7.1%} "
              f"{new_err/total:20.3f}")
        return 0
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
