"""
The primitive layer: the marks a slide is actually made of.

A slide's content used to be a closed union - text, image, table, chart - and
the block vocabulary above this module makes that union bigger without making
it open. Seven more members still cannot express "a hairline two thirds of the
way across", because that is not a block, it is a mark.

So there are three layers, and this is the floor:

* **regions** (``regions.py``) compose the canvas
* **blocks** (``blocks.py``) carry meaning - a stat callout, a row band - and
  **compile down to primitives**
* **primitives** (here) are marks: a text box, a rectangle, a line, an
  ellipse, a picture

Everything drawable on a slide is one of those five, so the primitive layer
has the same ceiling as writing python-pptx by hand. That is the point: the
escape hatch costs no expressivity, and unlike generated code it stays inside
the guard rails - containment, contrast and measurement all operate on
primitives, so a slide that drops to this layer is checked exactly as
thoroughly as one built from blocks.

Two properties are deliberate and load-bearing:

**Authored geometry is fractional, resolved geometry is absolute.** A
``PrimitiveSpec`` positions itself as fractions of the region it is bound to
(or of the canvas), which is what makes a composition survive a change of
canvas. ``resolve`` turns that into inches once, and everything downstream
works in inches.

**Every primitive has a stable id.** Nothing in this module needs it yet, and
that is fine: an addressable element is the difference between "move that
heading 4pt left" being a patch and being a full redesign of the slide.
"""

from __future__ import annotations

import math
import os
import tempfile
from dataclasses import dataclass, replace
from typing import Any, Iterable, Optional, Sequence

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app import colors
from app.renderer import text_fit
from app.renderer.default_template import _hex_to_rgbcolor
from app.renderer.layout import PositionConfig
from app.renderer.typography import TypeScale
from app.renderer.xml_utils import describe_picture

# The five marks. Anything a slide can carry is one of these; there is no
# sixth, which is why an open primitive layer is a bounded commitment rather
# than an ever-growing vocabulary.
PRIMITIVE_KINDS = frozenset({"text", "rect", "line", "ellipse", "image"})

ALIGNMENTS = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "centre": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}
ANCHORS = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "center": MSO_ANCHOR.MIDDLE,
    "centre": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}

# Shapes this layer puts on a slide are named with their primitive id, so an
# element can be found again.
#
# Two prefixes, because the distinction is load-bearing: geometry validation,
# dead-space checks and the density metrics all need to tell the deck's own
# furniture from the slide's content, and measuring a footer rule as if it
# were a body element is how a blank slide reads as full. Both prefixes live
# here rather than in the renderer because this is the layer that names
# shapes, and one place has to decide.
CONTENT_NAME_PREFIX = "ThinkSlides:"
CHROME_NAME_PREFIX = "ThinkSlides Chrome:"

# A mark thinner than this in either direction is a rule, not a box, and is
# drawn without being collapsed away by rounding.
MIN_MARK_IN = 0.004
# Fractional geometry rounds; a mark this far outside its frame is a mistake
# rather than a rounding artefact.
CONTAINMENT_TOLERANCE = 0.005
# Two text primitives overlapping by more than this share of the smaller one
# are drawn on top of each other. Matches the rendered-output check's own
# threshold, so the two passes agree about what an overlap is.
TEXT_OVERLAP_TOLERANCE = 0.35


def content_shape_name(primitive_id: str) -> str:
    return f"{CONTENT_NAME_PREFIX}{primitive_id}"


def chrome_shape_name(primitive_id: str) -> str:
    return f"{CHROME_NAME_PREFIX}{primitive_id}"


def is_chrome_shape(shape) -> bool:
    """True for shapes the renderer added as deck chrome, not slide content."""
    return str(getattr(shape, "name", "") or "").startswith(CHROME_NAME_PREFIX)


def is_content_shape(shape) -> bool:
    """True for shapes this layer placed as content, carrying a primitive id."""
    return str(getattr(shape, "name", "") or "").startswith(CONTENT_NAME_PREFIX)


def shape_primitive_id(shape) -> Optional[str]:
    """The primitive id a shape carries, whether it is content or chrome."""
    name = str(getattr(shape, "name", "") or "")
    for prefix in (CHROME_NAME_PREFIX, CONTENT_NAME_PREFIX):
        if name.startswith(prefix):
            return name[len(prefix):]
    return None


@dataclass(frozen=True)
class TextRunSpec:
    """One run inside a text primitive. Mirrors the shared TextRun schema."""
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    size_pt: Optional[float] = None
    color: Optional[str] = None
    font: Optional[str] = None
    # A line break before this run, so a primitive can carry a heading over a
    # body without becoming two primitives that have to be kept aligned.
    break_before: bool = False
    # A true superscript, set as a baseline offset rather than a smaller font
    # or a Unicode glyph, so it scales with the type around it.
    superscript: bool = False
    # Set by a type role rather than by an author: an eyebrow in small caps
    # with open tracking is a house style, not a per-run decision.
    smallcaps: bool = False
    tracking_per_mille: Optional[int] = None

    @classmethod
    def from_dict(cls, data: Any) -> "TextRunSpec":
        if isinstance(data, str):
            return cls(text=data)
        if not isinstance(data, dict):
            return cls(text=str(data))
        return cls(
            text=str(data.get("text", "")),
            bold=bool(data.get("bold", False)),
            italic=bool(data.get("italic", False)),
            underline=bool(data.get("underline", False)),
            size_pt=_positive_or_none(data.get("sizePt", data.get("fontSizePt"))),
            color=_hex_or_none(data.get("colorHex", data.get("color"))),
            font=str(data["fontName"]) if data.get("fontName") else None,
            break_before=bool(data.get("breakBefore", False)),
            superscript=bool(data.get("superscript", False)),
            smallcaps=bool(data.get("smallcaps", False)),
            tracking_per_mille=(
                int(data["trackingPerMille"])
                if data.get("trackingPerMille") is not None else None
            ),
        )


@dataclass(frozen=True)
class Primitive:
    """One resolved mark, positioned in inches on the slide."""
    id: str
    kind: str
    box: PositionConfig
    runs: tuple[TextRunSpec, ...] = ()
    fill: Optional[str] = None
    stroke: Optional[str] = None
    stroke_pt: float = 0.0
    align: str = "left"
    anchor: str = "top"
    wrap: bool = True
    autofit: bool = True
    size_pt: Optional[float] = None
    color: Optional[str] = None
    font: Optional[str] = None
    bold: bool = False
    image: Optional[dict] = None
    end: Optional[tuple[float, float]] = None    # line end, absolute inches
    alt_text: str = ""
    # ── typographic and geometric control the DesignDoc layer needs ────────
    # A picture is stretched into its box unless it says otherwise, which
    # distorts any photograph whose aspect does not match its slot - and no
    # slot's aspect is chosen to match a photograph. "cover" crops instead,
    # around a focal point; "contain" letterboxes.
    fill_alpha: float = 1.0
    fit: str = "fill"                            # fill | cover | contain
    focal: tuple[float, float] = (0.5, 0.5)
    crop: Optional[tuple[float, float, float, float]] = None  # l,t,r,b fractions
    rotation_deg: float = 0.0
    # Per-mille of the em: OOXML's a:rPr/@spc, in the unit it actually wants.
    tracking_per_mille: Optional[int] = None
    smallcaps: bool = False
    # Exact leading, so a page can hold a baseline rhythm rather than
    # whatever line spacing each box happened to inherit.
    line_spacing_pt: Optional[float] = None
    space_before_pt: float = 0.0
    space_after_pt: float = 0.0
    # Native multi-column text (a:bodyPr/@numCol), which stays editable.
    columns: int = 1
    column_gap_pt: float = 0.0
    # Deck furniture rather than slide content: named with the chrome prefix
    # so every downstream measurement keeps ignoring it.
    chrome: bool = False
    # The region this mark was positioned against, when it named one. Kept on
    # the resolved primitive because the ground it is coloured for follows
    # from it, and because an addressable element should say where it lives.
    region_id: Optional[str] = None

    @property
    def text(self) -> str:
        return "".join(run.text for run in self.runs)


@dataclass(frozen=True)
class PaintContext:
    """What a primitive needs to know that is not its own business."""
    ground: str
    default_color: str
    type_scale: TypeScale
    default_size_pt: float = 18.0
    default_font: Optional[str] = None


# --------------------------------------------------------------------------
# Authoring -> resolved
# --------------------------------------------------------------------------

def resolve(
    specs: Sequence[dict],
    frame_for: Any,
    canvas: PositionConfig,
    id_prefix: str = "p",
) -> list[Primitive]:
    """
    Turn authored primitive specs into resolved marks.

    ``frame_for`` maps a spec's ``regionId`` to the rectangle it is fractional
    against, returning None when the region is unknown - an unresolvable
    binding falls back to the canvas rather than dropping the mark, matching
    how a block bound to a missing region behaves.
    """
    out: list[Primitive] = []
    used_ids: set[str] = set()
    for index, spec in enumerate(specs or []):
        if not isinstance(spec, dict):
            continue
        kind = str(spec.get("kind") or "").strip()
        if kind not in PRIMITIVE_KINDS:
            continue

        frame = None
        region_id = spec.get("regionId")
        if region_id and frame_for is not None:
            frame = frame_for(str(region_id))
        frame = frame or canvas

        primitive = _resolve_one(
            spec, kind, frame,
            _unique(str(spec.get("id") or "").strip() or f"{id_prefix}-{index}",
                    used_ids),
            str(region_id) if region_id else None,
        )
        if primitive is not None:
            out.append(primitive)
    return out


def _resolve_one(spec: dict, kind: str, frame: PositionConfig,
                 primitive_id: str,
                 region_id: Optional[str]) -> Optional[Primitive]:
    if kind == "line":
        x1 = frame.left + _fraction(spec.get("x1"), 0.0) * frame.width
        y1 = frame.top + _fraction(spec.get("y1"), 0.0) * frame.height
        x2 = frame.left + _fraction(spec.get("x2"), 1.0) * frame.width
        y2 = frame.top + _fraction(spec.get("y2"), 0.0) * frame.height
        box = PositionConfig(left=min(x1, x2), top=min(y1, y2),
                             width=abs(x2 - x1), height=abs(y2 - y1))
        return Primitive(
            id=primitive_id, kind=kind, box=box, end=(x2, y2),
            stroke=_hex_or_none(spec.get("stroke")) or "#000000",
            stroke_pt=_positive(spec.get("strokeWeightPt"), 1.0),
            fill=None, region_id=region_id,
        )

    box = PositionConfig(
        left=frame.left + _fraction(spec.get("x"), 0.0) * frame.width,
        top=frame.top + _fraction(spec.get("y"), 0.0) * frame.height,
        width=_fraction(spec.get("w"), 1.0) * frame.width,
        height=_fraction(spec.get("h"), 1.0) * frame.height,
    )
    if box.width <= 0 or box.height <= 0:
        return None

    runs = tuple(
        TextRunSpec.from_dict(run) for run in (spec.get("runs") or [])
    ) if kind == "text" else ()
    if kind == "text" and not runs and spec.get("text"):
        runs = (TextRunSpec.from_dict({"text": spec.get("text")}),)

    return Primitive(
        id=primitive_id,
        kind=kind,
        box=box,
        runs=runs,
        fill=_hex_or_none(spec.get("fill")),
        stroke=_hex_or_none(spec.get("stroke")),
        stroke_pt=_positive(spec.get("strokeWeightPt"), 0.0),
        align=str(spec.get("align") or "left").lower(),
        anchor=str(spec.get("anchor") or "top").lower(),
        wrap=bool(spec.get("wrap", True)),
        autofit=bool(spec.get("autofit", True)),
        size_pt=_positive_or_none(spec.get("sizePt")),
        color=_hex_or_none(spec.get("color")),
        font=str(spec["font"]) if spec.get("font") else None,
        bold=bool(spec.get("bold", False)),
        image=spec.get("image") if isinstance(spec.get("image"), dict) else None,
        alt_text=str(spec.get("altText") or ""),
        region_id=region_id,
    )


# --------------------------------------------------------------------------
# Painting
# --------------------------------------------------------------------------

def paint_all(slide, primitives: Iterable[Primitive], context: PaintContext) -> list:
    """Paint primitives in order; later ones sit on top of earlier ones."""
    return [shape for shape in (paint(slide, p, context) for p in primitives)
            if shape is not None]


def paint(slide, primitive: Primitive, context: PaintContext):
    """Draw one primitive. Returns the shape, or None when nothing was drawn."""
    if slide is None:
        return None
    painter = _PAINTERS.get(primitive.kind)
    if painter is None:
        return None
    shape = painter(slide, primitive, context)
    if shape is not None:
        try:
            shape.name = (chrome_shape_name(primitive.id) if primitive.chrome
                          else content_shape_name(primitive.id))
        except (AttributeError, ValueError):
            pass
        if primitive.rotation_deg:
            try:
                shape.rotation = primitive.rotation_deg
            except (AttributeError, ValueError):
                pass
    return shape


def reorder_by_z(shapes_and_z: Sequence[tuple[object, int]]) -> None:
    """Repaint order, after the fact, by declared z.

    ``send_to_back`` already moves one shape in the tree; this generalises it
    to a full ordering, which is what a page with layered frames needs. The
    sort is stable, so shapes at equal z keep the order they were painted -
    a page renders the same way twice.
    """
    ordered = sorted(range(len(shapes_and_z)),
                     key=lambda i: (shapes_and_z[i][1], i))
    for shape, _ in shapes_and_z:
        element = getattr(shape, "_element", None)
        if element is None:
            return
    for index in ordered:
        shape = shapes_and_z[index][0]
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
            parent.append(element)


def _paint_text(slide, primitive: Primitive, context: PaintContext):
    if not primitive.text.strip():
        return None
    box = slide.shapes.add_textbox(
        Inches(primitive.box.left), Inches(primitive.box.top),
        Inches(max(MIN_MARK_IN, primitive.box.width)),
        Inches(max(MIN_MARK_IN, primitive.box.height)),
    )
    frame = box.text_frame
    frame.word_wrap = primitive.wrap
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = ANCHORS.get(primitive.anchor, MSO_ANCHOR.TOP)

    base_size = primitive.size_pt or context.default_size_pt
    scale = _autofit_scale(primitive, context, base_size)

    _apply_columns(frame, primitive)

    paragraph = frame.paragraphs[0]
    paragraph.alignment = ALIGNMENTS.get(primitive.align, PP_ALIGN.LEFT)
    _suppress_bullet(paragraph)
    _apply_paragraph_rhythm(paragraph, primitive, scale, first=True)

    for index, run_spec in enumerate(primitive.runs):
        if index and run_spec.break_before:
            paragraph = frame.add_paragraph()
            paragraph.alignment = ALIGNMENTS.get(primitive.align, PP_ALIGN.LEFT)
            _suppress_bullet(paragraph)
            _apply_paragraph_rhythm(paragraph, primitive, scale, first=False)
        run = paragraph.add_run()
        run.text = run_spec.text
        run.font.size = Pt(max(6.0, (run_spec.size_pt or base_size) * scale))
        run.font.bold = run_spec.bold or primitive.bold
        run.font.italic = run_spec.italic
        run.font.underline = run_spec.underline
        run.font.name = run_spec.font or primitive.font or context.default_font \
            or context.type_scale.body_font
        # Contrast is enforced here rather than trusted from upstream: it is
        # one comparison at the last point before the colour reaches the file,
        # and an unreadable run is a defect nothing later can see.
        preferred = run_spec.color or primitive.color or context.default_color
        try:
            run.font.color.rgb = _hex_to_rgbcolor(
                colors.readable_on(context.ground, preferred)
            )
        except ValueError:
            pass
        # Several stock Office layouts declare cap="all"; a block inheriting it
        # renders SHOUTED.
        rpr = run._r.get_or_add_rPr()
        rpr.set("cap", "none")
        if run_spec.smallcaps or primitive.smallcaps:
            rpr.set("cap", "small")
        tracking = run_spec.tracking_per_mille
        if tracking is None:
            tracking = primitive.tracking_per_mille
        if tracking:
            # OOXML counts character spacing in 1/100 pt, and a per-mille of
            # the em at this size is what a type role actually specifies.
            size_pt = (run_spec.size_pt or base_size) * scale
            rpr.set("spc", str(int(round(size_pt * tracking / 10.0))))
        if run_spec.superscript:
            from app.renderer.apparatus import apply_superscript

            apply_superscript(run)

    if primitive.fill:
        _apply_fill(box, primitive.fill, primitive.fill_alpha)
    return box


def _apply_columns(frame, primitive: Primitive) -> None:
    """Native multi-column text, which stays editable in PowerPoint.

    A story set in two columns is one text box with numCol=2, not two boxes
    the author has to keep in step - and PowerPoint reflows it correctly
    when the reader edits the copy.
    """
    if primitive.columns <= 1:
        return
    body = frame._txBody.bodyPr
    body.set("numCol", str(int(primitive.columns)))
    if primitive.column_gap_pt:
        body.set("spcCol", str(int(round(primitive.column_gap_pt * 12700))))


def _apply_paragraph_rhythm(paragraph, primitive: Primitive, scale: float,
                            *, first: bool) -> None:
    """Exact leading and paragraph spacing, in points.

    The primitive layer previously set no paragraph spacing at all, so every
    block took whatever leading it inherited and no page could hold a
    baseline rhythm. Stated in points rather than as a multiple because a
    rhythm is a fixed distance down the page, not a ratio of each block's
    own size.
    """
    if primitive.line_spacing_pt:
        paragraph.line_spacing = Pt(primitive.line_spacing_pt * scale)
    if primitive.space_after_pt:
        paragraph.space_after = Pt(primitive.space_after_pt * scale)
    if primitive.space_before_pt and not first:
        paragraph.space_before = Pt(primitive.space_before_pt * scale)


def _paint_rect(slide, primitive: Primitive, context: PaintContext):
    return _paint_autoshape(slide, primitive, MSO_SHAPE.RECTANGLE)


def _paint_ellipse(slide, primitive: Primitive, context: PaintContext):
    return _paint_autoshape(slide, primitive, MSO_SHAPE.OVAL)


def _paint_autoshape(slide, primitive: Primitive, shape_type):
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(primitive.box.left), Inches(primitive.box.top),
        Inches(max(MIN_MARK_IN, primitive.box.width)),
        Inches(max(MIN_MARK_IN, primitive.box.height)),
    )
    if primitive.fill:
        _apply_fill(shape, primitive.fill, primitive.fill_alpha)
    else:
        shape.fill.background()
    if primitive.stroke and primitive.stroke_pt > 0:
        try:
            shape.line.color.rgb = _hex_to_rgbcolor(primitive.stroke)
            shape.line.width = Pt(primitive.stroke_pt)
        except ValueError:
            shape.line.fill.background()
    else:
        shape.line.fill.background()
    _drop_theme_style(shape)
    return shape


def _paint_line(slide, primitive: Primitive, context: PaintContext):
    if primitive.end is None:
        return None
    x1, y1 = primitive.box.left, primitive.box.top
    x2, y2 = primitive.end
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    try:
        connector.line.color.rgb = _hex_to_rgbcolor(primitive.stroke or "#000000")
        connector.line.width = Pt(max(0.25, primitive.stroke_pt or 1.0))
    except ValueError:
        pass
    _drop_theme_style(connector)
    return connector


def _natural_aspect(data: bytes) -> Optional[float]:
    """Width over height of the actual image, or None if it will not open."""
    try:
        from io import BytesIO

        from PIL import Image

        with Image.open(BytesIO(data)) as img:
            width, height = img.size
        return (width / height) if height else None
    except Exception:
        return None


def cover_crop(image_aspect: float, box_aspect: float,
               focal: tuple[float, float]) -> tuple[float, float, float, float]:
    """The crop that makes an image fill a box without distorting it.

    Returned as (left, top, right, bottom) fractions of the original, which
    is the unit python-pptx's ``picture.crop_*`` wants. The focal point
    decides which part survives: a portrait cropped to a landscape slot
    should keep the face, and the centre is only the right answer when
    nothing better is known.
    """
    if image_aspect <= 0 or box_aspect <= 0:
        return (0.0, 0.0, 0.0, 0.0)
    if image_aspect > box_aspect:
        visible = box_aspect / image_aspect
        excess = max(0.0, 1.0 - visible)
        left = min(max(focal[0] - visible / 2.0, 0.0), excess)
        return (left, 0.0, excess - left, 0.0)
    visible = image_aspect / box_aspect
    excess = max(0.0, 1.0 - visible)
    top = min(max(focal[1] - visible / 2.0, 0.0), excess)
    return (0.0, top, 0.0, excess - top)


def _contain_box(image_aspect: float, box) -> "PositionConfig":
    """The largest rectangle of the image's aspect that fits inside the box."""
    from app.renderer.layout import PositionConfig

    if image_aspect <= 0:
        return box
    box_aspect = box.width / box.height if box.height else image_aspect
    if image_aspect > box_aspect:
        width = box.width
        height = width / image_aspect
    else:
        height = box.height
        width = height * image_aspect
    return PositionConfig(
        left=box.left + (box.width - width) / 2.0,
        top=box.top + (box.height - height) / 2.0,
        width=width,
        height=height,
    )


def _paint_image(slide, primitive: Primitive, context: PaintContext):
    from app.renderer.pptx_renderer import _resolve_image_bytes

    data = _resolve_image_bytes(primitive.image or {})
    if not data:
        return None

    box = primitive.box
    crop = primitive.crop
    if primitive.fit in ("cover", "contain"):
        aspect = _natural_aspect(data)
        if aspect:
            if primitive.fit == "contain":
                box = _contain_box(aspect, box)
            elif crop is None:
                box_aspect = (box.width / box.height) if box.height else aspect
                crop = cover_crop(aspect, box_aspect, primitive.focal)

    handle = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    try:
        handle.write(data)
        handle.close()
        picture = slide.shapes.add_picture(
            handle.name,
            Inches(box.left), Inches(box.top),
            Inches(max(MIN_MARK_IN, box.width)),
            Inches(max(MIN_MARK_IN, box.height)),
        )
    finally:
        os.unlink(handle.name)

    if crop:
        left, top, right, bottom = crop
        picture.crop_left = max(0.0, min(0.95, left))
        picture.crop_top = max(0.0, min(0.95, top))
        picture.crop_right = max(0.0, min(0.95, right))
        picture.crop_bottom = max(0.0, min(0.95, bottom))
    describe_picture(picture, primitive.alt_text)
    return picture


_PAINTERS = {
    "text": _paint_text,
    "rect": _paint_rect,
    "ellipse": _paint_ellipse,
    "line": _paint_line,
    "image": _paint_image,
}


# The first two children of a slide's shape tree are its group properties,
# not shapes; content starts after them.
_SPTREE_HEADER_ELEMENTS = 2


def send_to_back(shapes: Sequence) -> None:
    """
    Move ``shapes`` behind everything else on their slide, keeping their order.

    Painting order alone is not enough. ``add_slide(layout)`` clones the
    layout's placeholders onto the slide *first*, so anything drawn afterwards
    - a region's ground, a tinted panel a slide placed itself - lands on top
    of the placeholder text rather than behind it. A full-bleed ground would
    simply erase the slide's title.

    The shapes are reinserted as a block so their order relative to each other
    survives: a rail inside a panel still sits on top of its panel.
    """
    ordered = [shape for shape in shapes if shape is not None]
    if not ordered:
        return
    tree = ordered[0]._element.getparent()
    if tree is None:
        return
    for offset, shape in enumerate(ordered):
        element = shape._element
        parent = element.getparent()
        if parent is None:
            continue
        parent.remove(element)
        tree.insert(_SPTREE_HEADER_ELEMENTS + offset, element)


def _apply_fill(shape, color_hex: str, alpha: float = 1.0) -> None:
    """Solid fill, optionally translucent.

    Translucency is written as OOXML rather than through python-pptx, which
    exposes no alpha. It is needed for exactly one thing and that thing is
    not decoration: a scrim over a photograph, so type on top of it stays
    legible whatever the image happens to be doing underneath.
    """
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _hex_to_rgbcolor(color_hex)
    except (ValueError, AttributeError):
        return
    if alpha >= 1.0:
        return
    try:
        from pptx.oxml.ns import qn

        srgb = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
        if srgb is None:
            return
        for existing in srgb.findall(qn("a:alpha")):
            srgb.remove(existing)
        node = srgb.makeelement(qn("a:alpha"), {})
        node.set("val", str(int(round(max(0.0, min(1.0, alpha)) * 100000))))
        srgb.append(node)
    except Exception:
        pass


def _drop_theme_style(shape) -> None:
    """
    Remove an autoshape's ``<p:style>``.

    It points at the theme's shadow effect, and LibreOffice honours the
    effectRef even when the effect list is cleared - a hairline rule with a
    soft grey shadow under it reads as a rendering artefact rather than as a
    design element.
    """
    element = getattr(shape, "_element", None)
    if element is None:
        return
    style = element.find(qn("p:style"))
    if style is not None:
        element.remove(style)


def _suppress_bullet(paragraph) -> None:
    """A mark is not a list item; an inherited bullet turns a label into one."""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for existing in pPr.findall(qn(tag)):
            pPr.remove(existing)
    pPr.append(OxmlElement("a:buNone"))


def _autofit_scale(primitive: Primitive, context: PaintContext,
                   base_size: float) -> float:
    """
    How far this primitive's type has to shrink to fit its own box.

    Measured against the primitive's box, not the slide's content band: a mark
    in a quarter-width rail wraps at the rail's width, and measuring anywhere
    else is how copy ends up running out of the region it was composed into.
    """
    if not primitive.autofit or not primitive.runs:
        return 1.0
    blocks = [
        text_fit.TextBlock(text=run.text, font_size_pt=(run.size_pt or base_size))
        for run in primitive.runs
    ]
    # Zero insets, because ``_paint_text`` zeroes the frame's margins. Left at
    # OOXML's defaults, every mark would be judged against a box 7.2pt shorter
    # than the one it is drawn in.
    scale, _ = text_fit.compute_autofit_scale(
        blocks,
        max(1.0, primitive.box.width * 72.0),
        max(1.0, primitive.box.height * 72.0),
        inset_h_pt=0.0, inset_v_pt=0.0,
    )
    return scale


# --------------------------------------------------------------------------
# Measurement
# --------------------------------------------------------------------------

def bounds(primitives: Sequence[Primitive]) -> Optional[PositionConfig]:
    """The rectangle containing every primitive, or None when there are none."""
    boxes = [p.box for p in primitives if p.box.width > 0 and p.box.height > 0]
    boxes += [
        PositionConfig(left=min(p.box.left, p.end[0]), top=min(p.box.top, p.end[1]),
                       width=abs(p.end[0] - p.box.left),
                       height=abs(p.end[1] - p.box.top))
        for p in primitives if p.kind == "line" and p.end is not None
    ]
    if not boxes:
        return None
    left = min(b.left for b in boxes)
    top = min(b.top for b in boxes)
    right = max(b.left + b.width for b in boxes)
    bottom = max(b.top + b.height for b in boxes)
    return PositionConfig(left=left, top=top, width=right - left, height=bottom - top)


# Wrap estimation uses one average glyph width for every face and weight.
# Under-estimating a wrap costs an overlap - a line appears that the caller
# reserved no room for - while over-estimating costs a little white space, so
# the measurement used to *reserve* space is deliberately the conservative
# one. Bold gets more slack because a bold face is the wider of the two.
MEASURE_SLACK = 0.96
MEASURE_SLACK_BOLD = 0.90


def measure_text(text: str, size_pt: float, width_in: float,
                 bold: bool = False) -> float:
    """
    The height a wrapped string needs at ``size_pt`` in ``width_in``.

    Conservative on purpose - see ``MEASURE_SLACK``. This is the number blocks
    lay themselves out from, so an optimistic answer here is a stat callout
    whose figure wraps onto its own caption.
    """
    if not str(text or "").strip() or width_in <= 0 or size_pt <= 0:
        return 0.0
    from app.renderer.apparatus import strip_markers

    slack = MEASURE_SLACK_BOLD if bold else MEASURE_SLACK
    # Markers are four characters of syntax that render as one small digit;
    # measuring them as written would reserve room for text nobody sees.
    lines = text_fit.estimate_wrapped_lines(
        strip_markers(str(text)), size_pt, max(1.0, width_in * 72.0 * slack)
    )
    return lines * size_pt * text_fit.LINE_HEIGHT_RATIO / 72.0


def translated(primitives: Sequence[Primitive], dx: float,
               dy: float) -> list[Primitive]:
    """Move every primitive, keeping ids and styling - used to place a
    compiled block once its final rectangle is known."""
    moved: list[Primitive] = []
    for primitive in primitives:
        box = PositionConfig(left=primitive.box.left + dx, top=primitive.box.top + dy,
                             width=primitive.box.width, height=primitive.box.height)
        end = ((primitive.end[0] + dx, primitive.end[1] + dy)
               if primitive.end is not None else None)
        moved.append(replace(primitive, box=box, end=end))
    return moved


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------

def validate_primitives(raw: object, region_ids: Iterable[str] = ()) -> tuple[
        list[str], list[str]]:
    """
    Check authored primitives before render, returning ``(errors, warnings)``.

    The open layer is only safe because it is checked as closely as the closed
    one. Three failures make a file that opens and is wrong - a mark outside
    the frame it was composed into, a mark bound to a region that does not
    exist, and two labels drawn on top of each other - and all three are
    caught here rather than discovered in the PPTX.
    """
    errors: list[str] = []
    warnings: list[str] = []
    if raw is None:
        return errors, warnings
    if not isinstance(raw, list):
        return [f"'primitives' must be an array, got {type(raw).__name__}"], warnings

    known_regions = {str(r) for r in region_ids}
    seen: set[str] = set()
    text_boxes: list[tuple[str, tuple[float, float, float, float]]] = []

    for index, spec in enumerate(raw):
        where = f"primitive {index}"
        if not isinstance(spec, dict):
            errors.append(f"{where}: must be an object, got {type(spec).__name__}")
            continue

        primitive_id = str(spec.get("id") or "").strip()
        if primitive_id:
            if primitive_id in seen:
                errors.append(f"{where}: duplicate primitive id {primitive_id!r}")
            seen.add(primitive_id)
            where = f"primitive {primitive_id!r}"

        kind = str(spec.get("kind") or "").strip()
        if kind not in PRIMITIVE_KINDS:
            errors.append(
                f"{where}: unknown kind {kind!r}, must be one of "
                f"{', '.join(sorted(PRIMITIVE_KINDS))}"
            )
            continue

        region_id = spec.get("regionId")
        if region_id and str(region_id) not in known_regions:
            errors.append(
                f"{where}: bound to region {str(region_id)!r}, which the slide "
                f"does not declare"
            )

        if kind == "line":
            box = tuple(_fraction(spec.get(key), default) for key, default in
                        (("x1", 0.0), ("y1", 0.0), ("x2", 1.0), ("y2", 0.0)))
            if all(abs(box[i] - box[i + 2]) < 1e-6 for i in (0, 1)):
                errors.append(f"{where}: line has no length")
            for value, name in zip(box, ("x1", "y1", "x2", "y2")):
                if not -CONTAINMENT_TOLERANCE <= value <= 1 + CONTAINMENT_TOLERANCE:
                    errors.append(
                        f"{where}: {name}={value} escapes its frame "
                        f"(fractional geometry must sit within 0..1)"
                    )
            continue

        x, y, w, h = (_fraction(spec.get(key), default) for key, default in
                      (("x", 0.0), ("y", 0.0), ("w", 1.0), ("h", 1.0)))
        if w <= 0 or h <= 0:
            errors.append(f"{where}: has no area (w={w}, h={h})")
        if (x < -CONTAINMENT_TOLERANCE or y < -CONTAINMENT_TOLERANCE
                or x + w > 1 + CONTAINMENT_TOLERANCE
                or y + h > 1 + CONTAINMENT_TOLERANCE):
            errors.append(
                f"{where}: escapes its frame (x={x}, y={y}, w={w}, h={h}; "
                f"fractional geometry must sit within 0..1)"
            )

        if kind == "text":
            runs = spec.get("runs") or ([spec["text"]] if spec.get("text") else [])
            if not any(str(_run_text(run)).strip() for run in runs):
                errors.append(f"{where}: a text mark with no text is an empty box")
            text_boxes.append((where, (x, y, x + w, y + h)))
        elif kind == "image":
            image = spec.get("image")
            if not isinstance(image, dict) or not str(image.get("data") or "").strip():
                errors.append(f"{where}: image marks need an image with a data source")
        elif kind in ("rect", "ellipse"):
            if not spec.get("fill") and not spec.get("stroke"):
                warnings.append(
                    f"{where}: has neither a fill nor a stroke, so it renders as "
                    f"nothing"
                )

    for i in range(len(text_boxes)):
        name_i, box_i = text_boxes[i]
        for j in range(i + 1, len(text_boxes)):
            name_j, box_j = text_boxes[j]
            share = _overlap_share(box_i, box_j)
            if share > TEXT_OVERLAP_TOLERANCE:
                errors.append(
                    f"{name_i} and {name_j} overlap by {share:.0%} of the smaller "
                    f"one - one of them will be drawn over the other"
                )
    return errors, warnings


def _run_text(run: Any) -> str:
    if isinstance(run, str):
        return run
    if isinstance(run, dict):
        return str(run.get("text", ""))
    return str(run or "")


def _overlap_share(a: tuple[float, float, float, float],
                   b: tuple[float, float, float, float]) -> float:
    left, top = max(a[0], b[0]), max(a[1], b[1])
    right, bottom = min(a[2], b[2]), min(a[3], b[3])
    if right <= left or bottom <= top:
        return 0.0
    intersection = (right - left) * (bottom - top)
    smaller = min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1]))
    return intersection / smaller if smaller > 0 else 0.0


# --------------------------------------------------------------------------

def _unique(candidate: str, used: set[str]) -> str:
    if candidate not in used:
        used.add(candidate)
        return candidate
    suffix = 2
    while f"{candidate}-{suffix}" in used:
        suffix += 1
    unique = f"{candidate}-{suffix}"
    used.add(unique)
    return unique


def _fraction(value: object, default: float) -> float:
    try:
        number = float(value)                        # type: ignore[arg-type]
    except (TypeError, ValueError):
        return default
    if math.isnan(number) or math.isinf(number):
        return default
    return number


def _positive(value: object, default: float) -> float:
    try:
        number = float(value)                        # type: ignore[arg-type]
    except (TypeError, ValueError):
        return default
    return number if number > 0 and not math.isinf(number) else default


def _positive_or_none(value: object) -> Optional[float]:
    try:
        number = float(value)                        # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    return number if number > 0 and not math.isinf(number) else None


def _hex_or_none(value: object) -> Optional[str]:
    text = str(value or "").strip()
    if not text:
        return None
    if not text.startswith("#"):
        text = f"#{text}"
    return text.upper() if colors.parse_hex(text) else None
