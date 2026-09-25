"""
A table as a native PowerPoint table, on either render path.

A table is data a user edits after the deck is delivered - a figure revised,
a row added - so it is never drawn as a picture. The composed path used to
rasterise every table to a PNG and place it in its frame, where a short table
in a tall frame was stretched until it read as a condensed face, and none of
it could be edited. Both paths now draw a table here.

Fitting is done the way a typesetter does it: `table_layout.plan_table`
sizes the columns from their content, steps the type down until the rows fit
the band, and wraps cells - never by scaling.
"""

from __future__ import annotations

from typing import Literal, Optional

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app import colors
from app.renderer import apparatus as apparatus_markers
from app.renderer import deck_theme, table_layout, table_style
from app.renderer import primitives as prim
from app.renderer.default_template import _hex_to_rgbcolor
from app.renderer.layout import PositionConfig

Placement = Literal["optical", "top"]


def add_native_table(slide, table_data: dict, pos: PositionConfig,
                     design_system: Optional[dict], paint_context, *,
                     id_prefix: str, placement: Placement = "optical",
                     ) -> Optional[float]:
    """Draw this table natively in `pos`; return its height in inches.

    ``placement`` is where a table shorter than its band sits: "optical" hangs
    it a third of the way into the surplus, which reads as hanging from a
    title on a template slide; "top" sets it at the top of its frame, which is
    where a composed page's frame already is.
    """
    rows_data = table_data.get("rows", [])
    if not rows_data:
        return None
    n_rows = len(rows_data)
    n_cols = max((len(r) for r in rows_data), default=0)
    if n_cols == 0:
        return None

    style = table_style.resolve_table_style(design_system)
    header_row = bool(table_data.get("headerRow"))
    layout = table_layout.plan_table(
        [[str(c) for c in row] for row in rows_data],
        has_header=header_row,
        region_width_in=pos.width,
        region_height_in=pos.height,
        family=style.font_name,
    )
    if layout is None:
        return None

    # Sit the table in its band rather than filling it: a short table
    # stretched to a tall region is what left a dead strip above and below
    # with rows an inch apart.
    surplus = max(0.0, pos.height - layout.height_in)
    top = pos.top + (surplus / 3.0 if placement == "optical" else 0.0)

    graphic_frame = slide.shapes.add_table(
        n_rows, n_cols,
        Inches(pos.left), Inches(top),
        Inches(pos.width), Inches(layout.height_in),
    )
    table = graphic_frame.table

    for index, fraction in enumerate(layout.col_fractions):
        table.columns[index].width = Inches(pos.width * fraction)
    for index, height in enumerate(layout.row_heights_in):
        table.rows[index].height = Inches(height)

    # python-pptx applies a banded PowerPoint table style by default, which
    # paints its own accent1 header over the fills set below.
    table.first_row = header_row
    table.horz_banding = False

    if style.treatment == "hairline":
        _rule_rows(slide, layout, style, header_row,
                   PositionConfig(left=pos.left, top=top, width=pos.width,
                                  height=layout.height_in),
                   paint_context, id_prefix)

    header_rgb = _hex_to_rgbcolor(style.header_bg)
    header_fg_rgb = _hex_to_rgbcolor(style.header_fg)
    band_rgb = _hex_to_rgbcolor(style.band_bg)
    row_rgb = _hex_to_rgbcolor(style.row_bg)
    body_fg_rgb = _hex_to_rgbcolor(deck_theme.BODY_ON_LIGHT)

    for r_idx, row in enumerate(rows_data):
        is_header = header_row and r_idx == 0
        is_total = r_idx == layout.total_row_index
        for c_idx in range(n_cols):
            cell = table.cell(r_idx, c_idx)
            raw_cell = str(row[c_idx]) if c_idx < len(row) else ""
            cell.text = apparatus_markers.strip_markers(raw_cell)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if style.treatment == "hairline":
                # Rows separated by rules on the slide's own ground rather than
                # banded blocks of colour, as the reference hairline tables are.
                cell.fill.background()
                if is_header and style.header_filled:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = header_rgb
            else:
                cell.fill.solid()
                if is_header:
                    cell.fill.fore_color.rgb = header_rgb
                elif is_total:
                    # A tint of the header sets a summary row apart without
                    # introducing another colour.
                    cell.fill.fore_color.rgb = _hex_to_rgbcolor(
                        colors.mix(style.band_bg, style.header_bg, 0.22))
                elif r_idx % 2 == 1:
                    cell.fill.fore_color.rgb = band_rgb
                else:
                    cell.fill.fore_color.rgb = row_rgb

            # Figures right-align so their digits line up and can be compared
            # down the column; prose stays left.
            alignment = (PP_ALIGN.RIGHT if c_idx in layout.numeric_columns
                         else PP_ALIGN.LEFT)
            size = Pt(layout.header_pt if is_header else layout.body_pt)
            color = (header_fg_rgb if is_header and style.header_filled
                     else body_fg_rgb)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = alignment
                for run in paragraph.runs:
                    run.font.size = size
                    run.font.bold = is_header or is_total
                    run.font.name = style.font_name
                    run.font.color.rgb = color
                if apparatus_markers.has_marker(raw_cell) and paragraph.runs:
                    _mark_cell(paragraph, raw_cell, style.font_name, size,
                               header_fg_rgb if is_header else body_fg_rgb,
                               is_header or is_total)
    return layout.height_in


def _mark_cell(paragraph, text: str, font_name: str, size, color_rgb,
               bold: bool) -> None:
    """Set a table cell's footnote markers as real superscripts."""
    parts = apparatus_markers.split_markers(text)
    if not parts:
        return
    first = paragraph.runs[0]
    first.text = parts[0][0]
    if parts[0][1] is not None:
        apparatus_markers.apply_superscript(first)
    for part_text, marker in parts[1:]:
        run = paragraph.add_run()
        run.text = part_text
        run.font.name = font_name
        run.font.size = size
        run.font.bold = bold
        run.font.color.rgb = color_rgb
        if marker is not None:
            apparatus_markers.apply_superscript(run)


def _rule_rows(slide, layout, style, header_row: bool, pos: PositionConfig,
               paint_context, id_prefix: str) -> None:
    """Draw a hairline table's own rules.

    python-pptx has no cell-border API, so the rules are marks: a heavier one
    under the header, a light one under every other row.
    """
    top = pos.top
    marks = []
    for index, height in enumerate(layout.row_heights_in):
        top += height
        if index == len(layout.row_heights_in) - 1:
            break
        weight = (style.header_rule_weight_pt
                  if header_row and index == 0 else style.rule_weight_pt)
        marks.append(prim.Primitive(
            id=f"{id_prefix}-table-rule-{index}", kind="rect",
            box=PositionConfig(left=pos.left, top=top, width=pos.width,
                               height=max(0.008, weight / 72.0)),
            fill=style.border,
        ))
    prim.paint_all(slide, marks, paint_context)
