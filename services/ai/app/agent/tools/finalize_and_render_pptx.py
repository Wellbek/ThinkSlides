"""
Tool: finalize_and_render_pptx
Takes validated slide JSON array and a template, then renders the final PPTX
by cloning real layout XML from the template, filling placeholder elements
without touching rPr formatting attributes, and repacking.
Also generates preview images for each slide.
"""

from __future__ import annotations
from app.agent.registry import ToolRegistry, ToolResult, ToolParameter
from app.agent.tools._json_extract import coerce_json_arg
from app.agent import run_context
from app.design.run import Unplaceable
from app.design.vocabulary import PLAIN_ID
from app.renderer.citations import apply_source_citations
from app.titles import slide_title
from app.renderer.pptx_renderer import PPTXRenderer
from app.renderer.preview_generator import PreviewGenerator
from app.observability.logging import get_logger
from typing import NamedTuple
import time
import os
import uuid

logger = get_logger(__name__)


class Composed(NamedTuple):
    """What laying a deck out in its pack produced, besides the file."""

    path: str
    # Per source slide, the archetypes its pages were placed on.
    placed: dict
    # Per source slide, copy still over its frames' budget after repair.
    over_budget: dict
    # Slides no other arrangement could carry.
    unmovable: list
    # Pages and slides set below the readable floor (#368).
    below_floor: dict
    # Rendered page numbers, from 1, that are covers.
    covers: frozenset


class FinalizeAndRenderPptxTool:
    """Render the final PPTX from validated slide JSON and a template."""

    @classmethod
    def register(cls, registry: ToolRegistry) -> None:
        registry.register_tool(
            name="finalize_and_render_pptx",
            description=(
                "Render the final PPTX presentation from validated slide JSON. "
                "Clones real layout XML from the template, fills placeholders "
                "without touching rPr formatting, and produces a downloadable file. "
                "Returns file path and metadata."
            ),
            fn=cls.run,
            parameters=[
                ToolParameter(
                    name="slides_json",
                    type="string",
                    description="JSON array of validated slide objects",
                ),
                ToolParameter(
                    name="template_path",
                    type="string",
                    description="Path to the PPTX/POTX template file",
                    required=False,
                ),
                ToolParameter(
                    name="presentation_metadata",
                    type="string",
                    description="JSON with title, author, subject",
                    required=False,
                ),
                ToolParameter(
                    name="session_id",
                    type="string",
                    description="The chat session ID for associating the generated file",
                ),
                ToolParameter(
                    name="color_theme",
                    type="string",
                    description=(
                        "Color palette to use: either one of professional, warm, monochrome, "
                        "bold, nature, OR a specific brand hex color like '#00704A' - must "
                        "match whatever was passed to extract_template_design_system earlier "
                        "in this run. Only takes effect when no template_path is given - an "
                        "uploaded template's own colors are always preserved."
                    ),
                    required=False,
                ),
                ToolParameter(
                    name="design_system_json",
                    type="string",
                    description=(
                        "JSON design system from extract_template_design_system, used to style "
                        "generated tables/charts (header/band colors, fonts). Auto-supplied by "
                        "the agent loop - do not set manually."
                    ),
                    required=False,
                ),
            ],
        )

    @classmethod
    async def run(
        cls,
        slides_json: str,
        session_id: str,
        template_path: str = "",
        presentation_metadata: str = "{}",
        color_theme: str = "",
        design_system_json: str = "{}",
    ) -> ToolResult:
        start = time.time()
        try:
            # coerce_json_arg, not a bare json.loads: a tool-calling model may
            # send a native JSON value here despite the declared "string"
            # type, which raises TypeError rather than JSONDecodeError. A
            # coercion failure on slides falls through to the existing
            # "No slides provided" check below, which already fails loudly.
            slides = coerce_json_arg(slides_json, [])
            metadata = coerce_json_arg(presentation_metadata, {})
            design_system = coerce_json_arg(design_system_json, {})

            if not slides:
                return ToolResult(
                    success=False,
                    error="No slides provided for rendering",
                    duration_ms=int((time.time() - start) * 1000),
                )

            # Determine output paths
            output_dir = os.environ.get("OUTPUT_DIR", "/app/output")
            os.makedirs(output_dir, exist_ok=True)

            output_id = str(uuid.uuid4())
            output_path = os.path.join(output_dir, f"{output_id}.pptx")
            markdown_path = os.path.join(output_dir, f"{output_id}.md")

            # Put the citation trail in the file itself, before anything is
            # rendered or written, so the markdown fallback carries it too.
            # A run with the toggle off, or with nothing to cite, passes
            # straight through - see renderer/citations.py.
            # How many pages the model itself designed. Anything past this in
            # the rendered deck was injected here rather than by the agent, and
            # the outline below marks it so - see _rendered_outline.
            designed_count = len(slides)

            slides, cited_sources = apply_source_citations(
                slides,
                run_context.get_research_findings(),
                enabled=run_context.get_include_sources(),
            )

            # Generate fallback markdown first (so we always have something to give the user)
            fallback_markdown = cls._generate_markdown_fallback(slides, metadata)
            with open(markdown_path, "w", encoding="utf-8") as f:
                f.write(fallback_markdown)

            # Retry logic for rendering (3 attempts)
            result_path = None
            last_error = None
            max_retries = 3

            # A run composing against a design language pack lays its pages
            # out rather than filling placeholders. Both paths end in a real
            # .pptx with live, editable text; what differs is where the
            # geometry comes from. An uploaded template has no pack, and
            # keeps the XML-clone path that makes it faithful.
            pack = run_context.get_pack()
            # The user's deck-wide operations (#376). A recolour is applied to
            # the pack and the design system before anything is drawn, so
            # what the render paints from them - charts included, which reach
            # the file as pictures - is in the new colour; the same operation
            # then repaints the file itself after the render, for whatever
            # the template drew. Style is read before the recolour, so the
            # file pass knows which colour it replaces.
            operations, operation_problems = cls._deck_operations()
            style_pack, style_system = pack, design_system
            pack, design_system = cls._recoloured(pack, design_system, operations)
            covers: frozenset = frozenset(
                index + 1 for index, slide in enumerate(slides)
                if isinstance(slide, dict) and cls._is_cover(slide, index))

            # A run with no template and no pack is about to render stock
            # Office placeholders, which is never what anyone asked for -
            # it is what a template-less run looked like before design
            # languages existed. This was silent for the whole life of
            # issue #290: the pack was being dropped at the outline-review
            # pause and every consulted run rendered plain bullets, with
            # nothing in the logs saying so because the same branch is a
            # legitimate path for an uploaded template. Said out loud now,
            # so the next time it happens it is visible on the first run
            # rather than after a user compares their deck to a preview.
            if pack is None and not (
                template_path and os.path.exists(template_path)
            ):
                logger.warning(
                    "rendering_without_a_design_language",
                    session_id=session_id,
                    slides=len(slides),
                    detail=("no pack and no uploaded template - this deck "
                            "will render on stock Office placeholders"),
                )

            # What each slide landed on, returned rather than left in a
            # context variable: the loop awaits this tool as its own task,
            # which runs in a *copy* of the context, so anything set here
            # never reached the loop. Every "replace" verdict was answered
            # with an empty list to avoid, and the page came back on the
            # arrangement it had been flagged for (#344).
            placed: dict[int, list[str]] = {}
            # Copy still too long for its frames after repair, per slide, with
            # how much fits - the writer's budget (#348).
            over_budget: dict[int, dict] = {}
            unmovable: list[int] = []
            # Copy set below the readable floor because the page count was
            # fixed (#368): which pages, and how small, per source slide.
            below_floor: dict = {"pages": [], "slides": {}}
            # Which renderer the deck actually went through, so a run can be
            # checked for it rather than only logged: "composed" in its design
            # language, "template" on an uploaded template, "fallback" when
            # the pack could not place the deck, "stock" on neither.
            render_path = "template" if pack is None and template_path and \
                os.path.exists(template_path) else "stock"
            for attempt in range(1, max_retries + 1):
                try:
                    if pack is not None:
                        try:
                            composed = cls._render_composed(
                                slides, pack, metadata, output_path,
                                design_system,
                            )
                            result_path, placed, over_budget, unmovable, below_floor, covers = composed
                            render_path = "composed"
                        except Unplaceable as exc:
                            # Material no archetype can carry, even split.
                            # Retrying is pointless - the answer is
                            # deterministic - and shipping nothing is worse
                            # for the user than shipping a plainer deck, so
                            # this drops to the template renderer. Loudly:
                            # a pack that cannot carry ordinary material is a
                            # defect in the pack, and it has to be visible.
                            logger.warning(
                                "pack_could_not_place_deck",
                                pack=pack.id, slides=len(slides),
                                error=str(exc),
                            )
                            pack = None
                            style_pack = None
                            render_path = "fallback"
                    if pack is None:
                        renderer = PPTXRenderer(
                            template_path=template_path if template_path and os.path.exists(template_path) else None,
                            color_theme=color_theme or None,
                            design_system=design_system,
                        )
                        result_path = renderer.render(
                            slides=slides,
                            metadata=metadata,
                            output_path=output_path,
                        )

                    # Verify the file was created
                    if not os.path.exists(result_path) or os.path.getsize(result_path) == 0:
                        raise Exception("PPTX file was not created or is empty")

                    logger.info(
                        "pptx_render_success",
                        session_id=session_id,
                        slide_count=len(slides),
                        attempt=attempt,
                        max_retries=max_retries,
                    )
                    break

                except Exception as e:
                    last_error = e
                    logger.warning(
                        "pptx_render_attempt_failed",
                        session_id=session_id,
                        attempt=attempt,
                        max_retries=max_retries,
                        error=str(e),
                    )

                    # Clean up failed attempt
                    if os.path.exists(output_path):
                        os.remove(output_path)

                    if attempt < max_retries:
                        time.sleep(1)  # Brief pause before retry

            # If all retries failed, return the fallback
            if result_path is None:
                logger.error(
                    "pptx_render_all_retries_failed",
                    session_id=session_id,
                    error=str(last_error),
                )

                # Return the markdown fallback as the primary result
                return ToolResult(
                    success=False,
                    error=f"PPTX rendering failed after {max_retries} attempts: {str(last_error)}. A markdown fallback has been generated.",
                    data={
                        "fallbackMarkdown": fallback_markdown,
                        "fallbackMarkdownPath": markdown_path,
                        "fallbackSlidesJson": slides_json,
                        "slides": slides,
                        "metadata": metadata,
                        "sessionId": session_id,
                    },
                    duration_ms=int((time.time() - start) * 1000),
                )

            operation_report = operation_problems + cls._apply_deck_operations(
                result_path, operations, style_pack, style_system, covers,
                slides=len(slides))

            file_size = os.path.getsize(result_path) if os.path.exists(result_path) else 0

            # Generate preview images
            preview_gen = PreviewGenerator()
            preview_urls = preview_gen.generate_previews(result_path, session_id)

            # Individual chart/image assets, extracted from the rendered
            # deck itself rather than intercepted mid-render - this reads
            # already-embedded pictures back out with python-pptx, so it
            # never touches the render path's own XML handling (issue #187).
            # Best-effort only: a failure here must never turn a successful
            # render into a failed one.
            assets = cls._persist_slide_assets(result_path, slides, session_id)

            duration_ms = int((time.time() - start) * 1000)
            logger.info(
                "pptx_rendered",
                session_id=session_id,
                slide_count=len(slides),
                file_size=file_size,
                preview_count=len(preview_urls),
                asset_count=len(assets),
                duration_ms=duration_ms,
            )

            return ToolResult(
                success=True,
                data={
                    "filePath": result_path,
                    "fileName": metadata.get("title", "presentation") + ".pptx",
                    "fileSizeBytes": file_size,
                    "slideCount": len(slides),
                    "citedSources": cited_sources,
                    "sessionId": session_id,
                    "previewUrls": preview_urls,
                    "renderedOutline": cls._rendered_outline(slides, designed_count),
                    "fallbackMarkdownPath": markdown_path,
                    "assets": assets,
                    "placedArrangements": placed,
                    "renderPath": render_path,
                    "copyOverBudget": over_budget,
                    # Slides a `replace` verdict could only hand back as they
                    # are: no other arrangement carries them (#376).
                    "unmovableSlides": unmovable,
                    "belowFloor": below_floor,
                    # What each deck-wide operation made of this deck: the
                    # pages it landed on and the ones it had to leave (#376).
                    "deckOperations": operation_report,
                    # Slides set on the pack's plain page because none of its
                    # own arrangements could carry them (#322).
                    "plainPages": sorted(
                        index for index, archetypes in placed.items()
                        if PLAIN_ID in archetypes),
                },
                duration_ms=duration_ms,
            )
        except Exception as e:
            logger.error("pptx_render_failed", error=str(e))
            return ToolResult(success=False, error=str(e), duration_ms=int((time.time() - start) * 1000))

    @staticmethod
    def _rendered_outline(slides: list[dict], designed_count: int) -> list[dict]:
        """
        One entry per page of the deck that was actually rendered.

        The agent's own record of the deck stops at the slides it designed, but
        the file does not: ``apply_source_citations`` appends a Sources appendix
        after the model is finished, and previews are rasterised from every page
        of the result. Without this the UI receives a preview image for those
        pages and has nothing to attach it to, so they vanish from the preview
        while remaining in the downloaded file.

        Deliberately just the index, the headline and whether the page was
        injected. The preview shows the rendered image; a title is all the
        surrounding chrome needs, and shipping whole slide bodies here would put
        another copy of the deck - image bytes included - into the tool result.
        """
        return [
            {
                "slideIndex": (
                    slide.get("slideIndex")
                    if isinstance(slide, dict) and isinstance(slide.get("slideIndex"), int)
                    else position
                ),
                "title": slide_title(slide) if isinstance(slide, dict) else "",
                "injected": position >= designed_count,
            }
            for position, slide in enumerate(slides)
        ]

    @classmethod
    def _persist_slide_assets(cls, result_path: str, slides: list[dict], session_id: str) -> list[dict]:
        """Write each slide's embedded chart/image as its own downloadable
        asset (issue #187), so the workspace UI can list and replace them
        individually - separate from the single all-in-one PPTX row.

        Reads the pictures back out of the already-rendered deck with
        python-pptx rather than intercepting bytes mid-render, so this never
        touches the renderer's own XML-cloning logic. Every picture the
        render pipeline itself places is named "ts-asset-chart" or
        "ts-asset-image" (see pptx_renderer.py's _add_chart/_add_image) -
        that is what gets picked up here, not "the Nth picture on the
        slide". Position-based pairing mispairs the moment a template
        layout carries its own static pictures (a logo, say) ahead of the
        content pictures in document order; the name marker sidesteps that
        entirely instead of trying to out-guess it, and as a side effect
        never surfaces a template's own artwork as a "replaceable" asset.
        """
        try:
            from pptx import Presentation

            output_dir = os.environ.get("OUTPUT_DIR", "/app/output")
            assets_dir = os.path.join(output_dir, str(session_id), "assets")
            os.makedirs(assets_dir, exist_ok=True)

            prs = Presentation(result_path)
            assets: list[dict] = []
            asset_types = {"ts-asset-chart": "chart", "ts-asset-image": "image"}

            for slide_index, pptx_slide in enumerate(prs.slides):
                position = 0
                for shape in pptx_slide.shapes:
                    # shape_type/name access itself can raise (a shape
                    # python-pptx doesn't fully model - embedded OLE,
                    # SmartArt - can throw NotImplementedError just reading
                    # its type), so this has to be inside the per-shape
                    # guard too, not just the picture-extraction below it.
                    # One bad shape must cost that one shape, not silently
                    # abort every asset on every remaining slide.
                    try:
                        file_type = asset_types.get(shape.name)
                        if file_type is None:
                            continue
                        image = shape.image
                        ext = image.ext or "png"
                        position += 1
                        asset_id = str(uuid.uuid4())
                        asset_path = os.path.join(assets_dir, f"{asset_id}.{ext}")
                        with open(asset_path, "wb") as f:
                            f.write(image.blob)
                        assets.append({
                            "filePath": asset_path,
                            "fileName": f"slide-{slide_index + 1}-{file_type}-{position}.{ext}",
                            "fileSizeBytes": len(image.blob),
                            "fileType": file_type,
                            "slideIndex": slide_index,
                        })
                    except Exception as shape_error:
                        # An unusual shape python-pptx can't fully read, or
                        # an image format it can't decode, should not lose
                        # every other asset on the deck.
                        logger.warning(
                            "slide_asset_extract_failed",
                            session_id=session_id,
                            slide_index=slide_index,
                            error=str(shape_error),
                        )

            return assets
        except Exception as e:
            # Individual assets are a bonus, not the deliverable - the PPTX
            # itself has already rendered successfully by the time this
            # runs, and that must never be undone by an asset-extraction bug.
            logger.warning("slide_asset_persist_failed", session_id=session_id, error=str(e))
            return []

    @staticmethod
    def _deck_operations() -> tuple[list, list[dict]]:
        """The run's recorded operations, parsed; and one report for each
        that no longer parses, so a bad record costs that operation only."""
        from app.renderer import deck_ops

        parsed, problems = [], []
        for raw in run_context.get_deck_operations():
            try:
                parsed.append(deck_ops.parse(raw))
            except ValueError as exc:
                problems.append({"id": str((raw or {}).get("id") or ""),
                                 "kind": str((raw or {}).get("kind") or ""),
                                 "problem": str(exc)})
        return parsed, problems

    @staticmethod
    def _recoloured(pack, design_system: dict, operations: list):
        """The pack and design system with every recolour applied."""
        from app.renderer.deck_ops import Recolour

        recolours = [op for op in operations if isinstance(op, Recolour)]
        if not recolours:
            return pack, design_system
        # The design system names the roles by their palette names.
        palette_name = {"accent": "primary", "accent_alt": "accent"}
        system = dict(design_system or {})
        palette = dict(system.get("palette") or {})
        entries = [dict(entry) for entry in palette.get("colors") or []
                   if isinstance(entry, dict)]
        for op in recolours:
            if pack is not None:
                pack = pack.model_copy(update={"colors": pack.colors.model_copy(
                    update={op.role: op.to})})
            name = palette_name[op.role]
            for entry in entries:
                if entry.get("name") == name:
                    entry["value"] = op.to
            if entries and not any(entry.get("name") == name for entry in entries):
                entries.append({"name": name, "value": op.to})
        if entries:
            palette["colors"] = entries
            system["palette"] = palette
        return pack, system

    @staticmethod
    def _composed_covers(slides: list, doc) -> frozenset:
        """The pages of a composed deck a reader takes for its cover.

        A page set as a cover, and the first page of the slide the deck opens
        with when that slide is a cover by composition's own rule - including
        when composition set it on a content page because it carries a page
        of copy, which leaves it the title page all the same.
        """
        from app.design.run import page_role

        ordered = sorted((s for s in slides if isinstance(s, dict)),
                         key=lambda s: int(s.get("slideIndex") or 0))
        opens_on_cover = bool(ordered) and page_role(ordered[0], 0, len(ordered)) == "cover"
        covers = {number for number, page in enumerate(doc.pages, start=1)
                  if page.role == "cover"}
        if opens_on_cover:
            first = next((number for number, page in enumerate(doc.pages, start=1)
                          if page.provenance.source_slide == 0), None)
            if first is not None:
                covers.add(first)
        return frozenset(covers)

    @staticmethod
    def _is_cover(slide: dict, index: int) -> bool:
        """Whether a slide the stock or template path renders is a cover:
        the first slide, when it says it is a title page."""
        if index != 0:
            return False
        said = " ".join(str(slide.get(key) or "") for key in
                        ("role", "slideType", "layoutName", "layoutId")).lower()
        return any(word in said for word in ("cover", "title"))

    @staticmethod
    def _deck_style(pack, design_system: dict, covers: frozenset):
        """The deck's own colours, face and margin, for an operation to wear."""
        from app.renderer.deck_ops import DeckStyle

        if pack is not None:
            margin = pack.grid.margin
            small = next((pack.type_roles[name].size_pt for name in
                          ("caption", "footnote", "folio") if name in pack.type_roles), 10.0)
            return DeckStyle(
                ink=pack.colors.ink, ink_reversed=pack.colors.ink_reversed,
                muted=pack.colors.ink_muted, accent=pack.colors.accent,
                accent_alt=pack.colors.accent_alt, ground=pack.colors.ground,
                face=pack.body_face, small_pt=small,
                margin_in=min(margin.top, margin.right, margin.bottom, margin.left) or 0.5,
                side_in=min(margin.left, margin.right) or 0.5,
                cover_pages=covers)
        palette = {str(entry.get("name")): str(entry.get("value"))
                   for entry in ((design_system or {}).get("palette") or {}).get("colors") or []
                   if isinstance(entry, dict) and entry.get("value")}
        fonts = ((design_system or {}).get("typography") or {}).get("fonts") or {}
        defaults = DeckStyle()
        return DeckStyle(
            ink=palette.get("text", defaults.ink),
            muted=palette.get("gray", defaults.muted),
            accent=palette.get("primary", defaults.accent),
            accent_alt=palette.get("accent", defaults.accent_alt),
            ground=palette.get("background", defaults.ground),
            face=str(fonts.get("minorFont") or fonts.get("body") or defaults.face),
            cover_pages=covers)

    @classmethod
    def _apply_deck_operations(cls, path: str, operations: list, pack,
                               design_system: dict, covers: frozenset, *,
                               slides: int) -> list[dict]:
        """Apply the run's deck-wide operations to the file just written.

        Never fails the render: a deck without its overlay is still a deck,
        and the report says what went missing.
        """
        if not operations:
            return []
        from app.renderer import deck_ops

        try:
            return deck_ops.apply(path, operations,
                                  cls._deck_style(pack, design_system, covers),
                                  dict(run_context.get_assets()))
        except Exception as exc:   # noqa: BLE001 - the deck itself rendered
            logger.warning("deck_operations_failed", error=str(exc), slides=slides)
            return [{"id": op.id, "kind": op.kind, "problem": str(exc)[:300]}
                    for op in operations]

    @classmethod
    def _render_composed(cls, slides: list, pack, metadata: dict,
                         output_path: str, design_system: dict) -> Composed:
        """Lay the deck out in its pack and render it.

        Returns the written path; per source slide, the archetypes its pages
        were placed on and the copy over its budget; and the slides no other
        arrangement could carry, which a `replace` verdict cannot move.

        Composition, repair and validation all run here, in code. The
        validator this replaces was a tool the model could decline to call,
        and it was called zero times across every observed run.
        """
        from app.agent import run_context
        from app.design.run import compose, copy_over_budget
        from app.renderer.designdoc_renderer import render_doc

        # What the previous pass already tried and validation disliked. Empty
        # on a first render, which is every render until a layout verdict
        # comes back.
        avoid = run_context.get_arrangements_to_avoid()
        doc = compose(slides, pack, title=str(metadata.get("title") or ""),
                      avoid=avoid or None,
                      pages="fixed" if run_context.get_page_count_fixed() else "free")
        logger.info(
            "designdoc_composed",
            pack=pack.id,
            slides=len(slides),
            pages=len(doc.pages),
            archetypes=sorted({page.archetype for page in doc.pages}),
            avoided={str(k): sorted(v) for k, v in (avoid or {}).items()},
        )
        # What each slide actually landed on, so a layout verdict can say
        # "not that one again" without the loop having to re-derive it.
        placed: dict[int, list[str]] = {}
        movable: set[int] = set()
        below_floor: dict = {"pages": [], "slides": {}}
        for number, page in enumerate(doc.pages, start=1):
            if page.provenance.below_floor_pt is not None:
                below_floor["pages"].append(number)
                slide = page.provenance.source_slide
                if slide is not None:
                    below_floor["slides"][slide] = min(
                        below_floor["slides"].get(slide, 99.0),
                        page.provenance.below_floor_pt)
        for page in doc.pages:
            source = page.provenance.source_slide
            if source is not None:
                placed.setdefault(source, []).append(page.archetype)
                if page.provenance.alternatives:
                    movable.add(source)
        return Composed(
            render_doc(doc, pack, output_path, design_system=design_system),
            placed, copy_over_budget(doc, pack), sorted(set(placed) - movable),
            below_floor,
            cls._composed_covers(slides, doc))

    @classmethod
    def _generate_markdown_fallback(cls, slides: list[dict], metadata: dict) -> str:
        """Generate a markdown representation of the slides as fallback."""
        lines = []

        # Header with metadata
        title = metadata.get("title", "Presentation")
        author = metadata.get("author", "")
        subject = metadata.get("subject", "")

        lines.append(f"# {title}\n")
        if author:
            lines.append(f"**Author:** {author}\n")
        if subject:
            lines.append(f"**Subject:** {subject}\n")
        lines.append("\n---\n\n")

        # Add each slide
        for slide in slides:
            slide_idx = slide.get("slideIndex", 0)
            layout_name = slide.get("layoutName", slide.get("layoutId", "Unknown"))

            lines.append(f"## Slide {slide_idx + 1}: {layout_name}\n")

            # Add slide content
            contents = slide.get("contents", [])
            for content in contents:
                if content.get("type") == "text":
                    runs = content.get("runs", [])
                    for run in runs:
                        text = run.get("text", "")
                        is_bold = run.get("bold", False)
                        is_bullet = run.get("bullet", False)
                        level = run.get("level", 0)

                        if text:
                            if is_bullet:
                                indent = "  " * level
                                lines.append(f"{indent}- {text}")
                            elif is_bold:
                                lines.append(f"**{text}**")
                            else:
                                lines.append(text)

            # Add notes
            notes = slide.get("notes", "")
            if notes:
                lines.append(f"\n**Notes:** {notes}")

            lines.append("\n---\n\n")

        # Add footer with instructions
        lines.append("---\n\n")
        lines.append("*This presentation was generated by ThinkSlides. ")
        lines.append("Due to a technical issue with PPTX rendering, this markdown fallback is provided. ")
        lines.append("You can use this content to manually create your presentation in PowerPoint, Google Slides, or any other presentation software.*")

        return "\n".join(lines)